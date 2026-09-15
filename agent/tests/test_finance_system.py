from __future__ import annotations

import hashlib
import json
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app import finance_system
from app.auth import current_user
from app.database import get_db
from app.main import app
from app.models import (
    AuditLog,
    Base,
    BankStatementBatch,
    BankTransaction,
    BankTransactionPurposeCorrection,
    BusinessEntity,
    CashEntry,
    FinanceReceivablePayable,
    FinancialAccount,
    ManagedProject,
    ProjectCashflowPlan,
    User,
)


def _database() -> tuple[Session, dict[str, User]]:
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    users = {
        "finance": User(
            username="finance01", display_name="财务",
            password_hash="x", role="employee",
            organization_role="finance", confidentiality_ceiling="L4",
        ),
        "founder": User(
            username="founder", display_name="叶靖波",
            password_hash="x", role="founder",
            organization_role="management", confidentiality_ceiling="L5",
        ),
        "manager": User(
            username="other-manager", display_name="其他管理",
            password_hash="x", role="founder",
            organization_role="management", confidentiality_ceiling="L5",
        ),
        "business": User(
            username="business01", display_name="业务",
            password_hash="x", role="employee",
            organization_role="business", confidentiality_ceiling="L5",
        ),
    }
    db = Session(engine, expire_on_commit=False)
    db.add_all(users.values())
    db.commit()
    return db, users


def _xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["交易日期", "贷方发生额", "借方发生额", "交易后余额", "对方户名", "摘要", "流水号"])
    sheet.append(["2026-08-01", 10000, None, 50000, "客户A", "项目回款", "S001"])
    sheet.append(["2026-08-02", None, 3000, 47000, "供应商B", "场地费用", "S002"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _project_code_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["交易日期", "贷方发生额", "借方发生额", "交易后余额", "对方户名", "摘要", "流水号"])
    sheet.append(["2026-08-01", 12000, None, 52000, "赛事客户", "CC26B05 项目回款", "CC-IN"])
    sheet.append(["2026-08-02", None, 3500, 48500, "执行供应商", "执行费，项目代码 cc26-b05", "CC-OUT"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _company_xlsx(
    *,
    serial_prefix: str,
    income: Decimal,
    expense: Decimal,
    opening_balance: Decimal,
    income_counterparty: str,
    expense_counterparty: str,
) -> bytes:
    """Build a small, distinguishable statement for company-isolation tests."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["交易日期", "贷方发生额", "借方发生额", "交易后余额", "对方户名", "摘要", "流水号"])
    after_income = opening_balance + income
    closing_balance = after_income - expense
    sheet.append([
        "2026-08-01",
        income,
        None,
        after_income,
        income_counterparty,
        "项目回款",
        f"{serial_prefix}-IN",
    ])
    sheet.append([
        "2026-08-02",
        None,
        expense,
        closing_balance,
        expense_counterparty,
        "项目支出",
        f"{serial_prefix}-OUT",
    ])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _beijing_bank_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["明细下载"])
    sheet.append(["账号", "20000100355200177786862"])
    sheet.append([
        "序号", "交易时间", "币种", "借方发生额", "贷方发生额", "余额",
        "对手方", "对手方账号", "对手方银行", "收款人", "收款账号",
        "用途", "备注", "凭证号", "摘要", "流水号", "操作员",
    ])
    sheet.append([
        1, "2026-07-01 21:09:45", "人民币", "-", "17,659.60", "1,286,151.51",
        "上海哔哩哔哩电竞信息科技有限公司", "121928339410201", "招商银行",
        "京奥电竞（北京）科技有限公司", "20000100355200177786862",
        "服务费", "网银清算，贷记来帐79311262", "-", "本系统转帐",
        "ASOB001000535425969", "BOBQZY",
    ])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _multi_sheet_minsheng_xlsx() -> bytes:
    """Model the query-cover + detail-sheet structure used by Minsheng."""
    workbook = Workbook()
    cover = workbook.active
    cover.title = "查询条件"
    cover.append(["账户名称", "北京王牌猎豹电竞科技文化发展有限公司"])
    cover.append(["账号", "655984649"])
    cover.append(["起始日期", "2026-01-01"])
    cover.append(["截至日期", "2026-01-31"])
    cover.append([])
    cover.append(["借方累计笔数", 1, "借方累计发生额", 4000])
    cover.append(["贷方累计笔数", 1, "贷方累计发生额", 1000000])

    detail = workbook.create_sheet("交易明细")
    for index in range(35):
        detail.append([f"查询说明 {index + 1}"])
    detail.append([
        "交易日期", "交易时间", "借贷标志", "交易金额", "账户余额",
        "收/付方名称", "收/付方账号", "交易说明", "客户附言", "交易参考号",
    ])
    detail.append([
        "2026-01-05", "16:21:00", "借", 4000, 268056.01,
        "杜萌萌", "6214680050815388", "Cc2609办公室912刷漆", "本系统转账", "M001",
    ])
    detail.append([
        "2026-01-06", "09:15:00", "贷", 1000000, 1268056.01,
        "项目客户", "123456789", "项目回款", "", "M002",
    ])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _cmb_xlsx() -> bytes:
    """Model a CMB export whose detail header follows account metadata."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "交易查询"
    sheet.append(["交易查询"])
    sheet.append(["账户名称", "上海劲腾豹跃文化传播有限责任公司"])
    sheet.append(["账号", "121991350810001"])
    for index in range(32):
        sheet.append([f"打印条件 {index + 1}"])
    sheet.append([
        "交易日", "交易时间", "借方金额", "贷方金额", "账面余额",
        "收款方名称", "收款方账号", "交易摘要", "银行参考号",
    ])
    sheet.append([
        "2026-03-01", "08:30:00", None, 250000, 300000,
        "赛事客户", "88880001", "赛事服务收入", "CMB001",
    ])
    sheet.append([
        "2026-03-02", "10:20:00", 50000, None, 250000,
        "执行供应商", "88880002", "Cc2610活动执行费", "CMB002",
    ])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _split_header_bank_xlsx() -> bytes:
    """Model banks that split grouped labels and RMB units over two rows."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "交易明细"
    sheet.append(["交易", "交易", "金额（人民币元）", "金额（人民币元）", "账户", "对方信息", "对方信息", "交易"])
    sheet.append(["日期", "时间", "借方发生额", "贷方发生额", "余额", "户名", "账号", "用途"])
    sheet.append(["2026-02-01", "09:30:00", 750, None, 9000, "杜萌萌", "62220001", "项目物料"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _analysis_warning_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["交易日期", "贷方发生额", "借方发生额", "交易后余额", "对方户名", "摘要", "流水号"])
    sheet.append(["2026-08-01 10:00:00", 150000, None, 300000, "大额客户", "项目回款", "W001"])
    sheet.append(["2026-08-02 10:00:00", None, 8888.88, 291111.12, "同一供应商", "服务费", "W002"])
    sheet.append(["2026-08-02 11:00:00", None, 8888.88, 282222.24, "同一供应商", "服务费", "W003"])
    sheet.append(["2026-08-03 10:00:00", 100, None, 282322.24, None, None, "W004"])
    sheet.append(["2026-08-04 10:00:00", -50, None, 282272.24, "异常对方", "负数测试", "W005"])
    sheet.append(["2026-08-05 10:00:00", None, 100, 999, "余额异常对方", "余额测试", "W006"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_beijing_bank_column_priority_summary_and_timezone() -> None:
    rows, errors = finance_system.parse_statement(_beijing_bank_xlsx(), ".xlsx")
    assert errors == 0
    assert len(rows) == 1
    item = rows[0]
    assert item["transacted_at"].astimezone(finance_system.SHANGHAI_ZONE).isoformat() == "2026-07-01T21:09:45+08:00"
    assert item["counterparty"] == "上海哔哩哔哩电竞信息科技有限公司"
    assert item["counterparty_account"] == "121928339410201"
    assert item["summary"] == "服务费 · 网银清算，贷记来帐79311262 · 本系统转帐"
    assert item["business_note"] == "服务费"
    assert item["serial"] == "ASOB001000535425969"


def test_multisheet_minsheng_and_cmb_templates_are_parsed() -> None:
    minsheng, minsheng_errors = finance_system.parse_statement(
        _multi_sheet_minsheng_xlsx(), ".xlsx"
    )
    assert minsheng_errors == 0
    assert len(minsheng) == 2
    assert minsheng[0]["expense"] == Decimal("4000.00")
    assert minsheng[0]["income"] == Decimal("0")
    assert minsheng[0]["counterparty"] == "杜萌萌"
    assert minsheng[0]["business_note"] == "Cc2609办公室912刷漆"
    assert minsheng[0]["project_reference"] == "Cc2609"
    assert minsheng[0]["transacted_at"].astimezone(
        finance_system.SHANGHAI_ZONE
    ).isoformat() == "2026-01-05T16:21:00+08:00"

    cmb, cmb_errors = finance_system.parse_statement(_cmb_xlsx(), ".xlsx")
    assert cmb_errors == 0
    assert len(cmb) == 2
    assert cmb[0]["income"] == Decimal("250000.00")
    assert cmb[0]["counterparty"] == "赛事客户"
    assert cmb[0]["business_note"] == "赛事服务收入"
    assert cmb[1]["expense"] == Decimal("50000.00")
    assert cmb[1]["project_reference"] == "Cc2610"

    split, split_errors = finance_system.parse_statement(
        _split_header_bank_xlsx(), ".xlsx"
    )
    assert split_errors == 0
    assert len(split) == 1
    assert split[0]["expense"] == Decimal("750.00")
    assert split[0]["counterparty"] == "杜萌萌"
    assert split[0]["business_note"] == "项目物料"


def test_generic_reimbursement_metadata_does_not_fake_business_purpose() -> None:
    raw = (
        "渠道流水号：ANET001##业务细类：代发##协议编号：2503"
        "##对公回单摘要：网银报销 网银报销##备注："
    )
    assert finance_system._meaningful_business_note(raw, "本系统转帐") is None


def _configure(monkeypatch, tmp_path, db: Session, user: User) -> TestClient:
    monkeypatch.setattr(
        finance_system,
        "settings",
        SimpleNamespace(
            data_dir=tmp_path / "agent-data",
            knowledge_root=tmp_path / "knowledge",
            inbox_max_file_bytes=10 * 1024 * 1024,
        ),
    )

    def override_db():
        yield db

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[current_user] = lambda: user
    return TestClient(app)


def test_statement_batch_delete_permissions_and_hard_delete(tmp_path, monkeypatch) -> None:
    db, users = _database()
    client = _configure(monkeypatch, tmp_path, db, users["finance"])
    form = {
        "company_name": "京奥电竞（北京）科技有限公司",
        "bank_name": "测试银行",
        "account_name": "基本户",
        "account_number": "6222000011117777",
    }
    try:
        pending = client.post(
            "/v1/finance/statements/upload",
            data=form,
            files={"file": ("待确认流水.xlsx", _xlsx(), "application/octet-stream")},
        ).json()
        pending_batch = db.get(BankStatementBatch, pending["id"])
        assert pending_batch is not None
        pending_source = Path(pending_batch.source_path)
        assert pending_source.is_file()

        for role in ("business", "manager", "founder"):
            app.dependency_overrides[current_user] = lambda role=role: users[role]
            forbidden = client.delete(
                f"/v1/finance/statements/{pending['id']}",
                params={"entity_id": pending["entity_id"]},
            )
            assert forbidden.status_code == 403
            assert forbidden.json()["detail"] == "待确认流水仅限财务删除"

        app.dependency_overrides[current_user] = lambda: users["finance"]
        wrong_entity = client.delete(
            f"/v1/finance/statements/{pending['id']}",
            params={"entity_id": "another-entity"},
        )
        assert wrong_entity.status_code == 404
        deleted_pending = client.delete(
            f"/v1/finance/statements/{pending['id']}",
            params={"entity_id": pending["entity_id"]},
        )
        assert deleted_pending.status_code == 200
        assert deleted_pending.json() == {
            "status": "deleted",
            "batch_id": pending["id"],
            "previous_status": "pending",
            "row_count": 2,
            "source_removed": True,
        }
        assert not pending_source.exists()
        assert db.get(BankStatementBatch, pending["id"]) is None
        assert db.scalar(select(func.count(BankTransaction.id))) == 0

        # A hard-deleted file can be uploaded again because its fingerprints
        # and batch uniqueness record have also been removed.
        confirmed = client.post(
            "/v1/finance/statements/upload",
            data=form,
            files={"file": ("待确认流水.xlsx", _xlsx(), "application/octet-stream")},
        ).json()
        assert confirmed["duplicate_file"] is False
        assert client.post(
            f"/v1/finance/statements/{confirmed['id']}/confirm",
            params={"entity_id": confirmed["entity_id"]},
        ).status_code == 200
        transaction = client.get(
            f"/v1/finance/statements/{confirmed['id']}/transactions",
            params={"entity_id": confirmed["entity_id"]},
        ).json()[0]
        correction = client.post(
            f"/v1/finance/transactions/{transaction['id']}/purpose-corrections",
            params={"entity_id": confirmed["entity_id"]},
            json={"proposed_purpose": "待创始人复核的用途"},
        )
        assert correction.status_code == 200

        for role in ("finance", "manager", "business"):
            app.dependency_overrides[current_user] = lambda role=role: users[role]
            forbidden = client.delete(
                f"/v1/finance/statements/{confirmed['id']}",
                params={"entity_id": confirmed["entity_id"]},
            )
            assert forbidden.status_code == 403
            assert forbidden.json()["detail"] == "已确认流水仅限创始人删除"

        confirmed_source = Path(db.get(BankStatementBatch, confirmed["id"]).source_path)
        app.dependency_overrides[current_user] = lambda: users["founder"]
        deleted_confirmed = client.delete(
            f"/v1/finance/statements/{confirmed['id']}",
            params={"entity_id": confirmed["entity_id"]},
        )
        assert deleted_confirmed.status_code == 200
        assert deleted_confirmed.json()["previous_status"] == "confirmed"
        assert not confirmed_source.exists()
        assert db.scalar(select(func.count(BankStatementBatch.id))) == 0
        assert db.scalar(select(func.count(BankTransaction.id))) == 0
        assert db.scalar(select(func.count(BankTransactionPurposeCorrection.id))) == 0

        delete_audits = db.scalars(
            select(AuditLog).where(AuditLog.action == "finance_statement_delete")
        ).all()
        assert len(delete_audits) == 2
        assert json.loads(delete_audits[-1].details_json)["previous_status"] == "confirmed"
        assert json.loads(delete_audits[-1].details_json)["purpose_correction_count"] == 1
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_finance_statement_upload_dedup_confirm_and_dashboard(tmp_path, monkeypatch) -> None:
    db, users = _database()
    client = _configure(monkeypatch, tmp_path, db, users["finance"])
    form = {
        "company_name": "京奥电竞（北京）科技有限公司",
        "bank_name": "测试银行",
        "account_name": "基本户",
        "account_number": "6222000012345678",
    }
    statement_bytes = _xlsx()
    try:
        first = client.post(
            "/v1/finance/statements/upload",
            data=form,
            files={"file": ("2026年8月流水.xlsx", statement_bytes, "application/octet-stream")},
        )
        assert first.status_code == 200
        assert first.json()["row_count"] == 2
        assert first.json()["duplicate_file"] is False
        assert db.scalar(select(func.count(BankTransaction.id))) == 2
        assert list((tmp_path / "knowledge" / "财务系统").rglob("*.xlsx"))

        duplicate = client.post(
            "/v1/finance/statements/upload",
            data=form,
            files={"file": ("重复.xlsx", statement_bytes, "application/octet-stream")},
        )
        assert duplicate.status_code == 200
        assert duplicate.json()["duplicate_file"] is True
        assert db.scalar(select(func.count(BankStatementBatch.id))) == 1

        transactions = client.get(
            f"/v1/finance/statements/{first.json()['id']}/transactions"
        )
        assert transactions.status_code == 200
        assert transactions.json()[0]["note"] == "项目回款"
        assert transactions.json()[1]["note"] == "场地费用"
        transaction_rows = db.scalars(
            select(BankTransaction).where(
                BankTransaction.batch_id == first.json()["id"]
            )
        ).all()
        for transaction_row in transaction_rows:
            transaction_row.note = None
            transaction_row.project_reference = None
        db.commit()
        repeated = client.post(
            "/v1/finance/statements/upload",
            data=form,
            files={
                "file": (
                    "statement.xlsx",
                    statement_bytes,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert repeated.status_code == 200
        assert repeated.json()["duplicate_file"] is True
        assert repeated.json()["annotation_updates"]["note_updates"] == 2
        refreshed = client.get(
            f"/v1/finance/statements/{first.json()['id']}/transactions"
        )
        assert refreshed.json()[0]["note"] == "项目回款"
        assert refreshed.json()[1]["note"] == "场地费用"
        update = client.patch(
            f"/v1/finance/transactions/{transactions.json()[0]['id']}",
            json={
                "transacted_at": "2026-08-01T09:30:00",
                "counterparty": "客户A（已核对）",
                "category": "项目回款",
                "pm_project_id": None,
            },
        )
        assert update.status_code == 200
        updated = client.get(
            f"/v1/finance/statements/{first.json()['id']}/transactions"
        ).json()[0]
        assert updated["counterparty"] == "客户A（已核对）"
        assert updated["summary"] == "项目回款"
        confirmed = client.post(
            f"/v1/finance/statements/{first.json()['id']}/confirm"
        )
        assert confirmed.status_code == 200
        confirmed_batch = client.get("/v1/finance/statements").json()[0]
        assert confirmed_batch["status"] == "confirmed"
        assert confirmed_batch["confirmer"] == "财务"
        assert confirmed_batch["confirmed_at"] is not None

        # Confirmation locks the bank-originated facts, while finance may
        # still add the real reimbursement purpose and Feishu project segment.
        annotation = client.patch(
            f"/v1/finance/transactions/{transactions.json()[0]['id']}",
            json={
                "category": "办公场地",
                "project_reference": " Cc2609 ",
                "pm_project_id": None,
            },
        )
        assert annotation.status_code == 200
        annotated = client.get(
            f"/v1/finance/statements/{first.json()['id']}/transactions"
        ).json()[0]
        assert annotated["note"] == "项目回款"
        assert annotated["project_reference"] == "Cc2609"
        immutable_edit = client.patch(
            f"/v1/finance/transactions/{transactions.json()[0]['id']}",
            json={"summary": "覆盖银行原始附言"},
        )
        assert immutable_edit.status_code == 409

        dashboard = client.get(
            "/v1/finance/dashboard?from_date=2026-08-01&to_date=2026-08-07"
        )
        assert dashboard.status_code == 200
        assert dashboard.json()["income"] == "10000.00"
        assert dashboard.json()["expense"] == "3000.00"
        assert dashboard.json()["bank_balance"] == "47000.00"
        assert dashboard.json()["net"] == "7000.00"
        assert dashboard.json()["weekly_top_income"] == [{
            "name": "客户A（已核对）",
            "amount": "10000.00",
            "transaction_count": 1,
        }]
        assert dashboard.json()["weekly_top_expense"] == [{
            "name": "供应商B",
            "amount": "3000.00",
            "transaction_count": 1,
        }]
        assert dashboard.json()["anomaly_count"] == 2
        assert {
            item["rule_codes"][0] for item in dashboard.json()["anomalies"]
        } == {
            "first_seen_income_counterparty",
            "first_seen_expense_counterparty",
        }
        assert dashboard.json()["oa_rule_ready"] is False
        assert dashboard.json()["previous_month"] == {
            "period_start": "2026-07-01",
            "period_end": "2026-07-31",
            "income": "0",
            "expense": "0",
            "net": "0",
        }
        assert dashboard.json()["current_year"] == {
            "period_start": "2026-01-01",
            "period_end": "2026-08-07",
            "income": "10000.00",
            "expense": "3000.00",
            "net": "7000.00",
        }
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_cost_center_in_statement_auto_links_confirmed_project_cash(
    tmp_path, monkeypatch
) -> None:
    db, users = _database()
    client = _configure(monkeypatch, tmp_path, db, users["finance"])
    try:
        entities = client.get("/v1/finance/entities").json()
        entity = entities[0]
        project = ManagedProject(
            project_no="CC26B05",
            name="三角洲国际战队培训",
            entity_id=entity["id"],
            company_name=entity["name"],
            client="测试甲方",
            business_category="电竞培训",
            manager_user_id=users["business"].id,
            members_json='["执行成员"]',
            planned_start=date(2026, 8, 1),
            planned_end=date(2026, 9, 30),
            objective="验证流水代码自动归集",
            status="active",
            created_by_user_id=users["business"].id,
        )
        db.add(project)
        db.commit()

        uploaded = client.post(
            "/v1/finance/statements/upload",
            data={
                "entity_id": entity["id"],
                "bank_name": "项目归集测试银行",
                "account_name": "基本户",
                "account_number": "6222000099990001",
            },
            files={
                "file": (
                    "项目代码流水.xlsx",
                    _project_code_xlsx(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert uploaded.status_code == 200
        transactions = db.scalars(
            select(BankTransaction).where(
                BankTransaction.batch_id == uploaded.json()["id"]
            )
        ).all()
        assert {item.project_reference for item in transactions} == {"CC26B05"}
        assert {item.pm_project_id for item in transactions} == {project.id}

        confirmed = client.post(
            f"/v1/finance/statements/{uploaded.json()['id']}/confirm",
            params={"entity_id": entity["id"]},
        )
        assert confirmed.status_code == 200
        project_payload = client.get(f"/v1/pm/projects/{project.id}").json()
        assert project_payload["bank_received"] == "12000.00"
        assert project_payload["bank_spent"] == "3500.00"
        assert project_payload["bank_transaction_count"] == 2
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_dashboard_combines_project_and_finance_receivables_payables(
    tmp_path, monkeypatch
) -> None:
    db, users = _database()
    client = _configure(monkeypatch, tmp_path, db, users["finance"])
    try:
        entity = client.get("/v1/finance/entities").json()[0]
        project = ManagedProject(
            project_no="CC26B06",
            name="后勤保障项目",
            entity_id=entity["id"],
            company_name=entity["name"],
            client="赛事主办方",
            business_category="赛事保障",
            manager_user_id=users["business"].id,
            members_json='["执行成员"]',
            planned_start=date(2026, 8, 1),
            planned_end=date(2026, 9, 30),
            objective="验证应收应付自动汇总",
            status="active",
            created_by_user_id=users["business"].id,
        )
        db.add(project)
        db.flush()
        db.add(ProjectCashflowPlan(
            project_id=project.id,
            direction="receivable",
            due_date=date(2026, 9, 10),
            amount=Decimal("1000"),
            actual_amount=Decimal("200"),
            counterparty="项目客户",
            created_by_user_id=users["business"].id,
        ))
        db.commit()

        manual = client.post(
            "/v1/finance/receivables-payables",
            json={
                "entity_id": entity["id"],
                "direction": "payable",
                "due_date": "2026-09-20",
                "amount": "600.00",
                "actual_amount": "100.00",
                "counterparty": "非项目供应商",
                "note": "总部服务费",
            },
        )
        assert manual.status_code == 200
        dashboard = client.get(
            "/v1/finance/dashboard", params={"entity_id": entity["id"]}
        ).json()["receivables_payables"]
        assert dashboard["receivable"]["total"] == "800.00"
        assert dashboard["receivable"]["count"] == 1
        assert dashboard["receivable"]["items"][0]["project_no"] == "CC26B06"
        assert dashboard["payable"]["total"] == "500.00"
        assert dashboard["payable"]["items"][0]["source"] == "finance"

        updated = client.patch(
            f"/v1/finance/receivables-payables/{manual.json()['id']}",
            params={"entity_id": entity["id"]},
            json={
                "direction": "payable",
                "due_date": "2026-09-22",
                "amount": "600.00",
                "actual_amount": "250.00",
                "counterparty": "非项目供应商",
                "note": "总部服务费已部分支付",
            },
        )
        assert updated.status_code == 200
        assert updated.json()["outstanding_amount"] == "350.00"
        assert db.scalar(select(func.count(FinanceReceivablePayable.id))) == 1

        app.dependency_overrides[current_user] = lambda: users["founder"]
        forbidden = client.post(
            "/v1/finance/receivables-payables",
            json={
                "entity_id": entity["id"],
                "direction": "receivable",
                "due_date": "2026-09-30",
                "amount": "1.00",
            },
        )
        assert forbidden.status_code == 403
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_finance_purpose_correction_requires_founder_review(tmp_path, monkeypatch) -> None:
    db, users = _database()
    client = _configure(monkeypatch, tmp_path, db, users["finance"])
    try:
        batch = client.post(
            "/v1/finance/statements/upload",
            data={
                "company_name": "京奥电竞（北京）科技有限公司",
                "bank_name": "测试银行",
                "account_name": "基本户",
                "account_number": "6222000099990000",
            },
            files={"file": ("用途审批.xlsx", _xlsx(), "application/octet-stream")},
        ).json()
        assert client.post(
            f"/v1/finance/statements/{batch['id']}/confirm",
            params={"entity_id": batch["entity_id"]},
        ).status_code == 200
        item = client.get(
            f"/v1/finance/statements/{batch['id']}/transactions",
            params={"entity_id": batch["entity_id"]},
        ).json()[0]
        assert item["note"] == "项目回款"

        app.dependency_overrides[current_user] = lambda: users["business"]
        assert client.post(
            f"/v1/finance/transactions/{item['id']}/purpose-corrections",
            params={"entity_id": batch["entity_id"]},
            json={"proposed_purpose": "业务账号越权申请"},
        ).status_code == 403
        app.dependency_overrides[current_user] = lambda: users["finance"]

        # The generic finance editor can never overwrite the bank memo or the
        # effective purpose. Purpose changes must use the review workflow.
        assert client.patch(
            f"/v1/finance/transactions/{item['id']}",
            params={"entity_id": batch["entity_id"]},
            json={"note": "绕过审批直接修改"},
        ).status_code == 409
        assert client.patch(
            f"/v1/finance/transactions/{item['id']}",
            params={"entity_id": batch["entity_id"]},
            json={"summary": "覆盖银行原始附言"},
        ).status_code == 409

        request = client.post(
            f"/v1/finance/transactions/{item['id']}/purpose-corrections",
            params={"entity_id": batch["entity_id"]},
            json={"proposed_purpose": "  Cc2609 白楼912办公室刷漆  "},
        )
        assert request.status_code == 200
        correction = request.json()
        assert correction["status"] == "pending"
        assert correction["previous_purpose"] == "项目回款"
        assert correction["proposed_purpose"] == "Cc2609 白楼912办公室刷漆"
        assert correction["requested_by"] == "财务"
        assert correction["transaction"]["summary"] == "项目回款"
        assert correction["transaction"]["current_purpose"] == "项目回款"
        other_entity = BusinessEntity(
            name="王牌猎豹（JAG三角洲）",
            created_by_user_id=users["founder"].id,
        )
        db.add(other_entity)
        db.commit()
        assert client.get(
            "/v1/finance/purpose-corrections",
            params={"entity_id": other_entity.id},
        ).json()["items"] == []
        assert client.post(
            f"/v1/finance/transactions/{item['id']}/purpose-corrections",
            params={"entity_id": batch["entity_id"]},
            json={"proposed_purpose": "另一用途"},
        ).status_code == 409

        registry = client.get(
            "/v1/finance/purpose-corrections",
            params={"entity_id": batch["entity_id"]},
        )
        assert registry.status_code == 200
        assert registry.json()["pending_count"] == 1
        assert registry.json()["items"][0]["id"] == correction["id"]
        annotated = client.get(
            f"/v1/finance/statements/{batch['id']}/transactions",
            params={"entity_id": batch["entity_id"]},
        ).json()[0]
        assert annotated["note"] == "项目回款"
        assert annotated["purpose_correction"]["status"] == "pending"

        app.dependency_overrides[current_user] = lambda: users["manager"]
        assert client.patch(
            f"/v1/finance/purpose-corrections/{correction['id']}/review",
            params={"entity_id": batch["entity_id"]},
            json={"decision": "approve"},
        ).status_code == 403
        app.dependency_overrides[current_user] = lambda: users["finance"]
        assert client.patch(
            f"/v1/finance/purpose-corrections/{correction['id']}/review",
            params={"entity_id": batch["entity_id"]},
            json={"decision": "approve"},
        ).status_code == 403

        app.dependency_overrides[current_user] = lambda: users["founder"]
        assert client.patch(
            f"/v1/finance/purpose-corrections/{correction['id']}/review",
            params={"entity_id": other_entity.id},
            json={"decision": "approve"},
        ).status_code == 404
        approved = client.patch(
            f"/v1/finance/purpose-corrections/{correction['id']}/review",
            params={"entity_id": batch["entity_id"]},
            json={"decision": "approve", "review_comment": "用途凭证已核实"},
        )
        assert approved.status_code == 200
        assert approved.json()["status"] == "approved"
        assert approved.json()["reviewed_by"] == "叶靖波"
        assert approved.json()["transaction"]["current_purpose"] == "Cc2609 白楼912办公室刷漆"
        assert client.patch(
            f"/v1/finance/purpose-corrections/{correction['id']}/review",
            params={"entity_id": batch["entity_id"]},
            json={"decision": "approve"},
        ).status_code == 409

        # A later proposal is allowed after review. Rejecting it requires a
        # reason and leaves the last approved purpose unchanged.
        app.dependency_overrides[current_user] = lambda: users["finance"]
        second = client.post(
            f"/v1/finance/transactions/{item['id']}/purpose-corrections",
            params={"entity_id": batch["entity_id"]},
            json={"proposed_purpose": "未经证明的其他用途"},
        ).json()
        app.dependency_overrides[current_user] = lambda: users["founder"]
        assert client.patch(
            f"/v1/finance/purpose-corrections/{second['id']}/review",
            params={"entity_id": batch["entity_id"]},
            json={"decision": "reject"},
        ).status_code == 422
        rejected = client.patch(
            f"/v1/finance/purpose-corrections/{second['id']}/review",
            params={"entity_id": batch["entity_id"]},
            json={"decision": "reject", "review_comment": "缺少对应业务凭证"},
        )
        assert rejected.status_code == 200
        assert rejected.json()["status"] == "rejected"
        transaction = db.get(BankTransaction, item["id"])
        assert transaction is not None
        assert transaction.note == "Cc2609 白楼912办公室刷漆"
        assert db.scalar(select(func.count(BankTransactionPurposeCorrection.id))) == 2

        audit_payloads = [
            row.details_json
            for row in db.scalars(
                select(AuditLog).where(
                    AuditLog.action.in_({
                        "finance_purpose_correction_request",
                        "finance_purpose_correction_review",
                    })
                )
            ).all()
        ]
        assert audit_payloads
        assert all("白楼912办公室刷漆" not in payload for payload in audit_payloads)
        assert all("未经证明的其他用途" not in payload for payload in audit_payloads)
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_finance_statement_upload_reports_read_only_storage(tmp_path, monkeypatch) -> None:
    db, users = _database()
    client = _configure(monkeypatch, tmp_path, db, users["finance"])
    blocked_root = tmp_path / "knowledge-is-a-file"
    blocked_root.write_text("not a directory", encoding="utf-8")
    finance_system.settings.knowledge_root = blocked_root
    try:
        response = client.post(
            "/v1/finance/statements/upload",
            data={
                "company_name": "京奥电竞（北京）科技有限公司",
                "bank_name": "测试银行",
                "account_name": "基本户",
                "account_number": "6222000012345678",
            },
            files={"file": ("2026年8月流水.xlsx", _xlsx(), "application/octet-stream")},
        )
        assert response.status_code == 503
        assert response.json()["detail"] == "财务原件目录当前不可写，请联系系统管理员检查NAS挂载权限"
        assert db.scalar(select(func.count(BankStatementBatch.id))) == 0
        assert db.scalar(select(func.count(FinancialAccount.id))) == 0
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_failed_statement_is_quarantined_outside_knowledge_and_audited(
    tmp_path, monkeypatch
) -> None:
    db, users = _database()
    client = _configure(monkeypatch, tmp_path, db, users["finance"])
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "查询条件"
    sheet.append(["账户名称", "测试公司"])
    sheet.append(["借方累计发生额", 4000, "贷方累计发生额", 1000000])
    output = BytesIO()
    workbook.save(output)
    payload = output.getvalue()
    try:
        response = client.post(
            "/v1/finance/statements/upload",
            data={
                "company_name": "京奥电竞（北京）科技有限公司",
                "bank_name": "测试银行",
                "account_name": "基本户",
                "account_number": "6222000012345678",
            },
            files={"file": ("只有汇总页.xlsx", payload, "application/octet-stream")},
        )
        assert response.status_code == 422
        diagnostic_id = hashlib.sha256(payload).hexdigest().upper()[:12]
        assert diagnostic_id in response.json()["detail"]
        assert "无需反复上传" in response.json()["detail"]
        quarantined = list(
            (tmp_path / "agent-data" / "finance-import-failures").glob("*.xlsx")
        )
        assert len(quarantined) == 1
        assert quarantined[0].read_bytes() == payload
        assert not list((tmp_path / "knowledge").rglob("只有汇总页.xlsx"))
        assert db.scalar(select(func.count(BankStatementBatch.id))) == 0
        assert db.scalar(select(func.count(FinancialAccount.id))) == 0
        audit = db.scalar(
            select(AuditLog).where(
                AuditLog.action == "finance_statement_parse_failed"
            )
        )
        assert audit is not None
        details = json.loads(audit.details_json)
        assert details["diagnostic_id"] == diagnostic_id
        assert details["structure"] == [{
            "sheet": "查询条件",
            "row_count": 2,
            "max_columns": 4,
            "recognised_fields": [],
        }]
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_finance_health_excludes_exact_internal_transfer_but_keeps_real_balance(
    tmp_path, monkeypatch
) -> None:
    db, users = _database()
    client = _configure(monkeypatch, tmp_path, db, users["finance"])
    try:
        entities = client.get("/v1/finance/entities").json()
        jingao = next(item for item in entities if item["key"] == "jingao")
        ace = next(item for item in entities if item["key"] == "ace-leopard")
        source_account = FinancialAccount(
            entity_id=jingao["id"], bank_name="测试银行", account_name="基本户",
            account_number_masked="6222****1111", account_number_hash="SRC-HASH",
        )
        target_account = FinancialAccount(
            entity_id=ace["id"], bank_name="测试银行", account_name="基本户",
            account_number_masked="6222****2222", account_number_hash="TARGET-HASH",
        )
        secondary_source_account = FinancialAccount(
            entity_id=jingao["id"], bank_name="测试银行", account_name="一般户",
            account_number_masked="6222****3333", account_number_hash="SRC-SECONDARY",
        )
        db.add_all([source_account, target_account, secondary_source_account])
        db.flush()
        batch = BankStatementBatch(
            entity_id=jingao["id"], account_id=source_account.id,
            original_filename="health.xlsx", source_path="/health.xlsx",
            file_hash="HEALTH-FILE", period_start=date(2026, 8, 1),
            period_end=date(2026, 8, 7), status="confirmed", row_count=4,
            uploaded_by_user_id=users["finance"].id,
            confirmed_by_user_id=users["finance"].id,
            confirmed_at=datetime.now(timezone.utc),
        )
        db.add(batch)
        db.flush()
        secondary_batch = BankStatementBatch(
            entity_id=jingao["id"], account_id=secondary_source_account.id,
            original_filename="secondary.xlsx", source_path="/secondary.xlsx",
            file_hash="HEALTH-SECONDARY", period_start=date(2026, 8, 4),
            period_end=date(2026, 8, 4), status="confirmed", row_count=1,
            uploaded_by_user_id=users["finance"].id,
            confirmed_by_user_id=users["finance"].id,
            confirmed_at=datetime.now(timezone.utc),
        )
        db.add(secondary_batch)
        db.flush()

        def transaction(
            day: int, *, income: str = "0", expense: str = "0",
            balance: str, counterparty: str, category: str = "待确认",
            counterparty_hash: str | None = None,
        ) -> BankTransaction:
            return BankTransaction(
                batch_id=batch.id, entity_id=jingao["id"], account_id=source_account.id,
                transacted_at=datetime(2026, 8, day, 2, tzinfo=timezone.utc),
                income=Decimal(income), expense=Decimal(expense),
                balance=Decimal(balance), counterparty=counterparty,
                counterparty_account_hash=counterparty_hash,
                summary="测试", category=category,
                fingerprint=f"health-{day}-{counterparty}", status="confirmed",
            )

        external_income = transaction(
            1, income="10000", balance="50000", counterparty="客户A",
            category="项目回款",
        )
        internal_expense = transaction(
            2, expense="5000", balance="45000", counterparty="未知显示名",
            category="内部往来", counterparty_hash="TARGET-HASH",
        )
        external_expense = transaction(
            3, expense="3000", balance="42000", counterparty="供应商B"
        )
        fuzzy_candidate = transaction(
            4, expense="1000", balance="41000", counterparty="王牌猎豹集团项目部",
            category="其他",
        )
        late_account_balance = BankTransaction(
            batch_id=secondary_batch.id, entity_id=jingao["id"],
            account_id=secondary_source_account.id,
            transacted_at=datetime(2026, 8, 4, 3, tzinfo=timezone.utc),
            income=Decimal("0"), expense=Decimal("0"), balance=Decimal("100"),
            counterparty="余额初始化", summary="余额初始化", category="余额初始化",
            fingerprint="health-secondary", status="confirmed",
        )
        db.add_all([
            external_income, internal_expense, external_expense, fuzzy_candidate,
            late_account_balance,
        ])
        db.commit()

        excluded = client.get(
            "/v1/finance/dashboard",
            params={
                "entity_id": jingao["id"], "from_date": "2026-08-01",
                "to_date": "2026-08-07",
            },
        ).json()
        assert Decimal(excluded["income"]) == Decimal("10000.00")
        assert Decimal(excluded["expense"]) == Decimal("4000.00")
        assert Decimal(excluded["bank_balance"]) == Decimal("41100.00")
        # Days 1-3 only cover one of the two accounts and cannot create a
        # misleading partial-account minimum.
        assert Decimal(excluded["health"]["minimum_balance"]["amount"]) == Decimal("41100.00")
        assert excluded["health"]["minimum_balance"]["date"] == "2026-08-04"
        assert excluded["health"]["minimum_balance"]["coverage_complete"] is False
        assert excluded["health"]["minimum_balance"]["coverage_start"] == "2026-08-04"
        assert excluded["health"]["minimum_balance"]["covered_days"] == 4
        assert excluded["health"]["minimum_balance"]["total_days"] == 7
        assert Decimal(excluded["health"]["maximum_daily_net_outflow"]["amount"]) == Decimal("3000.00")
        assert excluded["health"]["maximum_daily_net_outflow"]["date"] == "2026-08-03"
        assert excluded["health"]["internal_transfers"]["count"] == 1
        assert Decimal(excluded["health"]["internal_transfers"]["gross_amount"]) == Decimal("5000.00")
        assert excluded["health"]["unclassified"]["count"] == 1

        included = client.get(
            "/v1/finance/dashboard",
            params={
                "entity_id": jingao["id"], "from_date": "2026-08-01",
                "to_date": "2026-08-07", "include_internal_transfers": "true",
            },
        ).json()
        assert Decimal(included["expense"]) == Decimal("9000.00")
        assert Decimal(included["bank_balance"]) == Decimal("41100.00")

        registry = client.get(
            "/v1/finance/internal-transfers",
            params={"entity_id": jingao["id"]},
        ).json()
        by_id = {item["id"]: item for item in registry["items"]}
        assert by_id[internal_expense.id]["internal_transfer_status"] == "auto_confirmed"
        assert by_id[internal_expense.id]["internal_transfer_source"] == "exact_cross_entity_account"
        assert by_id[fuzzy_candidate.id]["is_candidate"] is True
        candidate_confirmed = client.patch(
            f"/v1/finance/internal-transfers/{fuzzy_candidate.id}",
            json={"decision": "confirm"},
        )
        assert candidate_confirmed.status_code == 200
        assert candidate_confirmed.json()["target_entity_id"] == ace["id"]
        assert candidate_confirmed.json()["internal_transfer_status"] == "manual_confirmed"
        # Review metadata remains editable after the statement is confirmed.
        rejected = client.patch(
            f"/v1/finance/internal-transfers/{internal_expense.id}",
            json={"decision": "reject", "note": "财务核对为外部款项"},
        )
        assert rejected.status_code == 200
        assert rejected.json()["internal_transfer_status"] == "not_internal"
        after_reject = client.get(
            "/v1/finance/dashboard",
            params={
                "entity_id": jingao["id"], "from_date": "2026-08-01",
                "to_date": "2026-08-07",
            },
        ).json()
        assert Decimal(after_reject["expense"]) == Decimal("8000.00")
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_finance_dashboard_explainable_anomaly_rules(tmp_path, monkeypatch) -> None:
    db, users = _database()
    client = _configure(monkeypatch, tmp_path, db, users["finance"])
    try:
        uploaded = client.post(
            "/v1/finance/statements/upload",
            data={
                "company_name": "京奥电竞（北京）科技有限公司",
                "bank_name": "测试银行",
                "account_name": "基本户",
                "account_number": "6222000099999999",
            },
            files={"file": ("异常规则.xlsx", _analysis_warning_xlsx(), "application/octet-stream")},
        )
        assert uploaded.status_code == 200
        dashboard = client.get(
            "/v1/finance/dashboard?from_date=2026-08-01&to_date=2026-08-07"
        ).json()
        assert dashboard["weekly_top_income"] == []
        assert dashboard["weekly_top_expense"] == []
        # Uploaded batches are pending until finance confirms them.  Pending
        # rows must not create either totals or alerts.
        assert dashboard["anomaly_count"] == 0
        assert dashboard["anomalies"] == []
        assert dashboard["oa_rule_ready"] is False
        assert "无审批支出" in dashboard["planned_rules"]
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_finance_dashboard_counterparty_behaviour_rules(tmp_path, monkeypatch) -> None:
    db, users = _database()
    client = _configure(monkeypatch, tmp_path, db, users["finance"])
    try:
        entities = client.get("/v1/finance/entities").json()
        jingao = next(item for item in entities if item["key"] == "jingao")
        account = FinancialAccount(
            entity_id=jingao["id"],
            bank_name="规则测试银行",
            account_name="规则测试账户",
            account_number_masked="****0001",
            account_number_hash="alert-rule-account",
        )
        db.add(account)
        db.flush()
        confirmed_batch = BankStatementBatch(
            entity_id=jingao["id"],
            account_id=account.id,
            original_filename="confirmed.xlsx",
            source_path="/confirmed.xlsx",
            file_hash="alert-confirmed-file",
            period_start=date(2026, 7, 1),
            period_end=date(2026, 8, 7),
            status="confirmed",
            row_count=12,
            uploaded_by_user_id=users["finance"].id,
            confirmed_by_user_id=users["finance"].id,
        )
        pending_batch = BankStatementBatch(
            entity_id=jingao["id"],
            account_id=account.id,
            original_filename="pending.xlsx",
            source_path="/pending.xlsx",
            file_hash="alert-pending-file",
            period_start=date(2026, 8, 1),
            period_end=date(2026, 8, 7),
            status="pending",
            row_count=1,
            uploaded_by_user_id=users["finance"].id,
        )
        db.add_all([confirmed_batch, pending_batch])
        db.flush()

        serial = 0

        def add_transaction(
            *,
            day: datetime,
            counterparty: str,
            income: str = "0",
            expense: str = "0",
            status: str = "confirmed",
        ) -> str:
            nonlocal serial
            serial += 1
            item = BankTransaction(
                batch_id=(
                    confirmed_batch.id if status == "confirmed" else pending_batch.id
                ),
                entity_id=jingao["id"],
                account_id=account.id,
                transacted_at=day,
                income=Decimal(income),
                expense=Decimal(expense),
                balance=None,
                counterparty=counterparty,
                summary="规则测试摘要",
                bank_serial=f"ALERT-{serial:02d}",
                category="规则测试",
                fingerprint=f"alert-fingerprint-{serial:02d}",
                status=status,
            )
            db.add(item)
            db.flush()
            return item.id

        # Known parties establish history before the selected period.
        add_transaction(
            day=datetime(2026, 7, 1, 4, tzinfo=timezone.utc),
            counterparty="已知大额单位",
            expense="1000",
        )
        add_transaction(
            day=datetime(2026, 7, 2, 4, tzinfo=timezone.utc),
            counterparty="已知收款单位",
            income="5000",
        )
        for day, amount in ((3, "10000"), (4, "12000"), (5, "8000")):
            add_transaction(
                day=datetime(2026, 7, day, 4, tzinfo=timezone.utc),
                counterparty="历史稳定供应商",
                expense=amount,
            )

        # A large amount alone is not an anomaly, and internal transfers are
        # excluded from first-party alerts even when they are large.
        known_large_id = add_transaction(
            day=datetime(2026, 8, 1, 2, tzinfo=timezone.utc),
            counterparty="已知大额单位",
            expense="150000",
        )
        internal_id = add_transaction(
            day=datetime(2026, 8, 1, 3, tzinfo=timezone.utc),
            counterparty="京奥电竞（北京）科技有限公司",
            expense="200000",
        )
        first_new_id = add_transaction(
            day=datetime(2026, 8, 1, 4, tzinfo=timezone.utc),
            counterparty="首次供应商",
            expense="5000",
        )
        second_new_id = add_transaction(
            day=datetime(2026, 8, 2, 4, tzinfo=timezone.utc),
            counterparty="首次供应商",
            expense="6000",
        )
        first_income_id = add_transaction(
            day=datetime(2026, 8, 3, 4, tzinfo=timezone.utc),
            counterparty="首次客户",
            income="7000",
        )
        spike_id = add_transaction(
            day=datetime(2026, 8, 4, 4, tzinfo=timezone.utc),
            counterparty="历史稳定供应商",
            expense="50000",
        )
        pending_id = add_transaction(
            day=datetime(2026, 8, 5, 4, tzinfo=timezone.utc),
            counterparty="待确认供应商",
            expense="9999",
            status="pending",
        )
        db.commit()

        dashboard = client.get(
            "/v1/finance/dashboard",
            params={
                "entity_id": jingao["id"],
                "from_date": "2026-08-01",
                "to_date": "2026-08-07",
            },
        ).json()
        alerts_by_id = {item["id"]: item for item in dashboard["anomalies"]}

        assert dashboard["anomaly_count"] == 3
        assert set(alerts_by_id) == {first_new_id, first_income_id, spike_id}
        assert alerts_by_id[first_new_id]["alert_type"] == "new_counterparty"
        assert alerts_by_id[first_new_id]["rule_codes"] == [
            "first_seen_expense_counterparty"
        ]
        assert alerts_by_id[first_income_id]["rule_codes"] == [
            "first_seen_income_counterparty"
        ]
        assert alerts_by_id[spike_id]["alert_type"] == "behaviour_change"
        assert alerts_by_id[spike_id]["rule_codes"] == ["expense_amount_spike"]
        assert "此前 3 笔中位数为 10,000.00 元" in alerts_by_id[spike_id]["reasons"][0]
        assert known_large_id not in alerts_by_id
        assert internal_id not in alerts_by_id
        assert second_new_id not in alerts_by_id
        assert pending_id not in alerts_by_id
        assert all(item["status"] == "confirmed" for item in dashboard["anomalies"])
        assert all(
            "10 万元" not in reason
            for item in dashboard["anomalies"]
            for reason in item["reasons"]
        )
        assert dashboard["alert_rule_status"]["oa_rule_ready"] is False
        assert {
            rule["code"]
            for rule in dashboard["alert_rule_status"]["planned_rules"]
        } >= {
            "oa_expense_without_approval",
            "oa_payee_mismatch",
            "oa_amount_exceeded",
            "oa_duplicate_payment",
        }
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_finance_dashboard_previous_month_panels_are_independent_and_scoped(
    tmp_path, monkeypatch
) -> None:
    db, users = _database()
    client = _configure(monkeypatch, tmp_path, db, users["finance"])
    try:
        entities = client.get("/v1/finance/entities").json()
        jingao = next(item for item in entities if item["key"] == "jingao")
        ace = next(item for item in entities if item["key"] == "ace-leopard")
        jingao_account = FinancialAccount(
            entity_id=jingao["id"],
            bank_name="月度测试银行",
            account_name="月度测试账户",
            account_number_masked="****7001",
            account_number_hash="monthly-panel-jingao-account",
        )
        ace_account = FinancialAccount(
            entity_id=ace["id"],
            bank_name="隔离测试银行",
            account_name="隔离测试账户",
            account_number_masked="****7002",
            account_number_hash="monthly-panel-ace-account",
        )
        db.add_all([jingao_account, ace_account])
        db.flush()
        confirmed_batch = BankStatementBatch(
            entity_id=jingao["id"],
            account_id=jingao_account.id,
            original_filename="月度已确认.xlsx",
            source_path="/monthly-confirmed.xlsx",
            file_hash="monthly-confirmed-file",
            period_start=date(2026, 6, 1),
            period_end=date(2026, 8, 7),
            status="confirmed",
            row_count=30,
            uploaded_by_user_id=users["finance"].id,
            confirmed_by_user_id=users["finance"].id,
        )
        pending_batch = BankStatementBatch(
            entity_id=jingao["id"],
            account_id=jingao_account.id,
            original_filename="月度待确认.xlsx",
            source_path="/monthly-pending.xlsx",
            file_hash="monthly-pending-file",
            period_start=date(2026, 7, 1),
            period_end=date(2026, 7, 31),
            status="pending",
            row_count=1,
            uploaded_by_user_id=users["finance"].id,
        )
        other_entity_batch = BankStatementBatch(
            entity_id=ace["id"],
            account_id=ace_account.id,
            original_filename="其他公司七月.xlsx",
            source_path="/monthly-other-entity.xlsx",
            file_hash="monthly-other-entity-file",
            period_start=date(2026, 7, 1),
            period_end=date(2026, 7, 31),
            status="confirmed",
            row_count=1,
            uploaded_by_user_id=users["finance"].id,
            confirmed_by_user_id=users["finance"].id,
        )
        db.add_all([confirmed_batch, pending_batch, other_entity_batch])
        db.flush()

        serial = 0

        def add_transaction(
            *,
            day: datetime,
            counterparty: str | None,
            income: str = "0",
            expense: str = "0",
            summary: str | None = "月度规则测试",
            status: str = "confirmed",
            entity: dict = jingao,
        ) -> str:
            nonlocal serial
            serial += 1
            is_jingao = entity["id"] == jingao["id"]
            if status == "pending":
                batch = pending_batch
                account = jingao_account
            elif is_jingao:
                batch = confirmed_batch
                account = jingao_account
            else:
                batch = other_entity_batch
                account = ace_account
            item = BankTransaction(
                batch_id=batch.id,
                entity_id=entity["id"],
                account_id=account.id,
                transacted_at=day,
                income=Decimal(income),
                expense=Decimal(expense),
                balance=None,
                counterparty=counterparty,
                summary=summary,
                bank_serial=f"MONTH-{serial:02d}",
                category="月度规则测试",
                fingerprint=f"monthly-panel-fingerprint-{serial:02d}",
                status=status,
            )
            db.add(item)
            db.flush()
            return item.id

        # A three-row baseline strictly before July is available for the
        # behaviour-change rule.  July rows are then appended chronologically.
        for day, amount in ((2, "1000"), (3, "1100"), (4, "900")):
            add_transaction(
                day=datetime(2026, 6, day, 4, tzinfo=timezone.utc),
                counterparty="历史稳定供应商",
                expense=amount,
            )

        first_monthly_customer = add_transaction(
            day=datetime(2026, 7, 2, 4, tzinfo=timezone.utc),
            counterparty="月度客户A",
            income="20000",
        )
        second_monthly_customer = add_transaction(
            day=datetime(2026, 7, 3, 4, tzinfo=timezone.utc),
            counterparty="月度客户A",
            income="30000",
        )
        first_monthly_supplier = add_transaction(
            day=datetime(2026, 7, 4, 4, tzinfo=timezone.utc),
            counterparty="月度供应商A",
            expense="5000",
        )
        second_monthly_supplier = add_transaction(
            day=datetime(2026, 7, 5, 4, tzinfo=timezone.utc),
            counterparty="月度供应商A",
            expense="7000",
        )
        monthly_spike = add_transaction(
            day=datetime(2026, 7, 6, 4, tzinfo=timezone.utc),
            counterparty="历史稳定供应商",
            expense="50000",
        )
        missing_party = add_transaction(
            day=datetime(2026, 7, 7, 4, tzinfo=timezone.utc),
            counterparty=None,
            expense="123",
        )

        # Twelve distinct receipt sources make the Top 10 cap observable.
        for number in range(1, 13):
            add_transaction(
                day=datetime(2026, 7, 8 + number, 4, tzinfo=timezone.utc),
                counterparty=f"月度排行客户{number:02d}",
                income=str(number * 1000),
            )

        weekly_income_id = add_transaction(
            day=datetime(2026, 8, 2, 4, tzinfo=timezone.utc),
            counterparty="仅周度客户",
            income="99999",
        )
        weekly_expense_id = add_transaction(
            day=datetime(2026, 8, 3, 4, tzinfo=timezone.utc),
            counterparty="仅周度供应商",
            expense="88888",
        )
        pending_id = add_transaction(
            day=datetime(2026, 7, 20, 5, tzinfo=timezone.utc),
            counterparty="待确认百万客户",
            income="1000000",
            status="pending",
        )
        other_entity_id = add_transaction(
            day=datetime(2026, 7, 21, 5, tzinfo=timezone.utc),
            counterparty="其他公司百万客户",
            income="2000000",
            entity=ace,
        )
        db.commit()

        dashboard = client.get(
            "/v1/finance/dashboard",
            params={
                "entity_id": jingao["id"],
                "from_date": "2026-08-01",
                "to_date": "2026-08-07",
            },
        )
        assert dashboard.status_code == 200
        payload = dashboard.json()

        assert payload["weekly_top_income"] == [
            {"name": "仅周度客户", "amount": "99999.00", "transaction_count": 1}
        ]
        assert payload["weekly_top_expense"] == [
            {"name": "仅周度供应商", "amount": "88888.00", "transaction_count": 1}
        ]
        assert payload["previous_month"] == {
            "period_start": "2026-07-01",
            "period_end": "2026-07-31",
            "income": "128000.00",
            "expense": "62123.00",
            "net": "65877.00",
        }
        assert len(payload["monthly_top_income"]) == 10
        assert payload["monthly_top_income"][0] == {
            "name": "月度客户A",
            "amount": "50000.00",
            "transaction_count": 2,
        }
        assert {item["name"] for item in payload["monthly_top_income"]}.isdisjoint(
            {"仅周度客户", "待确认百万客户", "其他公司百万客户"}
        )
        assert payload["monthly_top_expense"][:2] == [
            {
                "name": "历史稳定供应商",
                "amount": "50000.00",
                "transaction_count": 1,
            },
            {
                "name": "月度供应商A",
                "amount": "12000.00",
                "transaction_count": 2,
            },
        ]

        monthly_alerts = {
            item["id"]: item for item in payload["monthly_anomalies"]
        }
        assert first_monthly_customer in monthly_alerts
        assert second_monthly_customer not in monthly_alerts
        assert first_monthly_supplier in monthly_alerts
        assert second_monthly_supplier not in monthly_alerts
        assert monthly_alerts[monthly_spike]["rule_codes"] == [
            "expense_amount_spike"
        ]
        assert monthly_alerts[missing_party]["rule_codes"] == [
            "missing_counterparty"
        ]
        assert weekly_income_id not in monthly_alerts
        assert weekly_expense_id not in monthly_alerts
        assert pending_id not in monthly_alerts
        assert other_entity_id not in monthly_alerts
        assert payload["monthly_anomaly_count"] == len(payload["monthly_anomalies"])
        assert all(item["status"] == "confirmed" for item in payload["monthly_anomalies"])
        assert all(
            "2026-07-01" <= item["transaction_date"] <= "2026-07-31"
            for item in payload["monthly_anomalies"]
        )
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_finance_permissions_and_manual_cash(tmp_path, monkeypatch) -> None:
    db, users = _database()
    client = _configure(monkeypatch, tmp_path, db, users["business"])
    payload = {
        "company_name": "京奥电竞（北京）科技有限公司",
        "cash_account": "备用金",
        "transaction_date": "2026-08-01",
        "direction": "opening",
        "amount": "5000.00",
        "category": "期初现金",
        "note": "试运行期初余额",
    }
    try:
        assert client.get("/v1/finance/dashboard").status_code == 403
        assert client.post("/v1/finance/cash", json=payload).status_code == 403

        app.dependency_overrides[current_user] = lambda: users["manager"]
        manager_dashboard = client.get("/v1/finance/dashboard")
        assert manager_dashboard.status_code == 200
        expected_end = date.today() - timedelta(days=7)
        assert manager_dashboard.json()["period_end"] == expected_end.isoformat()
        assert manager_dashboard.json()["period_start"] == (expected_end - timedelta(days=6)).isoformat()
        assert client.post("/v1/finance/cash", json=payload).status_code == 403

        app.dependency_overrides[current_user] = lambda: users["founder"]
        created = client.post("/v1/finance/cash", json=payload)
        assert created.status_code == 200
        expense_payload = {
            **payload,
            "transaction_date": "2026-08-02",
            "direction": "expense",
            "amount": "1333.00",
            "category": "账外支出",
            "note": "不得混入账内表盘",
        }
        assert client.post("/v1/finance/cash", json=expense_payload).status_code == 200
        assert db.scalar(select(func.count(CashEntry.id))) == 2
        dashboard = client.get(
            "/v1/finance/dashboard?from_date=2026-08-01&to_date=2026-08-07"
        )
        assert dashboard.json()["cash_balance"] == "3667.00"
        assert Decimal(dashboard.json()["income"]) == Decimal("0")
        assert Decimal(dashboard.json()["expense"]) == Decimal("0")
        assert Decimal(dashboard.json()["net"]) == Decimal("0")
        assert dashboard.json()["weekly"] == []
        assert Decimal(dashboard.json()["previous_month"]["income"]) == Decimal("0")
        assert Decimal(dashboard.json()["previous_month"]["expense"]) == Decimal("0")
        assert Decimal(dashboard.json()["current_year"]["income"]) == Decimal("0")
        assert Decimal(dashboard.json()["current_year"]["expense"]) == Decimal("0")
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_finance_entities_are_fixed_and_company_data_is_isolated(tmp_path, monkeypatch) -> None:
    db, users = _database()
    client = _configure(monkeypatch, tmp_path, db, users["finance"])
    try:
        entity_response = client.get("/v1/finance/entities")
        assert entity_response.status_code == 200
        entities = entity_response.json()
        assert [item["key"] for item in entities] == [
            "jingao",
            "ace-leopard",
            "power-leopard",
            "xingyao",
        ]
        assert [item["display_name"] for item in entities] == [
            "京奥电竞",
            "王牌猎豹",
            "劲腾豹跃",
            "星曜电竞",
        ]
        assert [item["business_name"] for item in entities] == [
            "总公司",
            "JAG三角洲",
            "JAG王者",
            "电竞教培",
        ]
        assert [item["show_cash"] for item in entities] == [True, False, False, False]
        assert client.get("/v1/finance/entities").json() == entities
        registry_audit = db.scalars(select(AuditLog).where(
            AuditLog.action == "business_entity_registry_normalized_v1"
        )).all()
        assert len(registry_audit) == 1
        assert json.loads(registry_audit[0].details_json)["active_entity_count"] == 4
        entity_by_key = {item["key"]: item for item in entities}
        jingao = entity_by_key["jingao"]
        ace_leopard = entity_by_key["ace-leopard"]
        power_leopard = entity_by_key["power-leopard"]

        jingao_upload = client.post(
            "/v1/finance/statements/upload",
            data={
                "entity_id": jingao["id"],
                "bank_name": "北京银行成寿寺支行",
                "account_name": "京奥基本户",
                "account_number": "20000100000000000001",
            },
            files={
                "file": (
                    "京奥流水.xlsx",
                    _company_xlsx(
                        serial_prefix="JINGAO",
                        income=Decimal("12000.00"),
                        expense=Decimal("2000.00"),
                        opening_balance=Decimal("50000.00"),
                        income_counterparty="京奥客户",
                        expense_counterparty="京奥供应商",
                    ),
                    "application/octet-stream",
                )
            },
        )
        ace_upload = client.post(
            "/v1/finance/statements/upload",
            data={
                "entity_id": ace_leopard["id"],
                "bank_name": "测试银行三角洲支行",
                "account_name": "三角洲基本户",
                "account_number": "20000100000000000002",
            },
            files={
                "file": (
                    "王牌猎豹流水.xlsx",
                    _company_xlsx(
                        serial_prefix="ACE",
                        income=Decimal("33000.00"),
                        expense=Decimal("7000.00"),
                        opening_balance=Decimal("100000.00"),
                        income_counterparty="三角洲客户",
                        expense_counterparty="三角洲供应商",
                    ),
                    "application/octet-stream",
                )
            },
        )
        assert jingao_upload.status_code == 200
        assert ace_upload.status_code == 200
        jingao_batch = jingao_upload.json()
        ace_batch = ace_upload.json()
        assert jingao_batch["entity_id"] == jingao["id"]
        assert ace_batch["entity_id"] == ace_leopard["id"]

        jingao_batches = client.get(
            "/v1/finance/statements", params={"entity_id": jingao["id"]}
        )
        ace_batches = client.get(
            "/v1/finance/statements", params={"entity_id": ace_leopard["id"]}
        )
        power_batches = client.get(
            "/v1/finance/statements", params={"entity_id": power_leopard["id"]}
        )
        assert [item["id"] for item in jingao_batches.json()] == [jingao_batch["id"]]
        assert [item["id"] for item in ace_batches.json()] == [ace_batch["id"]]
        assert power_batches.json() == []

        jingao_transactions = client.get(
            f"/v1/finance/statements/{jingao_batch['id']}/transactions",
            params={"entity_id": jingao["id"]},
        )
        ace_transactions = client.get(
            f"/v1/finance/statements/{ace_batch['id']}/transactions",
            params={"entity_id": ace_leopard["id"]},
        )
        assert jingao_transactions.status_code == 200
        assert ace_transactions.status_code == 200
        assert {item["counterparty"] for item in jingao_transactions.json()} == {
            "京奥客户",
            "京奥供应商",
        }
        assert {item["counterparty"] for item in ace_transactions.json()} == {
            "三角洲客户",
            "三角洲供应商",
        }
        assert {
            item["id"] for item in jingao_transactions.json()
        }.isdisjoint({item["id"] for item in ace_transactions.json()})

        # A valid but different company identifier must conceal the batch and
        # transaction instead of revealing that another company owns it.
        assert client.get(
            f"/v1/finance/statements/{jingao_batch['id']}/transactions",
            params={"entity_id": ace_leopard["id"]},
        ).status_code == 404
        assert client.patch(
            f"/v1/finance/transactions/{jingao_transactions.json()[0]['id']}",
            params={"entity_id": ace_leopard["id"]},
            json={
                "category": "越权修改",
                "note": None,
                "pm_project_id": None,
            },
        ).status_code == 404
        assert client.post(
            f"/v1/finance/statements/{jingao_batch['id']}/confirm",
            params={"entity_id": ace_leopard["id"]},
        ).status_code == 404

        assert client.post(
            f"/v1/finance/statements/{jingao_batch['id']}/confirm",
            params={"entity_id": jingao["id"]},
        ).status_code == 200
        assert client.post(
            f"/v1/finance/statements/{ace_batch['id']}/confirm",
            params={"entity_id": ace_leopard["id"]},
        ).status_code == 200

        period = {"from_date": "2026-08-01", "to_date": "2026-08-07"}
        jingao_dashboard = client.get(
            "/v1/finance/dashboard",
            params={**period, "entity_id": jingao["id"]},
        )
        ace_dashboard = client.get(
            "/v1/finance/dashboard",
            params={**period, "entity_id": ace_leopard["id"]},
        )
        power_dashboard = client.get(
            "/v1/finance/dashboard",
            params={**period, "entity_id": power_leopard["id"]},
        )
        assert jingao_dashboard.status_code == 200
        assert ace_dashboard.status_code == 200
        assert power_dashboard.status_code == 200
        assert jingao_dashboard.json()["income"] == "12000.00"
        assert jingao_dashboard.json()["expense"] == "2000.00"
        assert jingao_dashboard.json()["bank_balance"] == "60000.00"
        assert jingao_dashboard.json()["weekly_top_income"][0]["name"] == "京奥客户"
        assert ace_dashboard.json()["income"] == "33000.00"
        assert ace_dashboard.json()["expense"] == "7000.00"
        assert ace_dashboard.json()["bank_balance"] == "126000.00"
        assert ace_dashboard.json()["weekly_top_income"][0]["name"] == "三角洲客户"
        assert power_dashboard.json()["income"] == "0"
        assert power_dashboard.json()["expense"] == "0"
        assert power_dashboard.json()["accounts"] == []

        xingyao = entity_by_key["xingyao"]
        assert xingyao["default_bank_name"] == ""
        assert client.get("/v1/finance/statements", params={"entity_id": xingyao["id"]}).json() == []
        xingyao_upload = client.post(
            "/v1/finance/statements/upload",
            data={"entity_id": xingyao["id"], "bank_name": "测试银行",
                  "account_name": "星曜基本户", "account_number": "20000100000000000004"},
            files={"file": ("星曜流水.xlsx", _company_xlsx(
                serial_prefix="XINGYAO", income=Decimal("8000"), expense=Decimal("1500"),
                opening_balance=Decimal("20000"), income_counterparty="教培客户",
                expense_counterparty="教培供应商",
            ), "application/octet-stream")},
        )
        assert xingyao_upload.status_code == 200
        xingyao_batch = xingyao_upload.json()
        assert xingyao_batch["entity_id"] == xingyao["id"]
        assert client.post(
            f"/v1/finance/statements/{xingyao_batch['id']}/confirm",
            params={"entity_id": xingyao["id"]},
        ).status_code == 200
        xingyao_dashboard = client.get(
            "/v1/finance/dashboard", params={**period, "entity_id": xingyao["id"]},
        )
        assert xingyao_dashboard.status_code == 200
        assert xingyao_dashboard.json()["income"] == "8000.00"
        assert xingyao_dashboard.json()["expense"] == "1500.00"
        assert xingyao_dashboard.json()["bank_balance"] == "26500.00"
        assert client.get(
            f"/v1/finance/statements/{xingyao_batch['id']}/transactions",
            params={"entity_id": jingao["id"]},
        ).status_code == 404

        jingao_cash_payload = {
            "entity_id": jingao["id"],
            "company_name": jingao["name"],
            "cash_account": "公司现金",
            "transaction_date": "2026-08-03",
            "direction": "opening",
            "amount": "8888.00",
            "category": "期初现金",
            "note": "仅京奥总公司账外现金",
        }
        created_cash = client.post("/v1/finance/cash", json=jingao_cash_payload)
        assert created_cash.status_code == 200
        assert client.get(
            "/v1/finance/cash", params={"entity_id": jingao["id"]}
        ).json()[0]["amount"] == "8888.00"
        assert client.get(
            "/v1/finance/cash", params={"entity_id": ace_leopard["id"]}
        ).json() == []
        assert client.post(
            "/v1/finance/cash",
            json={
                **jingao_cash_payload,
                "entity_id": ace_leopard["id"],
                "company_name": ace_leopard["name"],
            },
        ).status_code == 403
        assert db.scalar(select(func.count(CashEntry.id))) == 1

        jingao_dashboard_after_cash = client.get(
            "/v1/finance/dashboard",
            params={**period, "entity_id": jingao["id"]},
        ).json()
        ace_dashboard_after_cash = client.get(
            "/v1/finance/dashboard",
            params={**period, "entity_id": ace_leopard["id"]},
        ).json()
        assert jingao_dashboard_after_cash["cash_balance"] == "8888.00"
        assert jingao_dashboard_after_cash["bank_balance"] == "60000.00"
        assert ace_dashboard_after_cash["cash_balance"] == "0"
        assert ace_dashboard_after_cash["bank_balance"] == "126000.00"
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_finance_entity_seed_normalizes_finance_projects_and_active_registry(
    tmp_path,
    monkeypatch,
) -> None:
    db, users = _database()
    canonical = BusinessEntity(
        name="京奥电竞（北京）科技有限公司",
        short_name="旧简称",
        created_by_user_id=users["founder"].id,
    )
    legacy_alias = BusinessEntity(
        name="京奥电竞",
        short_name="项目管理仍可能使用",
        created_by_user_id=users["founder"].id,
    )
    legacy_cash_owner = BusinessEntity(
        name="掏喵",
        created_by_user_id=users["founder"].id,
    )
    db.add_all([canonical, legacy_alias, legacy_cash_owner])
    db.flush()
    legacy_project = ManagedProject(
        project_no="FEISHU-LEGACY-01",
        name="旧主体项目",
        entity_id=legacy_alias.id,
        company_name="京奥电竞",
        client="测试甲方",
        business_category="电竞培训",
        manager_user_id=users["business"].id,
        members_json='["执行成员"]',
        planned_start=date(2026, 8, 1),
        planned_end=date(2026, 8, 31),
        objective="验证公司主体迁移",
        created_by_user_id=users["business"].id,
    )
    db.add(legacy_project)

    target_account = FinancialAccount(
        entity_id=canonical.id,
        bank_name="北京银行成寿寺支行",
        account_name="京奥基本户",
        account_number_masked="2000****0001",
        account_number_hash="SAME-ACCOUNT-HASH",
    )
    legacy_account = FinancialAccount(
        entity_id=legacy_alias.id,
        bank_name="北京银行成寿寺支行",
        account_name="旧别名基本户",
        account_number_masked="2000****0001",
        account_number_hash="SAME-ACCOUNT-HASH",
    )
    db.add_all([target_account, legacy_account])
    db.flush()
    target_batch = BankStatementBatch(
        entity_id=canonical.id,
        account_id=target_account.id,
        original_filename="已迁移流水.xlsx",
        source_path="/finance/canonical.xlsx",
        file_hash="SAME-FILE-HASH",
        period_start=date(2026, 7, 1),
        period_end=date(2026, 7, 2),
        status="confirmed",
        row_count=1,
        uploaded_by_user_id=users["finance"].id,
    )
    legacy_batch = BankStatementBatch(
        entity_id=legacy_alias.id,
        account_id=legacy_account.id,
        original_filename="旧别名流水.xlsx",
        source_path="/finance/legacy.xlsx",
        file_hash="SAME-FILE-HASH",
        period_start=date(2026, 7, 1),
        period_end=date(2026, 7, 2),
        status="confirmed",
        row_count=2,
        uploaded_by_user_id=users["finance"].id,
    )
    db.add_all([target_batch, legacy_batch])
    db.flush()
    db.add_all([
        BankTransaction(
            batch_id=target_batch.id,
            entity_id=canonical.id,
            account_id=target_account.id,
            transacted_at=datetime(2026, 7, 1, tzinfo=timezone.utc),
            income=Decimal("100.00"),
            expense=Decimal("0"),
            balance=Decimal("100.00"),
            counterparty="已存在客户",
            fingerprint="DUPLICATE-FINGERPRINT",
            status="confirmed",
        ),
        BankTransaction(
            batch_id=legacy_batch.id,
            entity_id=legacy_alias.id,
            account_id=legacy_account.id,
            transacted_at=datetime(2026, 7, 1, tzinfo=timezone.utc),
            income=Decimal("100.00"),
            expense=Decimal("0"),
            balance=Decimal("100.00"),
            counterparty="重复客户",
            fingerprint="DUPLICATE-FINGERPRINT",
            status="confirmed",
        ),
        BankTransaction(
            batch_id=legacy_batch.id,
            entity_id=legacy_alias.id,
            account_id=legacy_account.id,
            transacted_at=datetime(2026, 7, 2, tzinfo=timezone.utc),
            income=Decimal("200.00"),
            expense=Decimal("0"),
            balance=Decimal("300.00"),
            counterparty="旧别名独有客户",
            fingerprint="UNIQUE-LEGACY-FINGERPRINT",
            status="confirmed",
        ),
        CashEntry(
            entity_id=legacy_cash_owner.id,
            cash_account="公司现金",
            transaction_date=date(2026, 8, 1),
            direction="expense",
            amount=Decimal("1333.00"),
            category="历史账外现金",
            note="迁回京奥总部",
            created_by_user_id=users["founder"].id,
        ),
    ])
    db.commit()

    client = _configure(monkeypatch, tmp_path, db, users["finance"])
    try:
        response = client.get("/v1/finance/entities")
        assert response.status_code == 200
        headquarters = next(
            item for item in response.json() if item["key"] == "jingao"
        )
        assert headquarters["id"] == canonical.id

        accounts = db.scalars(select(FinancialAccount)).all()
        batches = db.scalars(select(BankStatementBatch)).all()
        transactions = db.scalars(select(BankTransaction)).all()
        cash_entries = db.scalars(select(CashEntry)).all()
        assert len(accounts) == 1
        assert accounts[0].id == target_account.id
        assert accounts[0].entity_id == canonical.id
        assert len(batches) == 1
        assert batches[0].id == target_batch.id
        assert batches[0].entity_id == canonical.id
        assert len(transactions) == 2
        assert {item.fingerprint for item in transactions} == {
            "DUPLICATE-FINGERPRINT",
            "UNIQUE-LEGACY-FINGERPRINT",
        }
        assert {item.entity_id for item in transactions} == {canonical.id}
        assert {item.account_id for item in transactions} == {target_account.id}
        assert {item.batch_id for item in transactions} == {target_batch.id}
        assert len(cash_entries) == 1
        assert cash_entries[0].entity_id == canonical.id

        migrated_project = db.get(ManagedProject, legacy_project.id)
        assert migrated_project.entity_id == canonical.id
        assert migrated_project.company_name == canonical.name

        # Legacy rows remain for audit history but are no longer selectable.
        assert db.get(BusinessEntity, legacy_alias.id) is not None
        assert db.get(BusinessEntity, legacy_cash_owner.id) is not None
        assert db.get(BusinessEntity, legacy_alias.id).short_name == "项目管理仍可能使用"
        assert db.get(BusinessEntity, legacy_alias.id).active is False
        assert db.get(BusinessEntity, legacy_cash_owner.id).active is False
        active_entities = db.scalars(
            select(BusinessEntity).where(BusinessEntity.active.is_(True))
        ).all()
        assert {item.name for item in active_entities} == {
            "京奥电竞（北京）科技有限公司",
            "王牌猎豹（JAG三角洲）",
            "劲腾豹跃（JAG王者）",
            "星曜电竞",
        }
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_annual_dashboard_is_natural_year_confirmed_only_and_entity_scoped(
    tmp_path,
    monkeypatch,
) -> None:
    db, users = _database()
    jingao = BusinessEntity(
        name="京奥电竞（北京）科技有限公司",
        short_name="京奥电竞",
        created_by_user_id=users["founder"].id,
    )
    ace = BusinessEntity(
        name="王牌猎豹（JAG三角洲）",
        short_name="王牌猎豹",
        created_by_user_id=users["founder"].id,
    )
    db.add_all([jingao, ace])
    db.flush()
    jingao_account = FinancialAccount(
        entity_id=jingao.id,
        bank_name="北京银行成寿寺支行",
        account_name="基本户",
        account_number_masked="2000****6862",
        account_number_hash="ANNUAL-JINGAO",
    )
    ace_account = FinancialAccount(
        entity_id=ace.id,
        bank_name="测试银行",
        account_name="基本户",
        account_number_masked="1000****0001",
        account_number_hash="ANNUAL-ACE",
    )
    db.add_all([jingao_account, ace_account])
    db.flush()
    confirmed = BankStatementBatch(
        entity_id=jingao.id,
        account_id=jingao_account.id,
        original_filename="2025前三季度.xlsx",
        source_path="/finance/2025-confirmed.xlsx",
        file_hash="ANNUAL-CONFIRMED",
        period_start=date(2025, 3, 10),
        period_end=date(2025, 9, 30),
        status="confirmed",
        row_count=4,
        uploaded_by_user_id=users["finance"].id,
        confirmed_by_user_id=users["finance"].id,
        confirmed_at=datetime.now(timezone.utc),
    )
    pending = BankStatementBatch(
        entity_id=jingao.id,
        account_id=jingao_account.id,
        original_filename="2025第四季度.xlsx",
        source_path="/finance/2025-pending.xlsx",
        file_hash="ANNUAL-PENDING",
        period_start=date(2025, 10, 1),
        period_end=date(2025, 12, 31),
        status="pending",
        row_count=1,
        uploaded_by_user_id=users["finance"].id,
    )
    other_batch = BankStatementBatch(
        entity_id=ace.id,
        account_id=ace_account.id,
        original_filename="2025王牌猎豹.xlsx",
        source_path="/finance/ace-2025.xlsx",
        file_hash="ANNUAL-ACE-FILE",
        period_start=date(2025, 1, 1),
        period_end=date(2025, 12, 31),
        status="confirmed",
        row_count=1,
        uploaded_by_user_id=users["finance"].id,
        confirmed_by_user_id=users["finance"].id,
        confirmed_at=datetime.now(timezone.utc),
    )
    db.add_all([confirmed, pending, other_batch])
    db.flush()

    def transaction(
        batch: BankStatementBatch,
        entity: BusinessEntity,
        account: FinancialAccount,
        when: datetime,
        income: str,
        expense: str,
        balance: str,
        counterparty: str,
        fingerprint: str,
        *,
        status: str = "confirmed",
        category: str = "待确认",
    ) -> BankTransaction:
        return BankTransaction(
            batch_id=batch.id,
            entity_id=entity.id,
            account_id=account.id,
            transacted_at=when,
            income=Decimal(income),
            expense=Decimal(expense),
            balance=Decimal(balance),
            counterparty=counterparty,
            summary="年度测试",
            category=category,
            fingerprint=fingerprint,
            status=status,
        )

    db.add_all([
        transaction(confirmed, jingao, jingao_account, datetime(2025, 3, 10, tzinfo=timezone.utc), "100", "0", "1100", "客户A", "A-1"),
        transaction(confirmed, jingao, jingao_account, datetime(2025, 3, 11, tzinfo=timezone.utc), "0", "20", "1080", "供应商A", "A-2", category="场地"),
        transaction(confirmed, jingao, jingao_account, datetime(2025, 7, 1, tzinfo=timezone.utc), "200", "0", "1280", "客户A", "A-3"),
        transaction(confirmed, jingao, jingao_account, datetime(2025, 9, 30, tzinfo=timezone.utc), "0", "80", "1200", "供应商B", "A-4", category="服务费"),
        transaction(pending, jingao, jingao_account, datetime(2025, 10, 1, tzinfo=timezone.utc), "9999", "0", "11199", "待确认客户", "P-1", status="pending"),
        transaction(other_batch, ace, ace_account, datetime(2025, 6, 1, tzinfo=timezone.utc), "8888", "0", "8888", "其他公司客户", "ACE-1"),
    ])
    db.commit()
    client = _configure(monkeypatch, tmp_path, db, users["finance"])
    try:
        years = client.get(
            "/v1/finance/annual-years", params={"entity_id": jingao.id}
        )
        assert years.status_code == 200
        assert years.json()["years"] == [2025]

        response = client.get(
            "/v1/finance/annual",
            params={"entity_id": jingao.id, "year": 2025},
        )
        assert response.status_code == 200
        payload = response.json()
        assert payload["income"] == "300.00"
        assert payload["expense"] == "100.00"
        assert payload["net"] == "200.00"
        assert payload["opening_balance"] == "1000.00"
        assert payload["closing_balance"] == "1200.00"
        assert payload["transaction_count"] == 4
        assert payload["monthly"][2]["income"] == "100.00"
        assert payload["monthly"][6]["income"] == "200.00"
        assert payload["top_income"][0]["name"] == "客户A"
        assert payload["top_income"][0]["amount"] == "300.00"
        assert payload["coverage"]["data_complete"] is False
        assert payload["coverage"]["pending_batch_count"] == 1
        assert payload["coverage"]["pending_batches"][0]["filename"] == "2025第四季度.xlsx"
        assert payload["unclassified"]["count"] == 2
        assert all(item["counterparty"] != "其他公司客户" for item in payload["anomalies"])
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_annual_transaction_search_is_keyword_year_entity_and_permission_scoped(
    tmp_path,
    monkeypatch,
) -> None:
    db, users = _database()
    entity = BusinessEntity(
        name="京奥电竞（北京）科技有限公司",
        short_name="京奥电竞",
        created_by_user_id=users["founder"].id,
    )
    other_entity = BusinessEntity(
        name="王牌猎豹（JAG三角洲）",
        short_name="王牌猎豹",
        created_by_user_id=users["founder"].id,
    )
    db.add_all([entity, other_entity])
    db.flush()
    account = FinancialAccount(
        entity_id=entity.id,
        bank_name="北京银行成寿寺支行",
        account_name="基本户",
        account_number_masked="2000****6862",
        account_number_hash="SEARCH-JINGAO",
    )
    other_account = FinancialAccount(
        entity_id=other_entity.id,
        bank_name="测试银行",
        account_name="基本户",
        account_number_masked="1000****0001",
        account_number_hash="SEARCH-OTHER",
    )
    db.add_all([account, other_account])
    db.flush()
    batch = BankStatementBatch(
        entity_id=entity.id,
        account_id=account.id,
        original_filename="京奥2025年度流水.xlsx",
        source_path="/finance/jingao-2025.xlsx",
        file_hash="SEARCH-BATCH",
        period_start=date(2025, 1, 1),
        period_end=date(2025, 12, 31),
        status="confirmed",
        row_count=2,
        uploaded_by_user_id=users["finance"].id,
        confirmed_by_user_id=users["finance"].id,
        confirmed_at=datetime.now(timezone.utc),
    )
    pending_batch = BankStatementBatch(
        entity_id=entity.id,
        account_id=account.id,
        original_filename="京奥待确认.xlsx",
        source_path="/finance/pending.xlsx",
        file_hash="SEARCH-PENDING",
        period_start=date(2025, 1, 1),
        period_end=date(2025, 1, 31),
        status="pending",
        row_count=1,
        uploaded_by_user_id=users["finance"].id,
    )
    other_batch = BankStatementBatch(
        entity_id=other_entity.id,
        account_id=other_account.id,
        original_filename="其他公司2025.xlsx",
        source_path="/finance/other.xlsx",
        file_hash="SEARCH-OTHER-BATCH",
        period_start=date(2025, 1, 1),
        period_end=date(2025, 12, 31),
        status="confirmed",
        row_count=1,
        uploaded_by_user_id=users["finance"].id,
    )
    db.add_all([batch, pending_batch, other_batch])
    db.flush()
    db.add_all([
        BankTransaction(
            batch_id=batch.id,
            entity_id=entity.id,
            account_id=account.id,
            transacted_at=datetime(2025, 6, 18, 8, 30, tzinfo=timezone.utc),
            income=Decimal("0"),
            expense=Decimal("100000"),
            balance=Decimal("520000"),
            counterparty="杭州星河会展有限公司",
            summary="场馆服务费",
            note="KPL发布会场租",
            category="场地费用",
            project_reference="Cc2508",
            bank_serial="SEARCH-001",
            fingerprint="SEARCH-TX-1",
            status="confirmed",
        ),
        BankTransaction(
            batch_id=pending_batch.id,
            entity_id=entity.id,
            account_id=account.id,
            transacted_at=datetime(2025, 1, 2, tzinfo=timezone.utc),
            income=Decimal("999"),
            expense=Decimal("0"),
            balance=Decimal("999"),
            counterparty="未确认星河客户",
            fingerprint="SEARCH-PENDING-TX",
            status="pending",
        ),
        BankTransaction(
            batch_id=other_batch.id,
            entity_id=other_entity.id,
            account_id=other_account.id,
            transacted_at=datetime(2025, 6, 18, tzinfo=timezone.utc),
            income=Decimal("100000"),
            expense=Decimal("0"),
            balance=Decimal("100000"),
            counterparty="其他星河客户",
            fingerprint="SEARCH-OTHER-TX",
            status="confirmed",
        ),
    ])
    db.commit()
    client = _configure(monkeypatch, tmp_path, db, users["finance"])
    try:
        by_party = client.get(
            "/v1/finance/annual-transactions/search",
            params={"entity_id": entity.id, "year": 2025, "q": "杭州星河"},
        )
        assert by_party.status_code == 200
        assert by_party.json()["total"] == 1
        assert by_party.json()["items"][0]["counterparty"] == "杭州星河会展有限公司"
        assert by_party.json()["items"][0]["matched_fields"] == ["往来单位"]

        by_purpose = client.get(
            "/v1/finance/annual-transactions/search",
            params={"entity_id": entity.id, "year": 2025, "q": "发布会场租"},
        )
        assert by_purpose.json()["items"][0]["note"] == "KPL发布会场租"

        by_amount = client.get(
            "/v1/finance/annual-transactions/search",
            params={"entity_id": entity.id, "year": 2025, "q": "100,000.00"},
        )
        assert by_amount.json()["total"] == 1
        assert by_amount.json()["items"][0]["expense"] == "100000.00"

        app.dependency_overrides[current_user] = lambda: users["business"]
        forbidden = client.get(
            "/v1/finance/annual-transactions/search",
            params={"entity_id": entity.id, "year": 2025, "q": "星河"},
        )
        assert forbidden.status_code == 403
    finally:
        app.dependency_overrides.clear()
        db.close()
