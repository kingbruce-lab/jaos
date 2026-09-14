from __future__ import annotations

from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from docx import Document as WordDocument
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app import ingest, main
from app.auth import current_user
from app.contracts import CONTRACT_CATEGORIES, ensure_contract_layout
from app.database import get_db
from app.main import app
from app.models import AuditLog, Base, ContractDocumentSource, Document, User
from app.retrieval import search


def _database() -> tuple[Session, dict[str, User]]:
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    users = {
        "l3": User(
            username="contract-l3",
            display_name="L3员工",
            password_hash="x",
            role="employee",
            organization_role="business",
            confidentiality_ceiling="L3",
        ),
        "l4": User(
            username="contract-l4",
            display_name="L4资料管理员",
            password_hash="x",
            role="knowledge_admin",
            organization_role="administrative",
            confidentiality_ceiling="L4",
        ),
        "l5": User(
            username="contract-l5",
            display_name="L5创始人",
            password_hash="x",
            role="founder",
            organization_role="management",
            confidentiality_ceiling="L5",
        ),
        "l4_management": User(
            username="contract-l4-management",
            display_name="L4管理",
            password_hash="x",
            role="founder",
            organization_role="management",
            confidentiality_ceiling="L4",
        ),
        "l5_admin": User(
            username="contract-l5-admin",
            display_name="L5行政",
            password_hash="x",
            role="knowledge_admin",
            organization_role="administrative",
            confidentiality_ceiling="L5",
        ),
        "l4_personnel": User(
            username="contract-l4-personnel",
            display_name="L4人事",
            password_hash="x",
            role="employee",
            organization_role="personnel",
            confidentiality_ceiling="L4",
        ),
        "l5_business": User(
            username="contract-l5-business",
            display_name="L5业务",
            password_hash="x",
            role="employee",
            organization_role="business",
            confidentiality_ceiling="L5",
        ),
        "l5_finance": User(
            username="contract-l5-finance",
            display_name="L5财务",
            password_hash="x",
            role="employee",
            organization_role="finance",
            confidentiality_ceiling="L5",
        ),
    }
    db = Session(engine, expire_on_commit=False)
    db.add_all(users.values())
    db.commit()
    return db, users


def _word_payload(text: str = "京奥电竞业务合同，合作期限一年。") -> bytes:
    document = WordDocument()
    document.add_heading("合同测试件", level=1)
    document.add_paragraph(text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _configure(
    monkeypatch,
    db: Session,
    user: User,
    knowledge_root,
) -> None:
    knowledge_root.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr(
        main,
        "settings",
        SimpleNamespace(
            knowledge_root=knowledge_root,
            inbox_max_file_bytes=4 * 1024 * 1024,
            policy_version="test-policy",
        ),
    )
    monkeypatch.setattr(
        ingest,
        "settings",
        SimpleNamespace(
            copy_sources=False,
            managed_source_dir=knowledge_root / "managed",
            inbox_dir=knowledge_root,
            inbox_settle_seconds=0,
            inbox_max_file_bytes=4 * 1024 * 1024,
        ),
    )

    def override_db():
        yield db

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[current_user] = lambda: user


def test_contract_layout_is_fixed_by_category_and_level(tmp_path) -> None:
    layout = ensure_contract_layout(tmp_path)

    assert set(layout) == set(CONTRACT_CATEGORIES)
    assert layout["administrative"] == (
        tmp_path / "合同档案库" / "行政合同" / "L4"
    ).resolve()
    assert layout["personnel"].is_dir()
    assert layout["business"].is_dir()
    assert layout["executive_office"] == (
        tmp_path / "合同档案库" / "总办合同" / "L5"
    ).resolve()


def test_contract_category_visibility_obeys_l4_l5_ceiling(
    tmp_path,
    monkeypatch,
) -> None:
    db, users = _database()
    _configure(monkeypatch, db, users["l3"], tmp_path / "knowledge")
    client = TestClient(app)
    try:
        assert client.get("/v1/contracts/categories").status_code == 403

        app.dependency_overrides[current_user] = lambda: users["l4"]
        l4 = client.get("/v1/contracts/categories")
        assert l4.status_code == 200
        assert {item["key"] for item in l4.json()} == {
            "administrative",
            "personnel",
            "business",
            "executive_office",
        }
        l4_policy = {item["key"]: item for item in l4.json()}
        assert l4_policy["administrative"]["can_search"] is True
        assert l4_policy["business"]["can_search"] is True
        assert l4_policy["personnel"]["can_search"] is False
        assert l4_policy["executive_office"]["can_search"] is True
        assert l4_policy["executive_office"]["search_scope"] == "own"
        assert all(item["can_upload"] for item in l4.json())

        app.dependency_overrides[current_user] = lambda: users["l5"]
        l5 = client.get("/v1/contracts/categories")
        assert l5.status_code == 200
        assert {item["key"] for item in l5.json()} == set(CONTRACT_CATEGORIES)
        assert all(item["can_search"] for item in l5.json())
        assert all(item["can_upload"] for item in l5.json())
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_contract_upload_is_forced_to_review_and_exact_nas_folder(
    tmp_path,
    monkeypatch,
) -> None:
    db, users = _database()
    knowledge_root = tmp_path / "knowledge"
    _configure(monkeypatch, db, users["l4"], knowledge_root)
    client = TestClient(app)
    try:
        response = client.post(
            "/v1/contracts/uploads",
            data={"category": "business"},
            files={
                "file": (
                    "赛事执行合同.docx",
                    _word_payload(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        document = db.scalar(select(Document))
        audit = db.scalar(
            select(AuditLog).where(AuditLog.action == "contract_upload")
        )

        assert response.status_code == 200
        assert response.json()["review_required"] is True
        assert document is not None
        assert document.role == "contract"
        assert document.confidentiality == "L4"
        assert document.project.confidentiality == "L4"
        assert document.project.domain == "contract_business"
        assert document.knowledge_status == "candidate"
        assert (
            knowledge_root
            / "合同档案库"
            / "业务合同"
            / "L4"
            / "赛事与活动"
            / "赛事执行合同.docx"
        ).is_file()
        assert response.json()["filing"]["mode"] == "local_ai"
        assert response.json()["filing"]["folder_path"] == "赛事与活动"
        assert audit is not None
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_administrative_can_list_create_folders_and_move_own_contract(
    tmp_path,
    monkeypatch,
) -> None:
    db, users = _database()
    knowledge_root = tmp_path / "knowledge"
    _configure(monkeypatch, db, users["l4"], knowledge_root)
    client = TestClient(app)
    try:
        upload = client.post(
            "/v1/contracts/uploads",
            data={
                "category": "administrative",
                "relative_path": "待整理/租赁原件.docx",
            },
            files={
                "file": (
                    "租赁原件.docx",
                    _word_payload("办公场地租赁、物业服务和租金约定。"),
                    "application/octet-stream",
                )
            },
        )
        assert upload.status_code == 200
        document_id = upload.json()["document_id"]

        mine = client.get("/v1/contracts/mine")
        assert mine.status_code == 200
        assert mine.json()["count"] == 1
        assert mine.json()["items"][0]["knowledge_status"] == "candidate"
        assert mine.json()["items"][0]["folder_path"] == "待整理"
        assert mine.json()["items"][0]["can_move"] is True

        created = client.post(
            "/v1/contracts/folders",
            json={"category": "administrative", "folder_path": "2026年/租赁合同"},
        )
        assert created.status_code == 200
        folders = client.get(
            "/v1/contracts/folders", params={"category": "administrative"}
        )
        assert "2026年/租赁合同" in folders.json()["items"]

        moved = client.patch(
            f"/v1/contracts/{document_id}/folder",
            json={"folder_path": "2026年/租赁合同"},
        )
        assert moved.status_code == 200
        assert moved.json()["folder_path"] == "2026年/租赁合同"
        document = db.get(Document, document_id)
        assert document is not None
        assert Path(document.file_blob.source_path).is_file()
        assert "2026年/租赁合同" in Path(document.file_blob.source_path).as_posix()
        assert not list(knowledge_root.rglob("待整理/租赁原件.docx"))
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_executive_office_internal_folder_is_always_available(
    tmp_path,
    monkeypatch,
) -> None:
    db, users = _database()
    knowledge_root = tmp_path / "knowledge"
    _configure(monkeypatch, db, users["l4"], knowledge_root)
    client = TestClient(app)
    try:
        folders = client.get(
            "/v1/contracts/folders", params={"category": "executive_office"}
        )
        assert folders.status_code == 200
        assert "内部资料（密）" in folders.json()["items"]
        assert (
            knowledge_root
            / "合同档案库"
            / "总办合同"
            / "L5"
            / "内部资料（密）"
        ).is_dir()
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_l4_administrative_can_upload_and_read_own_executive_office_contracts(
    tmp_path,
    monkeypatch,
) -> None:
    db, users = _database()
    _configure(monkeypatch, db, users["l4"], tmp_path / "knowledge")
    client = TestClient(app)
    try:
        empty = client.get(
            "/v1/contracts", params={"category": "executive_office"}
        )
        assert empty.status_code == 200
        assert empty.json()["items"] == []
        upload = client.post(
            "/v1/contracts/uploads",
            data={
                "category": "executive_office",
                "relative_path": "涉密系列/第一批/总办合同.docx",
            },
            files={"file": ("总办合同.docx", _word_payload(), "application/octet-stream")},
        )
        assert upload.status_code == 200
        document = db.get(Document, upload.json()["document_id"])
        waiting = client.get(
            "/v1/contracts", params={"category": "executive_office"}
        )
        assert waiting.json()["pending_count"] == 1
        assert waiting.json()["items"] == []
        document.knowledge_status = "approved"
        document.project.knowledge_status = "approved"
        db.commit()
        listed = client.get(
            "/v1/contracts", params={"category": "executive_office"}
        )
        assert [item["document_id"] for item in listed.json()["items"]] == [document.id]
        assert listed.json()["pending_count"] == 0
        assert listed.json()["items"][0]["folder_path"] == "涉密系列/第一批"
        assert list((tmp_path / "knowledge").rglob("涉密系列/第一批/总办合同.docx"))
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_management_l5_can_upload_contracts(
    tmp_path,
    monkeypatch,
) -> None:
    db, users = _database()
    knowledge_root = tmp_path / "knowledge"
    _configure(monkeypatch, db, users["l5"], knowledge_root)
    client = TestClient(app)
    try:
        assert client.get("/v1/contracts/categories").status_code == 200
        response = client.post(
            "/v1/contracts/uploads",
            data={"category": "administrative"},
            files={
                "file": (
                    "管理上传.docx",
                    _word_payload(),
                    "application/octet-stream",
                )
            },
        )

        assert response.status_code == 200
        assert response.json()["review_required"] is True
        assert list(knowledge_root.rglob("管理上传.docx"))
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_exact_duplicate_contract_does_not_leave_a_second_nas_file(
    tmp_path,
    monkeypatch,
) -> None:
    db, users = _database()
    knowledge_root = tmp_path / "knowledge"
    _configure(monkeypatch, db, users["l4"], knowledge_root)
    client = TestClient(app)
    payload = _word_payload("同一份系列合同。")
    try:
        first = client.post(
            "/v1/contracts/uploads",
            data={
                "category": "administrative",
                "relative_path": "系列合同/原件.docx",
            },
            files={"file": ("原件.docx", payload, "application/octet-stream")},
        )
        duplicate = client.post(
            "/v1/contracts/uploads",
            data={
                "category": "administrative",
                "relative_path": "系列合同/原件.docx",
            },
            files={"file": ("原件.docx", payload, "application/octet-stream")},
        )

        assert first.status_code == 200
        assert duplicate.status_code == 200
        assert duplicate.json()["duplicate_filtered"] is True
        assert [path.name for path in knowledge_root.rglob("原件*.docx")] == ["原件.docx"]
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_contract_search_scope_obeys_organization_role_matrix(
    tmp_path,
    monkeypatch,
) -> None:
    db, users = _database()
    _configure(monkeypatch, db, users["l5_admin"], tmp_path / "knowledge")
    client = TestClient(app)
    try:
        admin_categories = client.get("/v1/contracts/categories")
        assert admin_categories.status_code == 200
        admin_policy = {
            item["key"]: item for item in admin_categories.json()
        }
        assert {
            key for key, item in admin_policy.items() if item["can_search"]
        } == {"administrative", "business", "executive_office"}
        assert admin_policy["executive_office"]["search_scope"] == "own"
        assert set(admin_policy) == set(CONTRACT_CATEGORIES)

        for denied_category in ("personnel",):
            assert client.get(
                "/v1/contracts", params={"category": denied_category}
            ).status_code == 403
            assert client.post(
                "/v1/contracts/search",
                json={"category": denied_category, "query": "合同"},
            ).status_code == 403

        app.dependency_overrides[current_user] = lambda: users["l4_personnel"]
        personnel_categories = client.get("/v1/contracts/categories")
        assert personnel_categories.status_code == 200
        assert personnel_categories.json() == [
            {
                "key": "personnel",
                "name": "人事合同",
                "confidentiality": "L4",
                "can_search": True,
                "can_upload": False,
                "search_scope": "all",
            },
            {
                "key": "executive_office",
                "name": "总办合同",
                "confidentiality": "L5",
                "can_search": True,
                "can_upload": True,
                "search_scope": "own",
            }
        ]
        assert client.get(
            "/v1/contracts", params={"category": "personnel"}
        ).status_code == 200
        assert client.get(
            "/v1/contracts", params={"category": "administrative"}
        ).status_code == 403

        app.dependency_overrides[current_user] = lambda: users["l5_business"]
        assert client.get("/v1/contracts/categories").status_code == 403
        for category in CONTRACT_CATEGORIES:
            assert client.get(
                "/v1/contracts", params={"category": category}
            ).status_code == 403

        app.dependency_overrides[current_user] = lambda: users["l5_finance"]
        finance_categories = client.get("/v1/contracts/categories")
        assert finance_categories.status_code == 200
        assert [item["key"] for item in finance_categories.json()] == ["executive_office"]
        assert finance_categories.json()[0]["search_scope"] == "own"
        for category in ("administrative", "personnel", "business"):
            assert client.get(
                "/v1/contracts", params={"category": category}
            ).status_code == 403

        for actor in (users["l5_business"],):
            app.dependency_overrides[current_user] = lambda actor=actor: actor
            assert client.get("/v1/contracts/categories").status_code == 403
            for category in CONTRACT_CATEGORIES:
                assert client.get(
                    "/v1/contracts", params={"category": category}
                ).status_code == 403
                assert client.post(
                    "/v1/contracts/search",
                    json={"category": category, "query": "合同"},
                ).status_code == 403

        for actor in (
            users["l4_management"],
            users["l4_personnel"],
            users["l5_business"],
            users["l5_finance"],
        ):
            app.dependency_overrides[current_user] = lambda actor=actor: actor
            denied_upload = client.post(
                "/v1/contracts/uploads",
                data={"category": "administrative"},
                files={
                    "file": (
                        f"{actor.organization_role}-denied.docx",
                        _word_payload(),
                        "application/octet-stream",
                    )
                },
            )
            assert denied_upload.status_code == 403
            assert denied_upload.json()["detail"] == "无权向该合同分类上传资料"
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_same_file_hash_gets_separate_contract_security_context(
    tmp_path,
    monkeypatch,
) -> None:
    db, users = _database()
    knowledge_root = tmp_path / "knowledge"
    _configure(monkeypatch, db, users["l5_admin"], knowledge_root)
    payload = _word_payload("同一物理内容用于权限隔离测试。")
    ordinary = knowledge_root / "公共资料" / "L1" / "示例.docx"
    ordinary.parent.mkdir(parents=True)
    ordinary.write_bytes(payload)
    ingest.ingest_document(
        db,
        {
            "path": str(ordinary),
            "project_name": "公开示例",
            "domain": "company",
            "role": "proposal",
            "confidentiality": "L1",
            "knowledge_status": "candidate",
        },
    )
    client = TestClient(app)
    try:
        response = client.post(
            "/v1/contracts/uploads",
            data={"category": "administrative"},
            files={"file": ("行政示例.docx", payload, "application/octet-stream")},
        )
        documents = db.scalars(select(Document).order_by(Document.confidentiality)).all()

        assert response.status_code == 200
        assert len(documents) == 2
        assert {(item.role, item.confidentiality) for item in documents} == {
            ("proposal", "L1"),
            ("contract", "L4"),
        }
        assert documents[0].content_hash == documents[1].content_hash
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_same_contract_bytes_can_move_independently_across_l4_l5_archives(
    tmp_path,
    monkeypatch,
) -> None:
    db, users = _database()
    knowledge_root = tmp_path / "knowledge"
    _configure(monkeypatch, db, users["l5_admin"], knowledge_root)
    client = TestClient(app)
    payload = _word_payload("同一份股权协议分别保存在业务合同和总办合同。")
    try:
        business = client.post(
            "/v1/contracts/uploads",
            data={
                "category": "business",
                "relative_path": "合作方/股权转让协议.docx",
            },
            files={"file": ("股权转让协议.docx", payload, "application/octet-stream")},
        )
        executive = client.post(
            "/v1/contracts/uploads",
            data={"category": "executive_office"},
            files={"file": ("股权转让协议.docx", payload, "application/octet-stream")},
        )

        assert business.status_code == 200
        assert executive.status_code == 200
        assert business.json()["document_id"] != executive.json()["document_id"]
        business_document = db.get(Document, business.json()["document_id"])
        executive_document = db.get(Document, executive.json()["document_id"])
        assert business_document.content_hash == executive_document.content_hash
        business_source = db.get(ContractDocumentSource, business_document.id)
        executive_source = db.get(ContractDocumentSource, executive_document.id)
        assert business_source is not None
        assert executive_source is not None
        business_path = Path(business_source.source_path)
        assert business_path.is_file()
        assert Path(executive_source.source_path).is_file()
        # Simulate a production row created before per-document source paths
        # existed.  The correct L5 copy remains on NAS and must be recovered.
        db.delete(executive_source)
        db.commit()
        db.expire(executive_document, ["contract_source"])
        assert db.get(ContractDocumentSource, executive_document.id) is None

        moved = client.patch(
            f"/v1/contracts/{executive_document.id}/folder",
            json={"folder_path": "内部资料（密）"},
        )

        assert moved.status_code == 200
        assert moved.json()["folder_path"] == "内部资料（密）"
        db.expire_all()
        executive_source = db.get(ContractDocumentSource, executive_document.id)
        assert "内部资料（密）" in Path(executive_source.source_path).as_posix()
        assert Path(executive_source.source_path).is_file()
        assert business_path.is_file()
        assert Path(business_source.source_path).resolve() == business_path.resolve()
        mine = client.get("/v1/contracts/mine")
        executive_item = next(
            item
            for item in mine.json()["items"]
            if item["document_id"] == executive_document.id
        )
        assert executive_item["folder_path"] == "内部资料（密）"
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_contracts_are_excluded_from_general_search_but_available_locally(
    tmp_path,
    monkeypatch,
) -> None:
    db, users = _database()
    knowledge_root = tmp_path / "knowledge"
    _configure(monkeypatch, db, users["l5_admin"], knowledge_root)
    client = TestClient(app)
    try:
        upload = client.post(
            "/v1/contracts/uploads",
            data={"category": "executive_office"},
            files={
                "file": (
                    "董事会保密合同.docx",
                    _word_payload("董事会特别决议与保密安排。"),
                    "application/octet-stream",
                )
            },
        )
        assert upload.status_code == 200
        document = db.get(Document, upload.json()["document_id"])
        document.knowledge_status = "approved"
        document.project.knowledge_status = "approved"
        db.commit()

        general = search(
            db,
            user=users["l5"],
            query="董事会特别决议",
            requested_scope="history",
            requested_retrieval="exact",
            generate=False,
        )
        assert general["results"] == []

        app.dependency_overrides[current_user] = lambda: users["l5"]
        dedicated = client.post(
            "/v1/contracts/search",
            json={
                "category": "executive_office",
                "query": "董事会特别决议",
                "limit": 5,
            },
        )
        assert dedicated.status_code == 200
        assert dedicated.json()["local_only"] is True
        assert len(dedicated.json()["results"]) == 1
        assert dedicated.json()["results"][0]["document_id"] == document.id
        listed = client.get(
            "/v1/contracts", params={"category": "executive_office"}
        )
        assert listed.status_code == 200
        assert [item["document_id"] for item in listed.json()["items"]] == [
            document.id
        ]
        audit = db.scalar(
            select(AuditLog).where(AuditLog.action == "contract_search")
        )
        assert audit is not None
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_executive_office_contracts_are_isolated_by_uploader(
    tmp_path,
    monkeypatch,
) -> None:
    db, users = _database()
    knowledge_root = tmp_path / "knowledge"
    _configure(monkeypatch, db, users["l4"], knowledge_root)
    client = TestClient(app)
    try:
        own_upload = client.post(
            "/v1/contracts/uploads",
            data={"category": "executive_office"},
            files={
                "file": (
                    "行政涉密.docx",
                    _word_payload("行政专属暗号甲号。"),
                    "application/octet-stream",
                )
            },
        )
        assert own_upload.status_code == 200

        app.dependency_overrides[current_user] = lambda: users["l5_finance"]
        other_upload = client.post(
            "/v1/contracts/uploads",
            data={"category": "executive_office"},
            files={
                "file": (
                    "财务涉密.docx",
                    _word_payload("财务专属暗号乙号。"),
                    "application/octet-stream",
                )
            },
        )
        assert other_upload.status_code == 200

        documents = [
            db.get(Document, own_upload.json()["document_id"]),
            db.get(Document, other_upload.json()["document_id"]),
        ]
        for document in documents:
            document.knowledge_status = "approved"
            document.project.knowledge_status = "approved"
        db.commit()

        app.dependency_overrides[current_user] = lambda: users["l4"]
        admin_list = client.get(
            "/v1/contracts", params={"category": "executive_office"}
        )
        assert [item["document_id"] for item in admin_list.json()["items"]] == [
            own_upload.json()["document_id"]
        ]
        assert client.get(
            f"/v1/documents/{other_upload.json()['document_id']}"
        ).status_code == 404
        hidden_search = client.post(
            "/v1/contracts/search",
            json={"category": "executive_office", "query": "财务专属暗号乙号"},
        )
        assert hidden_search.status_code == 200
        assert other_upload.json()["document_id"] not in {
            item["document_id"] for item in hidden_search.json()["results"]
        }

        own_search = client.post(
            "/v1/contracts/search",
            json={"category": "executive_office", "query": "行政专属暗号甲号"},
        )
        assert [item["document_id"] for item in own_search.json()["results"]] == [
            own_upload.json()["document_id"]
        ]

        app.dependency_overrides[current_user] = lambda: users["l5"]
        management_list = client.get(
            "/v1/contracts", params={"category": "executive_office"}
        )
        assert {item["document_id"] for item in management_list.json()["items"]} == {
            own_upload.json()["document_id"],
            other_upload.json()["document_id"],
        }
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_business_and_finance_cannot_bypass_contract_policy_by_document_id(
    tmp_path,
    monkeypatch,
) -> None:
    db, users = _database()
    _configure(monkeypatch, db, users["l5_admin"], tmp_path / "knowledge")
    client = TestClient(app)
    try:
        upload = client.post(
            "/v1/contracts/uploads",
            data={"category": "business"},
            files={
                "file": (
                    "业务保密合同.docx",
                    _word_payload("业务合同直接访问隔离测试。"),
                    "application/octet-stream",
                )
            },
        )
        assert upload.status_code == 200
        document = db.get(Document, upload.json()["document_id"])
        document.knowledge_status = "approved"
        document.project.knowledge_status = "approved"
        db.commit()

        app.dependency_overrides[current_user] = lambda: users["l5"]
        assert client.get(f"/v1/documents/{document.id}").status_code == 200

        for actor in (
            users["l4_personnel"],
            users["l5_business"],
            users["l5_finance"],
        ):
            app.dependency_overrides[current_user] = lambda actor=actor: actor
            assert client.get(f"/v1/documents/{document.id}").status_code == 404
            assert client.post(
                f"/v1/documents/{document.id}/preview-ticket?page=1"
            ).status_code == 404
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_scanner_enforces_contract_path_level_and_marks_invalid_layout(
    tmp_path,
    monkeypatch,
) -> None:
    db, users = _database()
    knowledge_root = tmp_path / "knowledge"
    _configure(monkeypatch, db, users["l5"], knowledge_root)
    valid = knowledge_root / "合同档案库" / "人事合同" / "L4"
    invalid = knowledge_root / "合同档案库" / "总办合同" / "L4"
    valid.mkdir(parents=True)
    invalid.mkdir(parents=True)
    (valid / "员工保密协议.docx").write_bytes(_word_payload("员工保密义务。"))
    (invalid / "错误密级.docx").write_bytes(_word_payload("总办事项。"))

    try:
        result = ingest.scan_inbox(db, knowledge_root, settle_seconds=0)
        document = db.scalar(select(Document))

        assert result["counts"]["ingested"] == 1
        assert result["counts"]["failed"] == 1
        assert document.role == "contract"
        assert document.confidentiality == "L4"
        assert document.project.domain == "contract_personnel"
        assert document.knowledge_status == "candidate"
        assert any(
            item.get("error_code") == "contract_layout_invalid"
            for item in result["items"]
        )
    finally:
        app.dependency_overrides.clear()
        db.close()
