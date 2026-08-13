from __future__ import annotations

import json
import os
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import pymupdf
from openpyxl import Workbook
from PIL import Image
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app import ingest
from app.models import (
    AuditLog,
    Base,
    Document,
    FileBlob,
    InboxIssue,
    KnowledgeCategory,
    Project,
)
from app.parsers import clean_text, parse_file


def scanner_db() -> Session:
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    return Session(engine)


def configure_scanner(monkeypatch, inbox) -> None:
    monkeypatch.setattr(
        ingest,
        "settings",
        SimpleNamespace(
            copy_sources=False,
            inbox_dir=inbox,
            inbox_settle_seconds=120,
            inbox_max_file_bytes=1024 * 1024,
        ),
    )


def make_old(path, seconds: int = 300) -> None:
    timestamp = (datetime.now(timezone.utc) - timedelta(seconds=seconds)).timestamp()
    os.utime(path, (timestamp, timestamp))


def test_scan_ingests_candidate_and_deduplicates_by_content(
    tmp_path,
    monkeypatch,
) -> None:
    inbox = tmp_path / "99_AI入库待审核"
    project_dir = inbox / "电竞培训" / "2026" / "高校电竞课程" / "02_提案"
    project_dir.mkdir(parents=True)
    source = project_dir / "2026_高校电竞课程_v1.2.md"
    source.write_text("# 课程方案\n课程目标与训练安排", encoding="utf-8")
    make_old(source)
    configure_scanner(monkeypatch, inbox)
    db = scanner_db()
    try:
        first = ingest.scan_inbox(db)
        duplicate = project_dir / "方案副本.md"
        duplicate.write_bytes(source.read_bytes())
        make_old(duplicate)
        second = ingest.scan_inbox(db)
        document = db.scalar(select(Document))
        scan_audit = db.scalars(
            select(AuditLog)
            .where(AuditLog.action == "inbox_scan")
                .order_by(AuditLog.created_at.desc(), AuditLog.id.desc())
        ).first()

        assert first["status"] == "ok"
        assert first["counts"]["ingested"] == 1
        assert second["counts"]["unchanged"] == 2
        assert db.scalar(select(func.count(Document.id))) == 1
        assert document is not None
        assert document.knowledge_status == "candidate"
        assert document.project.name == "高校电竞课程"
        assert document.project.domain == "training"
        assert document.project.confirmed is False
        assert document.version == "v1.2"
        assert document.is_final is False
        assert scan_audit is not None
        assert "高校电竞课程" not in scan_audit.details_json
        assert json.loads(scan_audit.details_json)["counts"]["unchanged"] == 2
    finally:
        db.close()


def test_scan_skips_system_managed_finance_records(tmp_path, monkeypatch) -> None:
    library = tmp_path / "knowledge"
    finance_file = library / "财务系统" / "银行流水" / "京奥电竞" / "流水.xlsx"
    finance_file.parent.mkdir(parents=True)
    finance_file.write_bytes(b"system-managed-finance-record")
    make_old(finance_file)
    configure_scanner(monkeypatch, library)
    db = scanner_db()
    try:
        result = ingest.scan_inbox(db, library, settle_seconds=0)

        assert result["status"] == "ok"
        assert result["counts"]["discovered"] == 0
        assert result["counts"]["skipped"] == 1
        assert db.scalar(select(func.count(Document.id))) == 0
        assert db.scalar(select(func.count(InboxIssue.id))) == 0
    finally:
        db.close()


def test_exact_copy_is_filtered_across_different_projects(
    tmp_path,
    monkeypatch,
) -> None:
    root = tmp_path / "library"
    root.mkdir()
    first = root / "项目甲.md"
    second = root / "项目乙.md"
    first.write_text("# 完全相同的方案\n相同内容", encoding="utf-8")
    second.write_bytes(first.read_bytes())
    configure_scanner(monkeypatch, root)
    db = scanner_db()
    try:
        first_result = ingest.ingest_document(
            db,
            {
                "path": str(first),
                "project_name": "项目甲",
                "domain": "training",
                "role": "proposal",
                "confidentiality": "L2",
            },
        )
        second_result = ingest.ingest_document(
            db,
            {
                "path": str(second),
                "project_name": "项目乙",
                "domain": "team",
                "role": "closing_report",
                "confidentiality": "L2",
            },
        )

        assert first_result["status"] == "ingested"
        assert second_result["status"] == "unchanged"
        assert second_result["duplicate_filtered"] is True
        assert second_result["document_id"] == first_result["document_id"]
        assert db.scalar(select(func.count(Document.id))) == 1
    finally:
        db.close()


def test_existing_active_duplicates_are_hidden_but_sources_are_kept(
    tmp_path,
) -> None:
    source = tmp_path / "报价.xlsx"
    source.write_bytes(b"same-workbook-bytes")
    content_hash = ingest.file_sha256(source)
    db = scanner_db()
    first_project = Project(name="项目甲", domain="training", confidentiality="L3")
    second_project = Project(name="项目乙", domain="training", confidentiality="L3")
    third_project = Project(name="项目丙", domain="training", confidentiality="L3")
    blob = FileBlob(
        content_hash=content_hash,
        size_bytes=source.stat().st_size,
        source_path=str(source),
    )
    first = Document(
        project=first_project,
        file_blob=blob,
        content_hash=content_hash,
        title="报价.xlsx",
        role="proposal",
        confidentiality="L3",
        knowledge_status="current",
    )
    second = Document(
        project=second_project,
        file_blob=blob,
        content_hash=content_hash,
        title="报价副本.xlsx",
        role="closing_report",
        confidentiality="L3",
        knowledge_status="candidate",
    )
    third = Document(
        project=third_project,
        file_blob=blob,
        content_hash=content_hash,
        title="已入库报价副本.xlsx",
        role="proposal",
        confidentiality="L3",
        knowledge_status="approved",
    )
    db.add_all([
        first_project,
        second_project,
        third_project,
        blob,
        first,
        second,
        third,
    ])
    db.commit()
    try:
        filtered = ingest.filter_exact_duplicate_documents(db)
        statuses = db.scalars(
            select(Document.knowledge_status).order_by(Document.knowledge_status)
        ).all()
        audit = db.scalar(
            select(AuditLog).where(
                AuditLog.action == "exact_duplicates_filtered"
            )
        )

        assert filtered == 2
        assert statuses == ["current", "duplicate", "duplicate"]
        assert source.is_file()
        assert audit is not None
        assert json.loads(audit.details_json)["source_files_deleted"] is False
    finally:
        db.close()


def test_library_scan_registers_files_copied_directly_to_category_folder(
    tmp_path,
    monkeypatch,
) -> None:
    library = tmp_path / "01_知识资料"
    source_dir = library / "校园合作" / "L2"
    source_dir.mkdir(parents=True)
    office_file = source_dir / "高校电竞课程方案.md"
    office_file.write_text("# 课程方案\n课程目标与教学安排", encoding="utf-8")
    asset_file = source_dir / "活动照片.jpg"
    Image.new("RGB", (32, 32), "white").save(asset_file)
    make_old(office_file)
    make_old(asset_file)
    configure_scanner(monkeypatch, library)
    db = scanner_db()
    db.add(KnowledgeCategory(key="campus", name="校园合作", active=True))
    db.commit()
    try:
        result = ingest.scan_inbox(db, library)
        documents = db.scalars(select(Document).order_by(Document.title)).all()

        assert result["status"] == "ok"
        assert result["counts"]["ingested"] == 2
        assert len(documents) == 2
        assert all(item.project.domain == "campus" for item in documents)
        assert all(item.confidentiality == "L2" for item in documents)
        assert next(
            item for item in documents if item.title.endswith(".md")
        ).knowledge_status == "candidate"
        assert next(
            item for item in documents if item.title.endswith(".jpg")
        ).knowledge_status == "approved"
    finally:
        db.close()


def test_scan_defers_recent_file_and_tracks_resolvable_issues(
    tmp_path,
    monkeypatch,
) -> None:
    inbox = tmp_path / "99_AI入库待审核"
    inbox.mkdir()
    recent = inbox / "仍在上传.md"
    recent.write_text("尚未稳定", encoding="utf-8")
    unsupported = inbox / "恶意程序.exe"
    unsupported.write_bytes(b"synthetic")
    configure_scanner(monkeypatch, inbox)
    db = scanner_db()
    try:
        first = ingest.scan_inbox(db)
        issues = db.scalars(
            select(InboxIssue).where(InboxIssue.resolved_at.is_(None))
        ).all()

        assert first["status"] == "warning"
        assert first["counts"]["deferred"] == 1
        assert first["counts"]["unsupported"] == 1
        assert {issue.error_code for issue in issues} == {
            "still_uploading",
            "unsupported_type",
        }
        assert db.scalar(select(func.count(Document.id))) == 0

        make_old(recent)
        unsupported.unlink()
        second = ingest.scan_inbox(db)
        unresolved = db.scalars(
            select(InboxIssue).where(InboxIssue.resolved_at.is_(None))
        ).all()
        resolved = db.scalars(
            select(InboxIssue).where(InboxIssue.resolved_at.is_not(None))
        ).all()

        assert second["status"] == "ok"
        assert second["counts"]["ingested"] == 1
        assert unresolved == []
        assert len(resolved) == 2
    finally:
        db.close()


def test_scan_blocks_oversized_file_without_parsing(
    tmp_path,
    monkeypatch,
) -> None:
    inbox = tmp_path / "99_AI入库待审核"
    inbox.mkdir()
    oversized = inbox / "超限资料.md"
    oversized.write_text("1234567890", encoding="utf-8")
    make_old(oversized)
    configure_scanner(monkeypatch, inbox)
    db = scanner_db()
    try:
        result = ingest.scan_inbox(db, max_file_bytes=5)
        issue = db.scalar(select(InboxIssue))

        assert result["status"] == "warning"
        assert result["counts"]["failed"] == 1
        assert issue is not None
        assert issue.error_code == "file_too_large"
        assert db.scalar(select(func.count(Document.id))) == 0
    finally:
        db.close()


def test_scan_skips_fnos_temporary_sidecar_files(tmp_path, monkeypatch) -> None:
    inbox = tmp_path / "99_AI入库待审核"
    inbox.mkdir()
    sidecar = inbox / "活动方案.pptx.~#0"
    sidecar.write_bytes(b"incomplete-sidecar")
    make_old(sidecar)
    configure_scanner(monkeypatch, inbox)
    db = scanner_db()
    try:
        result = ingest.scan_inbox(db)

        assert result["status"] == "ok"
        assert result["counts"]["discovered"] == 0
        assert result["counts"]["skipped"] == 1
        assert db.scalar(select(func.count(InboxIssue.id))) == 0
    finally:
        db.close()


def test_blank_image_registers_as_metadata_asset(tmp_path) -> None:
    image_path = tmp_path / "活动照片.png"
    Image.new("RGB", (120, 60), "white").save(image_path)

    pages, citation_basis, page_count = parse_file(image_path)

    assert pages == []
    assert citation_basis == "asset-metadata"
    assert page_count == 1


def test_legacy_doc_is_converted_to_page_citable_pdf(
    tmp_path,
    monkeypatch,
) -> None:
    source = tmp_path / "旧版项目报告.doc"
    source.write_bytes(b"legacy-doc")

    def fake_convert(_source, output_dir, target_extension):
        assert target_extension == "pdf"
        output = output_dir / "旧版项目报告.pdf"
        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "JINGAO LEGACY DOC EVIDENCE")
        pdf.save(output)
        return output

    monkeypatch.setattr("app.parsers._convert_legacy_office", fake_convert)
    pages, citation_basis, page_count = parse_file(source)

    assert page_count == 1
    assert citation_basis == "reference-page"
    assert "JINGAO LEGACY DOC EVIDENCE" in pages[0][1]


def test_macro_enabled_doc_is_read_without_executing_macros(
    tmp_path,
    monkeypatch,
) -> None:
    source = tmp_path / "会议通知.docm"
    source.write_bytes(b"macro-enabled-doc")

    def fake_convert(_source, output_dir, target_extension):
        assert target_extension == "pdf"
        output = output_dir / "会议通知.pdf"
        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "MACRO CONTENT READ ONLY")
        pdf.save(output)
        return output

    monkeypatch.setattr("app.parsers._convert_legacy_office", fake_convert)
    pages, citation_basis, page_count = parse_file(source)

    assert page_count == 1
    assert citation_basis == "reference-page"
    assert "MACRO CONTENT READ ONLY" in pages[0][1]


def test_legacy_xls_is_converted_and_indexed_by_worksheet(
    tmp_path,
    monkeypatch,
) -> None:
    source = tmp_path / "旧版验收单.xls"
    source.write_bytes(b"legacy-xls")

    def fake_convert(_source, output_dir, target_extension):
        assert target_extension == "xlsx"
        output = output_dir / "旧版验收单.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "验收结果"
        worksheet.append(["项目", "状态"])
        worksheet.append(["京奥电竞", "通过"])
        workbook.save(output)
        return output

    monkeypatch.setattr("app.parsers._convert_legacy_office", fake_convert)
    pages, citation_basis, page_count = parse_file(source)

    assert page_count == 1
    assert citation_basis == "worksheet"
    assert pages[0][2] == "验收结果"
    assert "京奥电竞 | 通过" in pages[0][1]


def test_clean_text_replaces_invalid_unicode_surrogates() -> None:
    cleaned = clean_text("\ud83c标题\n正文")

    assert cleaned == "�标题\n正文"
    assert cleaned.encode("utf-8")


def test_clean_text_removes_database_invalid_control_characters() -> None:
    cleaned = clean_text("官方入口\x00与明星入口\x7f形成飞轮")

    assert cleaned == "官方入口 与明星入口 形成飞轮"
    assert "\x00" not in cleaned


def test_non_text_assets_register_as_metadata_assets(
    tmp_path,
    monkeypatch,
) -> None:
    inbox = tmp_path / "99_AI入库待审核"
    project_dir = inbox / "电竞战队" / "俱乐部LOGO"
    project_dir.mkdir(parents=True)
    sources = [
        project_dir / "JAG-可编辑LOGO.ai",
        project_dir / "提交logo.zip",
        project_dir / "观赛训导.mp3",
        project_dir / "风险预案.mm",
    ]
    for source in sources:
        source.write_bytes(f"synthetic-{source.suffix}".encode())
        make_old(source)
    configure_scanner(monkeypatch, inbox)
    db = scanner_db()
    try:
        result = ingest.scan_inbox(db)
        documents = db.scalars(select(Document).order_by(Document.title)).all()

        assert result["status"] == "ok"
        assert result["counts"]["ingested"] == 4
        assert result["counts"]["unsupported"] == 0
        assert len(documents) == 4
        assert all(document.role == "asset" for document in documents)
        assert all(document.citation_basis == "asset-metadata" for document in documents)
        assert all(document.knowledge_status == "approved" for document in documents)
        assert all(document.page_count == 1 for document in documents)
        assert all(document.chunks == [] for document in documents)
    finally:
        db.close()


def test_policy_backfill_removes_existing_assets_from_review_queue(
    tmp_path,
    monkeypatch,
) -> None:
    inbox = tmp_path / "99_AI入库待审核"
    inbox.mkdir()
    source = inbox / "历史活动视频.mp4"
    source.write_bytes(b"synthetic-video")
    make_old(source)
    configure_scanner(monkeypatch, inbox)
    db = scanner_db()
    try:
        ingest.scan_inbox(db)
        document = db.scalar(select(Document))
        document.knowledge_status = "candidate"
        db.commit()

        updated = ingest.auto_approve_pending_assets(db)
        db.refresh(document)
        audit = db.scalar(
            select(AuditLog).where(
                AuditLog.action == "asset_review_policy_backfill"
            )
        )

        assert updated == 1
        assert document.knowledge_status == "approved"
        assert audit is not None
    finally:
        db.close()
