from __future__ import annotations

import hashlib
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.auth import current_user
from app.database import get_db
from app.main import app
from app.models import AuditLog, Base, Chunk, Document, FileBlob, Project, User


def seeded_preview_db(source_path: str) -> tuple[Session, User, User, Document]:
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    db = Session(engine)
    founder = User(
        username="founder-preview",
        display_name="Founder",
        password_hash="x",
        role="founder",
        confidentiality_ceiling="L4",
    )
    employee = User(
        username="employee-preview",
        display_name="Employee",
        password_hash="x",
        role="employee",
        confidentiality_ceiling="L2",
    )
    project = Project(
        name="Preview test",
        confidentiality="L3",
        knowledge_status="candidate",
    )
    payload = Path(source_path).read_bytes()
    blob = FileBlob(
        content_hash=hashlib.sha256(payload).hexdigest().upper(),
        size_bytes=len(payload),
        source_path=source_path,
    )
    document = Document(
        project=project,
        file_blob=blob,
        content_hash=blob.content_hash,
        title="preview.pdf",
        confidentiality="L3",
        knowledge_status="candidate",
        page_count=2,
    )
    db.add_all([founder, employee, document])
    db.commit()
    return db, founder, employee, document


def test_authorized_preview_streams_pdf_and_writes_audit(tmp_path) -> None:
    source = tmp_path / "preview.pdf"
    source.write_bytes(b"%PDF-1.4\n%%EOF")
    db, founder, _employee, document = seeded_preview_db(str(source))

    def override_db():
        yield db

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[current_user] = lambda: founder
    try:
        response = TestClient(app).get(
            f"/v1/documents/{document.id}/preview?page=2"
        )
        audit = db.scalar(
            select(AuditLog)
            .where(AuditLog.action == "open_source")
            .order_by(AuditLog.created_at.desc())
        )
        assert response.status_code == 200
        assert response.headers["content-type"] == "application/pdf"
        assert response.content.startswith(b"%PDF")
        assert audit is not None
        assert document.id in audit.document_ids_json
        assert '"page": 2' in audit.details_json
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_preview_ticket_supports_native_pdf_range_requests(tmp_path) -> None:
    source = tmp_path / "preview.pdf"
    source.write_bytes(b"%PDF-1.4\n%%EOF")
    db, founder, _employee, document = seeded_preview_db(str(source))

    def override_db():
        yield db

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[current_user] = lambda: founder
    try:
        client = TestClient(app)
        issued = client.post(
            f"/v1/documents/{document.id}/preview-ticket?page=1"
        )
        assert issued.status_code == 200
        assert issued.json()["path"].startswith("/v1/previews/")

        preview = client.get(
            issued.json()["path"],
            headers={"Range": "bytes=0-7"},
        )
        assert preview.status_code == 206
        assert preview.headers["accept-ranges"] == "bytes"
        assert preview.content == b"%PDF-1.4"
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_preview_masks_unauthorized_document(tmp_path) -> None:
    source = tmp_path / "preview.pdf"
    source.write_bytes(b"%PDF-1.4\n%%EOF")
    db, _founder, employee, document = seeded_preview_db(str(source))

    def override_db():
        yield db

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[current_user] = lambda: employee
    try:
        response = TestClient(app).get(
            f"/v1/documents/{document.id}/preview?page=1"
        )
        assert response.status_code == 404
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_authorized_preview_renders_extracted_text(tmp_path) -> None:
    source = tmp_path / "preview.txt"
    source.write_text("原始文本", encoding="utf-8")
    db, founder, _employee, document = seeded_preview_db(str(source))
    document.title = "preview.txt"
    document.page_count = 1
    db.add(
        Chunk(
            document=document,
            page=1,
            chunk_index=0,
            text="可直接审核的提取正文。",
        )
    )
    db.commit()

    def override_db():
        yield db

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[current_user] = lambda: founder
    try:
        response = TestClient(app).get(
            f"/v1/documents/{document.id}/preview?page=1"
        )
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("text/plain")
        assert "可直接审核的提取正文" in response.text
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_authorized_preview_renders_legacy_xls_extracted_text(tmp_path) -> None:
    source = tmp_path / "项目验收.xls"
    source.write_bytes(b"synthetic-legacy-xls")
    db, founder, _employee, document = seeded_preview_db(str(source))
    document.title = "项目验收.xls"
    document.page_count = 1
    db.add(
        Chunk(
            document=document,
            page=1,
            chunk_index=0,
            text="项目验收状态：通过。",
        )
    )
    db.commit()

    def override_db():
        yield db

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[current_user] = lambda: founder
    try:
        response = TestClient(app).get(
            f"/v1/documents/{document.id}/preview?page=1"
        )
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("text/plain")
        assert "项目验收状态" in response.text
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_preview_ticket_renders_xlsx_as_safe_table_and_downloads_original(
    tmp_path,
) -> None:
    source = tmp_path / "项目费用.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "总览"
    first.append(["仅第一页", 100])
    second = workbook.create_sheet("人员费用")
    second.append(["岗位", "单价", "备注"])
    second.append(["导演", 3000, "<script>alert('x')</script>"])
    workbook.save(source)
    workbook.close()

    db, founder, _employee, document = seeded_preview_db(str(source))
    document.title = "项目费用.xlsx"
    document.page_count = 2
    db.commit()

    def override_db():
        yield db

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[current_user] = lambda: founder
    try:
        client = TestClient(app)
        issued = client.post(
            f"/v1/documents/{document.id}/preview-ticket?page=2"
        )
        assert issued.status_code == 200

        preview = client.get(issued.json()["path"])
        assert preview.status_code == 200
        assert preview.headers["content-type"].startswith("text/html")
        assert "人员费用" in preview.text
        assert "工作表 2/2" in preview.text
        assert "岗位" in preview.text
        assert "仅第一页" not in preview.text
        assert "&lt;script&gt;alert" in preview.text
        assert "<script>alert" not in preview.text
        assert "?download=true" in preview.text
        assert "style-src 'unsafe-inline'" in preview.headers["content-security-policy"]

        downloaded = client.get(f'{issued.json()["path"]}?download=true')
        assert downloaded.status_code == 200
        assert downloaded.content == source.read_bytes()
        assert "attachment" in downloaded.headers["content-disposition"]
        audit = db.scalar(
            select(AuditLog)
            .where(AuditLog.action == "download_source")
            .order_by(AuditLog.created_at.desc())
        )
        assert audit is not None
        assert document.id in audit.document_ids_json
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_authorized_preview_streams_image(tmp_path) -> None:
    source = tmp_path / "preview.png"
    source.write_bytes(b"\x89PNG\r\n\x1a\nsynthetic")
    db, founder, _employee, document = seeded_preview_db(str(source))
    document.title = "preview.png"
    document.page_count = 1
    db.commit()

    def override_db():
        yield db

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[current_user] = lambda: founder
    try:
        response = TestClient(app).get(
            f"/v1/documents/{document.id}/preview?page=1"
        )
        assert response.status_code == 200
        assert response.headers["content-type"] == "image/png"
        assert response.content.startswith(b"\x89PNG")
    finally:
        app.dependency_overrides.clear()
        db.close()


def test_preview_ticket_supports_video_range_playback(tmp_path) -> None:
    source = tmp_path / "preview.mp4"
    source.write_bytes(b"0123456789-video-payload")
    db, founder, _employee, document = seeded_preview_db(str(source))
    document.title = "preview.mp4"
    document.role = "asset"
    document.page_count = 1
    db.commit()

    def override_db():
        yield db

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[current_user] = lambda: founder
    try:
        client = TestClient(app)
        issued = client.post(
            f"/v1/documents/{document.id}/preview-ticket?page=1"
        )
        preview = client.get(
            issued.json()["path"],
            headers={"Range": "bytes=2-7"},
        )

        assert issued.status_code == 200
        assert preview.status_code == 206
        assert preview.headers["content-type"] == "video/mp4"
        assert preview.headers["accept-ranges"] == "bytes"
        assert preview.content == b"234567"
    finally:
        app.dependency_overrides.clear()
        db.close()
