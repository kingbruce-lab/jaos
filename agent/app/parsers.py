from __future__ import annotations

import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document as WordDocument
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation


SPACE_RE = re.compile(r"[ \t\r\f\v]+")
BLANK_RE = re.compile(r"\n{3,}")
INVALID_SURROGATE_RE = re.compile(r"[\ud800-\udfff]")
INVALID_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
SENTENCE_END_RE = re.compile(r"(?<=[。！？!?；;])")


def clean_text(text: str) -> str:
    # Some PDF text layers expose isolated UTF-16 surrogate code points.
    # They are not valid Unicode scalar values and database drivers cannot
    # encode them as UTF-8, so preserve their position with the replacement
    # character before any extracted text reaches persistence or the LLM.
    text = INVALID_SURROGATE_RE.sub("\ufffd", text)
    # PDF text layers can also contain NUL and other non-printing controls.
    # PostgreSQL rejects NUL in text fields, while the remaining controls make
    # search snippets unreadable, so normalize them before chunk persistence.
    text = INVALID_CONTROL_RE.sub(" ", text)
    lines = [SPACE_RE.sub(" ", line).strip() for line in text.splitlines()]
    return BLANK_RE.sub("\n\n", "\n".join(line for line in lines if line)).strip()


ParsedPage = tuple[int, str, str | None]
ParseResult = tuple[list[ParsedPage], str, int]


def parse_pdf(path: Path) -> ParseResult:
    reader = PdfReader(str(path))
    pages: list[ParsedPage] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = clean_text(page.extract_text() or "")
        if text:
            pages.append((page_number, text, None))
    return pages, "source-page", len(reader.pages)


def parse_pptx(path: Path) -> ParseResult:
    presentation = Presentation(str(path))
    pages: list[ParsedPage] = []
    for page_number, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                parts.append(shape.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
        try:
            notes = slide.notes_slide.notes_text_frame.text
            if notes:
                parts.append(f"演讲备注：{notes}")
        except (AttributeError, KeyError):
            pass
        text = clean_text("\n".join(parts))
        if text:
            first_part = next((clean_text(part) for part in parts if clean_text(part)), "")
            section = first_part.splitlines()[0][:120] if first_part else None
            pages.append((page_number, text, section))
    return pages, "slide", len(presentation.slides)


def parse_docx(path: Path) -> ParseResult:
    document = WordDocument(str(path))
    blocks: list[str] = []
    section: str | None = None
    pages: list[ParsedPage] = []
    page_number = 1
    for paragraph in document.paragraphs:
        text = clean_text(paragraph.text)
        if not text:
            continue
        if paragraph.style and paragraph.style.name.startswith("Heading"):
            if blocks:
                pages.append((page_number, clean_text("\n".join(blocks)), section))
                page_number += 1
                blocks = []
            section = text[:120]
        blocks.append(text)
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(clean_text(cell.text) for cell in row.cells))
    if blocks:
        pages.append((page_number, clean_text("\n".join(blocks)), section))
    return pages, "logical-section", len(pages)


def parse_xlsx(path: Path) -> ParseResult:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        pages: list[ParsedPage] = []
        worksheet_count = len(workbook.worksheets)
        for sheet_number, worksheet in enumerate(workbook.worksheets, start=1):
            rows: list[str] = []
            for row in worksheet.iter_rows(values_only=True):
                values = [
                    str(value).strip()
                    for value in row
                    if value not in (None, "")
                ]
                if values:
                    rows.append(" | ".join(values))
            text = clean_text("\n".join(rows))
            if text:
                pages.append((sheet_number, text, worksheet.title))
        return pages, "worksheet", worksheet_count
    finally:
        workbook.close()


def _convert_legacy_office(
    path: Path,
    output_dir: Path,
    target_extension: str,
) -> Path:
    soffice = (
        os.getenv("JINGAO_SOFFICE_PATH")
        or shutil.which("soffice")
        or shutil.which("libreoffice")
    )
    if not soffice:
        raise RuntimeError("legacy_office_converter_unavailable")
    profile = (output_dir / "profile").resolve()
    subprocess.run(
        [
            soffice,
            "--headless",
            f"-env:UserInstallation={profile.as_uri()}",
            "--convert-to",
            target_extension,
            "--outdir",
            str(output_dir),
            str(path),
        ],
        check=True,
        timeout=180,
        capture_output=True,
    )
    generated = output_dir / f"{path.stem}.{target_extension}"
    if not generated.is_file():
        raise RuntimeError("legacy_office_output_missing")
    return generated


def parse_legacy_doc(path: Path) -> ParseResult:
    with tempfile.TemporaryDirectory(prefix="jingao-legacy-doc-") as temporary:
        output = _convert_legacy_office(path, Path(temporary), "pdf")
        pages, _basis, page_count = parse_pdf(output)
    return pages, "reference-page", page_count


def parse_legacy_xls(path: Path) -> ParseResult:
    with tempfile.TemporaryDirectory(prefix="jingao-legacy-xls-") as temporary:
        output = _convert_legacy_office(path, Path(temporary), "xlsx")
        return parse_xlsx(output)


def parse_image_ocr(path: Path) -> ParseResult:
    from rapidocr import RapidOCR

    result = RapidOCR()(str(path))
    texts = list(getattr(result, "txts", None) or [])
    scores = [float(value) for value in (getattr(result, "scores", None) or [])]
    accepted = [
        text
        for text, score in zip(texts, scores, strict=False)
        if score >= 0.45 and clean_text(text)
    ]
    text = clean_text("\n".join(accepted))
    if not text:
        return [], "asset-metadata", 1
    return [(1, text, None)], "ocr-page", 1


def parse_image(_path: Path) -> ParseResult:
    # Image OCR is intentionally deferred to the derivative worker so a
    # folder containing many event photos can be registered quickly without
    # blocking the scanner and its issue reconciliation.
    return [], "asset-metadata", 1


def parse_file(path: Path) -> ParseResult:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix == ".pptx":
        return parse_pptx(path)
    if suffix == ".docx":
        return parse_docx(path)
    if suffix in {".doc", ".docm"}:
        return parse_legacy_doc(path)
    if suffix in {".xlsx", ".xlsm"}:
        return parse_xlsx(path)
    if suffix == ".xls":
        return parse_legacy_xls(path)
    if suffix in {".txt", ".md"}:
        return [(1, clean_text(path.read_text(encoding="utf-8")), None)], "source-page", 1
    if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        return parse_image(path)
    if suffix in {".mp4", ".webm", ".mov", ".m4v"}:
        return [], "asset-metadata", 1
    if suffix in {".ai", ".zip", ".mp3", ".m4a", ".wav", ".flac", ".mm"}:
        # Design sources, archives, audio and mind maps are catalogued for
        # discovery and download only. Never unpack or execute their contents
        # during ingestion. Audio transcription remains a later-phase feature.
        return [], "asset-metadata", 1
    raise ValueError(f"暂不支持的文件类型：{suffix}")


def _split_oversized_unit(value: str, target: int) -> list[str]:
    sentences = [
        item.strip()
        for item in SENTENCE_END_RE.split(value)
        if item.strip()
    ]
    if len(sentences) == 1 and len(sentences[0]) > target:
        return [
            sentences[0][start:start + target].strip()
            for start in range(0, len(sentences[0]), target)
            if sentences[0][start:start + target].strip()
        ]

    units: list[str] = []
    current: list[str] = []
    current_size = 0
    for sentence in sentences:
        if len(sentence) > target:
            if current:
                units.append("".join(current))
                current = []
                current_size = 0
            units.extend(
                sentence[start:start + target].strip()
                for start in range(0, len(sentence), target)
                if sentence[start:start + target].strip()
            )
            continue
        if current and current_size + len(sentence) > target:
            units.append("".join(current))
            current = []
            current_size = 0
        current.append(sentence)
        current_size += len(sentence)
    if current:
        units.append("".join(current))
    return units


def chunk_text(text: str, target: int = 900, overlap: int = 120) -> list[str]:
    """Split on document structure before falling back to sentence boundaries.

    PPT shapes, Word paragraphs and spreadsheet rows arrive as separate lines.
    Keeping those units intact improves retrieval while preserving page-level
    citations. Overlap carries complete trailing units instead of arbitrary
    character fragments.
    """
    normalized = clean_text(text)
    if not normalized:
        return []
    if len(normalized) <= target:
        return [normalized]

    units: list[str] = []
    for line in normalized.splitlines():
        value = line.strip()
        if not value:
            continue
        if len(value) <= target:
            units.append(value)
        else:
            units.extend(_split_oversized_unit(value, target))

    chunks: list[str] = []
    current: list[str] = []
    current_size = 0
    for unit in units:
        separator_size = 1 if current else 0
        if current and current_size + separator_size + len(unit) > target:
            chunks.append("\n".join(current))
            carry: list[str] = []
            carry_size = 0
            for previous in reversed(current):
                required = len(previous) + (1 if carry else 0)
                if carry and carry_size + required > overlap:
                    break
                if not carry and len(previous) > overlap:
                    break
                carry.insert(0, previous)
                carry_size += required
            current = carry
            current_size = len("\n".join(current))
        current.append(unit)
        current_size += len(unit) + (1 if len(current) > 1 else 0)

    if current:
        chunks.append("\n".join(current))
    return [item for item in chunks if item]
