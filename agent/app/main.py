from __future__ import annotations

import hashlib
import html
import json
import os
import re
import secrets
from contextlib import asynccontextmanager
from datetime import datetime, timezone
from pathlib import Path
from threading import Lock
from urllib.parse import quote

from fastapi import (
    Depends,
    FastAPI,
    File,
    Form,
    HTTPException,
    Query,
    Request,
    UploadFile,
    status,
)
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, Response
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import delete, func, or_, select
from sqlalchemy.exc import OperationalError, TimeoutError as DatabasePoolTimeout
from sqlalchemy.orm import Session, joinedload, selectinload

from .auth import (
    authenticate,
    bootstrap_admin,
    create_session,
    current_user,
    hash_password,
    login_rate_limiter,
    raw_bearer_token,
    revoke_session,
    revoke_user_sessions,
    verify_password,
)
from .config import settings
from .confidentiality import suggest_document_confidentiality
from .contracts import (
    CONTRACT_CATEGORIES,
    CONTRACT_CATEGORIES_BY_DOMAIN,
    CONTRACT_DEFAULT_FOLDERS,
    CONTRACT_DOMAINS,
    CONTRACT_UPLOAD_SUFFIXES,
    ContractCategory,
    ensure_contract_layout,
)
from .contract_filing import suggest_contract_folder
from .database import SessionLocal, get_db, init_database
from .evaluation import (
    approve_business_gold_cases,
    business_gold_registry,
    generate_business_gold_candidates,
    load_latest_evaluation,
    run_concurrency_evaluation,
    run_technical_evaluation,
    upsert_business_gold_drafts,
)
from .evolution import create_evolution_digest
from .generation import (
    GenerationEvidence,
    GenerationServiceError,
    evidence_allowed,
    generate_grounded_draft,
)
from .maintained_artifacts import (
    EvolutionWorkflowError,
    close_review,
    create_artifact,
    create_artifact_review,
    create_change_plan,
    decide_candidate,
    get_artifact,
    get_review_run,
    list_artifacts,
    list_baseline_documents,
    update_artifact_baseline,
)
from .models import (
    AuditLog,
    Chunk,
    ChunkEmbedding,
    ContractDocumentOwner,
    ContractDocumentSource,
    Document,
    InboxIssue,
    KnowledgeCategory,
    ManagedProject,
    Project,
    ReviewProposal,
    SourceHealth,
    User,
    WritingDraft,
)
from .operations import collect_operations_status
from .proposal_pptx import ProposalPptxError, render_proposal_pptx
from .retrieval import is_authorized, search
from .ingest import (
    SUPPORTED_SUFFIXES,
    auto_approve_pending_assets,
    filter_exact_duplicate_documents,
    ingest_document,
    prefill_entry,
    requires_ingestion_review,
    scan_inbox,
)
from .source_integrity import (
    file_sha256,
    reconcile_sources,
    source_is_available,
    verify_and_record,
)
from .storage_layout import (
    StorageLayoutError,
    ensure_category_layout,
    rename_category_layout,
    validate_category_folder_name,
)
from .schemas import (
    AdminUserCreate,
    AdminUserUpdate,
    BatchReviewConfirmRequest,
    BatchReviewRejectRequest,
    BusinessGoldApprovalRequest,
    BusinessGoldGenerateRequest,
    BusinessGoldImportRequest,
    EvolutionArtifactRunRequest,
    EvolutionCandidateDecisionRequest,
    EvolutionChangePlanRequest,
    EvolutionDigestRequest,
    EvolutionReviewCloseRequest,
    GovernanceAIClassifyRequest,
    GovernanceConfidentialityUpdateRequest,
    ContractSearchRequest,
    ContractFolderCreateRequest,
    ContractFolderMoveRequest,
    LoginRequest,
    MaintainedArtifactBaselineUpdate,
    MaintainedArtifactCreate,
    KnowledgeCategoryCreate,
    KnowledgeCategoryUpdate,
    OwnPasswordChange,
    PasswordReset,
    PublicationRequest,
    ProposalBrief,
    ReviewApplyRequest,
    ReviewIngestionRejectRequest,
    ReviewProposalRequest,
    ReviewRejectRequest,
    SearchRequest,
    WritingDraftRequest,
    WritingDraftSaveRequest,
)
from .taxonomy import seed_default_categories
from .finance_system import (
    normalize_business_entity_registry,
    router as finance_router,
)
from .project_system import router as project_management_router
from .education_system import router as education_router
from .education_enrollment import router as education_enrollment_router
from .education_assessments import router as education_assessment_router
from .education_ledger import router as education_ledger_router
from .education_operations import router as education_operations_router
from .writing_drafts import (
    MAX_SAVED_WRITING_DRAFTS,
    list_writing_drafts,
    save_writing_snapshot,
    serialize_writing_draft,
    upsert_latest_writing_draft,
)


@asynccontextmanager
async def lifespan(_app: FastAPI):
    init_database()
    with SessionLocal() as db:
        bootstrap_admin(db)
        normalize_business_entity_registry(db)
        seed_default_categories(db)
        filter_exact_duplicate_documents(db)
        auto_approve_pending_assets(db)
        ensure_contract_layout(settings.knowledge_root)
    yield


version_file = Path(__file__).resolve().parents[2] / "VERSION"
app_version = (
    version_file.read_text(encoding="utf-8").strip()
    if version_file.is_file()
    else "0.0.0"
)

app = FastAPI(
    title=settings.api_title,
    version=app_version,
    docs_url="/docs",
    redoc_url=None,
    lifespan=lifespan,
)
app.add_middleware(
    CORSMiddleware,
    allow_origins=list(settings.cors_origins),
    allow_credentials=False,
    allow_methods=["GET", "POST", "PATCH", "DELETE"],
    allow_headers=["Authorization", "Content-Type"],
)
app.include_router(finance_router)
app.include_router(project_management_router)
app.include_router(education_router)
app.include_router(education_enrollment_router)
app.include_router(education_assessment_router)
app.include_router(education_ledger_router)
app.include_router(education_operations_router)


@app.exception_handler(DatabasePoolTimeout)
@app.exception_handler(OperationalError)
async def database_unavailable_handler(_request: Request, _exc: Exception) -> JSONResponse:
    # Connection/lock failures are temporary service errors, not bad passwords.
    # Never include SQL, connection strings or bound parameters in the reply.
    return JSONResponse(
        status_code=503,
        content={"detail": "数据库暂时繁忙，请稍后重试；无需更改密码。"},
        headers={"Retry-After": "5"},
    )


@app.exception_handler(EvolutionWorkflowError)
async def evolution_workflow_error_handler(
    _request: Request,
    exc: EvolutionWorkflowError,
) -> JSONResponse:
    return JSONResponse(
        status_code=exc.status_code,
        content={"detail": exc.detail},
    )


def user_payload(user: User) -> dict:
    return {
        "id": user.id,
        "username": user.username,
        "display_name": user.display_name,
        "role": user.role,
        "organization_role": user.organization_role or "business",
        "confidentiality_ceiling": user.confidentiality_ceiling,
        "departments": json.loads(user.departments_json),
    }


def admin_user_payload(user: User) -> dict:
    return {
        **user_payload(user),
        "active": user.active,
        "external_identity": user.external_identity,
        "created_at": user.created_at,
    }


def require_founder(user: User) -> None:
    if user.role != "founder":
        raise HTTPException(status_code=403, detail="仅创始人可管理本地账号")


def _safe_upload_filename(value: str | None) -> str:
    raw = (value or "").replace("\\", "/")
    name = Path(raw).name.strip()
    name = re.sub(r"[\x00-\x1f<>:\"/\\|?*]+", "_", name)
    name = name.strip(" .")
    if not name:
        raise HTTPException(status_code=422, detail="文件名无效")
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise HTTPException(
            status_code=415,
            detail=f"暂不支持该文件类型：{suffix or '无扩展名'}",
        )
    if len(name) > 180:
        name = f"{Path(name).stem[:150]}{suffix}"
    return name


def _safe_upload_folder_name(value: str | None) -> str | None:
    raw = (value or "").strip()
    if not raw:
        return None
    try:
        name = validate_category_folder_name(raw)
    except StorageLayoutError as exc:
        raise HTTPException(status_code=422, detail="素材文件夹名称无效") from exc
    if len(name) > 120:
        raise HTTPException(status_code=422, detail="素材文件夹名称不能超过 120 个字符")
    return name


def _safe_upload_relative_parent(value: str | None) -> Path:
    raw = (value or "").replace("\\", "/").strip(" /")
    if not raw:
        return Path()
    parts = raw.split("/")
    # The final component is the uploaded file name and is handled separately.
    parents = parts[:-1]
    safe_parts: list[str] = []
    for part in parents:
        try:
            safe_parts.append(validate_category_folder_name(part))
        except StorageLayoutError as exc:
            raise HTTPException(status_code=422, detail="素材文件夹内部路径无效") from exc
    return Path(*safe_parts)


def _contract_document_folder_path(source_path: str, category_key: str) -> str:
    """Return the user-created folder path below a contract category root."""
    try:
        category_root = ensure_contract_layout(settings.knowledge_root)[category_key].resolve()
        source_parent = Path(source_path).resolve().parent
        relative = source_parent.relative_to(category_root)
    except (KeyError, OSError, ValueError):
        return ""
    return "" if relative == Path(".") else relative.as_posix()


def _path_inside_root(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
    except (OSError, ValueError):
        return False
    return True


def _contract_document_source_path(
    document: Document,
    category_key: str,
    *,
    recover_legacy: bool = False,
) -> Path | None:
    """Resolve the NAS copy that belongs to this contract document.

    Older rows only have the shared ``FileBlob.source_path``.  When identical
    bytes were uploaded into another contract security category, that shared
    path can point outside this document's category.  During an explicit move
    we recover the legacy copy by exact filename, size and SHA-256 within the
    correct category root.
    """
    try:
        root = ensure_contract_layout(settings.knowledge_root)[category_key].resolve()
    except (KeyError, OSError):
        return None
    contract_source = getattr(document, "contract_source", None)
    if contract_source and contract_source.source_path:
        assigned = Path(contract_source.source_path).resolve()
        if _path_inside_root(assigned, root):
            return assigned
    canonical_path = (
        Path(document.file_blob.source_path).resolve()
        if document.file_blob and document.file_blob.source_path
        else None
    )
    if canonical_path is not None and _path_inside_root(canonical_path, root):
        return canonical_path
    if not recover_legacy:
        return None

    filenames = {
        Path(document.title).name,
        canonical_path.name if canonical_path is not None else "",
    }
    expected_size = document.file_blob.size_bytes if document.file_blob else None
    for filename in sorted(item for item in filenames if item):
        for candidate in root.rglob(filename):
            try:
                relative = candidate.resolve().relative_to(root)
                if any(part.startswith(".") for part in relative.parts):
                    continue
                if not candidate.is_file():
                    continue
                if expected_size is not None and candidate.stat().st_size != expected_size:
                    continue
                if file_sha256(candidate) == document.content_hash:
                    return candidate.resolve()
            except OSError:
                continue
    return None


def _set_contract_document_source(
    db: Session,
    document: Document,
    source: Path,
) -> ContractDocumentSource:
    assigned = db.get(ContractDocumentSource, document.id)
    if assigned is None:
        assigned = ContractDocumentSource(
            document_id=document.id,
            source_path=str(source.resolve()),
        )
        db.add(assigned)
    else:
        assigned.source_path = str(source.resolve())
    document.contract_source = assigned
    return assigned


def _safe_contract_folder_path(value: str | None) -> Path:
    raw = (value or "").replace("\\", "/").strip(" /")
    if not raw:
        return Path()
    parts: list[str] = []
    for item in raw.split("/"):
        try:
            parts.append(validate_category_folder_name(item))
        except StorageLayoutError as exc:
            raise HTTPException(status_code=422, detail="合同文件夹名称无效") from exc
    if len(parts) > 6:
        raise HTTPException(status_code=422, detail="合同文件夹最多支持 6 层")
    return Path(*parts)


def _move_nas_file_without_overwrite(source: Path, target_dir: Path) -> Path:
    source = source.resolve()
    target_dir = target_dir.resolve()
    target_dir.mkdir(parents=True, exist_ok=True)
    if source.parent == target_dir:
        return source
    for attempt in range(100):
        candidate = (
            target_dir / source.name
            if attempt == 0
            else target_dir / f"{source.stem}_{secrets.token_hex(3)}{source.suffix}"
        )
        try:
            os.link(source, candidate)
            source.unlink()
            return candidate
        except FileExistsError:
            continue
    raise HTTPException(status_code=409, detail="目标文件夹中同名合同过多")


def _can_manage_contract_folders(user: User) -> bool:
    role = user.organization_role or "business"
    ceiling = CONFIDENTIALITY_RANK.get(user.confidentiality_ceiling, 0)
    return bool(
        user.active
        and (
            (role == "administrative" and ceiling >= CONFIDENTIALITY_RANK["L4"])
            or (role == "management" and ceiling >= CONFIDENTIALITY_RANK["L5"])
        )
    )


def _contract_status_label(value: str) -> str:
    return {
        "candidate": "待审核",
        "approved": "已审核",
        "current": "当前版本",
        "superseded": "历史版本",
        "archived": "已归档",
        "rejected": "已拒绝",
    }.get(value, value)


def _upload_department_allowed(user: User, department: str) -> bool:
    return bool(user.active and department)


def _can_upload_contract_category(user: User, category_key: str) -> bool:
    """Apply the narrow upload exception for highest-secret contracts."""
    if not user.active:
        return False
    organization_role = user.organization_role or "business"
    ceiling = CONFIDENTIALITY_RANK.get(user.confidentiality_ceiling, 0)
    if organization_role == "management":
        return ceiling >= CONFIDENTIALITY_RANK["L5"]
    if organization_role == "administrative":
        return category_key in CONTRACT_CATEGORIES
    return bool(
        organization_role in {"personnel", "finance"}
        and category_key == "executive_office"
        and ceiling >= CONFIDENTIALITY_RANK["L4"]
    )


def _contract_scope_for_user(user: User, category_key: str) -> str | None:
    organization_role = user.organization_role or "business"
    ceiling = CONFIDENTIALITY_RANK.get(user.confidentiality_ceiling, 0)
    if not user.active:
        return None
    if category_key == "executive_office":
        if organization_role == "management" and ceiling >= CONFIDENTIALITY_RANK["L5"]:
            return "all"
        if (
            organization_role in {"administrative", "personnel", "finance"}
            and ceiling >= CONFIDENTIALITY_RANK["L4"]
        ):
            return "own"
        return None
    allowed_categories = {
        "administrative": {"administrative", "business"},
        "personnel": {"personnel"},
        "business": set(),
        "finance": set(),
        "management": set(CONTRACT_CATEGORIES),
    }
    category = CONTRACT_CATEGORIES[category_key]
    if (
        category_key in allowed_categories.get(organization_role, set())
        and ceiling >= CONFIDENTIALITY_RANK[category.confidentiality]
    ):
        return "all"
    return None


def _owned_contract_document_ids(db: Session, user: User) -> set[str]:
    """Return explicit ownership plus legacy contract-upload audit records."""
    owned = set(
        db.scalars(
            select(ContractDocumentOwner.document_id).where(
                ContractDocumentOwner.uploaded_by_user_id == user.id
            )
        ).all()
    )
    legacy_logs = db.scalars(
        select(AuditLog).where(
            AuditLog.user_id == user.id,
            AuditLog.action == "contract_upload",
        )
    ).all()
    for log in legacy_logs:
        try:
            values = json.loads(log.document_ids_json or "[]")
        except (TypeError, json.JSONDecodeError):
            continue
        owned.update(
            str(value) for value in values if isinstance(value, str) and value
        )
    return owned


def _contract_category_for_user(
    user: User,
    category_key: str,
    *,
    operation: str = "search",
) -> ContractCategory:
    category = CONTRACT_CATEGORIES.get(category_key)
    if category is None:
        raise HTTPException(status_code=404, detail="合同分类不存在")
    if operation == "upload":
        allowed = _can_upload_contract_category(user, category_key)
    else:
        allowed = _contract_scope_for_user(user, category_key) is not None
    if not allowed:
        raise HTTPException(status_code=403, detail="无权访问该合同档案")
    return category


ORGANIZATION_ACCESS_ROLE = {
    "education": "employee",
    "administrative": "knowledge_admin",
    "personnel": "employee",
    "business": "employee",
    "finance": "employee",
    "management": "founder",
}


def _contract_policy_for_document(document: Document) -> ContractCategory | None:
    category = CONTRACT_CATEGORIES_BY_DOMAIN.get(document.project.domain)
    if document.role == "contract" and category is None:
        raise HTTPException(
            status_code=409,
            detail="合同档案分类异常，请先由系统管理员修复目录归属",
        )
    return category


def _commit_uploaded_file(
    temporary: Path,
    target_dir: Path,
    filename: str,
) -> Path:
    """Publish an upload without ever replacing an existing NAS file."""
    stem = Path(filename).stem
    suffix = Path(filename).suffix
    for attempt in range(100):
        candidate = (
            target_dir / filename
            if attempt == 0
            else target_dir / f"{stem}_{secrets.token_hex(3)}{suffix}"
        )
        try:
            # Hard-link creation is atomic and fails when the destination
            # already exists. Both paths are in the same NAS directory.
            os.link(temporary, candidate)
            temporary.unlink()
            return candidate
        except FileExistsError:
            continue
    raise HTTPException(status_code=409, detail="同名文件过多，请修改文件名后重试")


@app.get("/healthz")
def healthz() -> dict:
    return {"status": "ok"}


@app.post("/v1/uploads")
async def upload_to_inbox(
    file: UploadFile = File(...),
    department: str = Form(default="training"),
    confidentiality: str = Form(default="L2"),
    folder_name: str = Form(default=""),
    relative_path: str = Form(default=""),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    """Stream one employee file into the NAS intake area and register it."""
    category = db.scalar(
        select(KnowledgeCategory).where(
            KnowledgeCategory.key == department,
            KnowledgeCategory.active.is_(True),
        )
    )
    if category is None:
        raise HTTPException(status_code=422, detail="资料分类不存在或已停用")
    if not _upload_department_allowed(user, department):
        raise HTTPException(status_code=403, detail="无权向该部门上传资料")
    requested_rank = CONFIDENTIALITY_RANK.get(confidentiality, 99)
    ceiling_rank = CONFIDENTIALITY_RANK.get(user.confidentiality_ceiling, 0)
    if requested_rank > ceiling_rank:
        raise HTTPException(status_code=403, detail="上传密级超过账号权限")

    filename = _safe_upload_filename(file.filename)
    safe_folder_name = _safe_upload_folder_name(folder_name)
    relative_parent = _safe_upload_relative_parent(relative_path)
    knowledge_root = settings.knowledge_root.resolve()
    if not knowledge_root.is_dir():
        raise HTTPException(status_code=503, detail="知识资料目录暂不可用")
    try:
        target_dir = (
            ensure_category_layout(knowledge_root, category.name)
            / confidentiality
        )
        if safe_folder_name:
            target_dir = target_dir / safe_folder_name / relative_parent
            target_dir.mkdir(parents=True, exist_ok=True)
    except (OSError, StorageLayoutError) as exc:
        raise HTTPException(status_code=503, detail="资料分类目录暂不可写入") from exc
    target_dir.resolve().relative_to(knowledge_root)
    temporary = target_dir / f".{secrets.token_hex(12)}.part"
    maximum = settings.inbox_max_file_bytes
    size = 0
    digest = hashlib.sha256()
    try:
        with temporary.open("wb") as destination:
            while True:
                chunk = await file.read(1024 * 1024)
                if not chunk:
                    break
                size += len(chunk)
                if size > maximum:
                    raise HTTPException(
                        status_code=413,
                        detail="文件超过当前单文件大小上限",
                    )
                digest.update(chunk)
                destination.write(chunk)
        if size == 0:
            raise HTTPException(status_code=422, detail="不能上传空文件")

        final_path = _commit_uploaded_file(temporary, target_dir, filename)

        upload_status = "uploaded_pending"
        document_id = None
        duplicate_filtered = False
        review_required = requires_ingestion_review(final_path)
        try:
            stat = final_path.stat()
            entry = prefill_entry(final_path, knowledge_root)
            entry["domain"] = department
            entry["confidentiality"] = confidentiality
            if safe_folder_name:
                entry["title"] = f"{safe_folder_name} / {filename}"
                entry["project_name"] = safe_folder_name
            result = ingest_document(
                db,
                entry,
                expected_stat=(stat.st_size, stat.st_mtime_ns),
            )
            upload_status = result["status"]
            document_id = result["document_id"]
            duplicate_filtered = result.get("duplicate_filtered", False)
            review_required = result.get("review_required", review_required)
        except Exception as exc:
            db.rollback()
            db.add(
                AuditLog(
                    user_id=user.id,
                    action="web_upload_parse_deferred",
                    details_json=json.dumps(
                        {
                            "department": department,
                            "confidentiality": confidentiality,
                            "suffix": final_path.suffix.lower(),
                            "size_bytes": size,
                            "error_code": type(exc).__name__[:80],
                        },
                        ensure_ascii=False,
                    ),
                )
            )
            db.commit()

        db.add(
            AuditLog(
                user_id=user.id,
                action="web_upload",
                document_ids_json=json.dumps(
                    [document_id] if document_id else []
                ),
                details_json=json.dumps(
                    {
                        "department": department,
                        "confidentiality": confidentiality,
                        "suffix": final_path.suffix.lower(),
                        "size_bytes": size,
                        "content_hash": digest.hexdigest().upper(),
                        "status": upload_status,
                        "folder_name": safe_folder_name,
                        "relative_path": relative_path if safe_folder_name else "",
                    },
                    ensure_ascii=False,
                ),
            )
        )
        db.commit()
        return {
            "filename": filename,
            "folder_name": safe_folder_name,
            "department": department,
            "confidentiality": confidentiality,
            "size_bytes": size,
            "status": upload_status,
            "document_id": document_id,
            "review_required": review_required,
            "duplicate_filtered": duplicate_filtered,
            "notice": (
                "检测到内容完全重复，已自动过滤，不重复入库。"
                if duplicate_filtered
                else "上传成功，资料已进入审核队列。"
                if document_id and review_required
                else "上传成功，素材已自动入库，无需审核。"
                if document_id
                else "上传成功，系统将在后台继续解析。"
            ),
        }
    finally:
        temporary.unlink(missing_ok=True)
        await file.close()


def _public_health(payload: dict) -> dict:
    return {
        "status": payload["status"],
        "components": {
            name: {"status": item["status"]}
            for name, item in payload["components"].items()
            if name in {"database", "storage", "memory", "backup"}
        },
        "generated_at": payload["generated_at"],
    }


@app.get("/v1/contracts/categories")
def list_contract_categories(
    user: User = Depends(current_user),
) -> list[dict]:
    ceiling = CONFIDENTIALITY_RANK.get(user.confidentiality_ceiling, 0)
    searchable = {
        key
        for key in CONTRACT_CATEGORIES
        if _contract_scope_for_user(user, key) is not None
    }
    uploadable = {
        key
        for key in CONTRACT_CATEGORIES
        if _can_upload_contract_category(user, key)
    }
    if (
        ceiling < CONFIDENTIALITY_RANK["L4"]
        or not (searchable or uploadable)
    ):
        raise HTTPException(status_code=403, detail="无权访问合同档案库")
    return [
        {
            "key": item.key,
            "name": item.name,
            "confidentiality": item.confidentiality,
            "can_search": item.key in searchable,
            "can_upload": item.key in uploadable,
            "search_scope": _contract_scope_for_user(user, item.key),
        }
        for item in CONTRACT_CATEGORIES.values()
        if item.key in searchable | uploadable
    ]


def _contract_document_item(
    document: Document,
    *,
    category: ContractCategory,
    can_move: bool,
) -> dict:
    source = _contract_document_source_path(document, category.key)
    source_path = str(source) if source is not None else ""
    return {
        "document_id": document.id,
        "title": document.title,
        "category": category.key,
        "category_name": category.name,
        "confidentiality": document.confidentiality,
        "knowledge_status": document.knowledge_status,
        "status_label": _contract_status_label(document.knowledge_status),
        "folder_path": (
            _contract_document_folder_path(source_path, category.key)
            if source_path
            else ""
        ),
        "page_count": document.page_count,
        "created_at": document.ingested_at,
        "source_available": bool(source_path and Path(source_path).is_file()),
        "can_move": can_move,
    }


@app.get("/v1/contracts/mine")
def list_my_contract_documents(
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    """List every non-deleted contract explicitly uploaded by this user."""
    uploadable = {
        key for key in CONTRACT_CATEGORIES if _can_upload_contract_category(user, key)
    }
    if not uploadable:
        raise HTTPException(status_code=403, detail="当前账号没有合同上传权限")
    owned_ids = _owned_contract_document_ids(db, user)
    if not owned_ids:
        return {"count": 0, "items": []}
    documents = db.scalars(
        select(Document)
        .join(Document.project)
        .options(
            joinedload(Document.project),
            joinedload(Document.file_blob),
            joinedload(Document.contract_source),
        )
        .where(
            Document.id.in_(owned_ids),
            Project.domain.in_(CONTRACT_DOMAINS),
            Document.knowledge_status.notin_(("deleted", "quarantined")),
        )
        .order_by(Document.ingested_at.desc())
        .limit(500)
    ).all()
    items: list[dict] = []
    for document in documents:
        category = CONTRACT_CATEGORIES_BY_DOMAIN.get(document.project.domain)
        if category is None:
            continue
        items.append(
            _contract_document_item(
                document,
                category=category,
                can_move=_can_manage_contract_folders(user),
            )
        )
    db.add(
        AuditLog(
            user_id=user.id,
            action="contract_own_archive_list",
            document_ids_json=json.dumps([item["document_id"] for item in items]),
            details_json=json.dumps({"result_count": len(items)}),
        )
    )
    db.commit()
    return {"count": len(items), "items": items}


@app.get("/v1/contracts/folders")
def list_contract_folders(
    category: str = Query(..., min_length=1, max_length=40),
    user: User = Depends(current_user),
) -> dict:
    if not _can_upload_contract_category(user, category):
        raise HTTPException(status_code=403, detail="无权查看该合同分类的文件夹")
    root = ensure_contract_layout(settings.knowledge_root)[category].resolve()
    folders = sorted(
        set(CONTRACT_DEFAULT_FOLDERS.get(category, ()))
        | {
            path.relative_to(root).as_posix()
            for path in root.rglob("*")
            if path.is_dir()
            and not any(part.startswith(".") for part in path.relative_to(root).parts)
        }
    )[:500]
    return {"category": category, "items": folders}


@app.post("/v1/contracts/folders")
def create_contract_folder(
    payload: ContractFolderCreateRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    if not _can_manage_contract_folders(user):
        raise HTTPException(status_code=403, detail="仅行政或最高管理账号可整理合同文件夹")
    if not _can_upload_contract_category(user, payload.category):
        raise HTTPException(status_code=403, detail="无权整理该合同分类")
    relative = _safe_contract_folder_path(payload.folder_path)
    root = ensure_contract_layout(settings.knowledge_root)[payload.category].resolve()
    target = (root / relative).resolve()
    try:
        target.relative_to(root)
        target.mkdir(parents=True, exist_ok=True)
    except (OSError, ValueError) as exc:
        raise HTTPException(status_code=503, detail="合同文件夹暂时无法创建") from exc
    db.add(
        AuditLog(
            user_id=user.id,
            action="contract_folder_create",
            details_json=json.dumps(
                {"category": payload.category, "folder_path": relative.as_posix()},
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    return {"category": payload.category, "folder_path": relative.as_posix()}


@app.patch("/v1/contracts/{document_id}/folder")
def move_contract_document(
    document_id: str,
    payload: ContractFolderMoveRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    if not _can_manage_contract_folders(user):
        raise HTTPException(status_code=403, detail="仅行政或最高管理账号可移动合同")
    document = db.scalar(
        select(Document)
        .options(
            joinedload(Document.project),
            joinedload(Document.file_blob),
            joinedload(Document.contract_source),
        )
        .where(Document.id == document_id)
    )
    if not document or not document.project or not document.file_blob:
        raise HTTPException(status_code=404, detail="合同不存在")
    category = CONTRACT_CATEGORIES_BY_DOMAIN.get(document.project.domain)
    if category is None:
        raise HTTPException(status_code=404, detail="合同不存在")
    if (user.organization_role or "business") != "management" and document.id not in _owned_contract_document_ids(db, user):
        raise HTTPException(status_code=404, detail="合同不存在")
    relative = _safe_contract_folder_path(payload.folder_path)
    root = ensure_contract_layout(settings.knowledge_root)[category.key].resolve()
    source = _contract_document_source_path(
        document,
        category.key,
        recover_legacy=True,
    )
    if source is None or not source.is_file():
        raise HTTPException(
            status_code=409,
            detail="未找到该合同在当前密级目录中的原件，请联系管理员核对NAS文件",
        )
    previous_folder = _contract_document_folder_path(str(source), category.key)
    moved = _move_nas_file_without_overwrite(source, root / relative)
    _set_contract_document_source(db, document, moved)
    if Path(document.file_blob.source_path).resolve() == source.resolve():
        document.file_blob.source_path = str(moved)
    db.add(
        AuditLog(
            user_id=user.id,
            action="contract_folder_move",
            document_ids_json=json.dumps([document.id]),
            details_json=json.dumps(
                {
                    "category": category.key,
                    "from": previous_folder,
                    "to": relative.as_posix(),
                },
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    return {
        "document_id": document.id,
        "category": category.key,
        "folder_path": relative.as_posix(),
    }


@app.post("/v1/contracts/uploads")
async def upload_contract(
    file: UploadFile = File(...),
    category: str = Form(...),
    relative_path: str = Form(default=""),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    if not _can_upload_contract_category(user, category):
        raise HTTPException(
            status_code=403,
            detail="无权向该合同分类上传资料",
        )
    contract_category = _contract_category_for_user(
        user,
        category,
        operation="upload",
    )
    filename = _safe_upload_filename(file.filename)
    if Path(filename).suffix.lower() not in CONTRACT_UPLOAD_SUFFIXES:
        raise HTTPException(
            status_code=415,
            detail="合同档案仅支持PDF和Word文件",
        )
    try:
        target_dir = ensure_contract_layout(settings.knowledge_root)[category]
        relative_parent = _safe_upload_relative_parent(relative_path)
        if relative_parent.parts:
            target_dir = (target_dir / relative_parent).resolve()
            target_dir.relative_to(settings.knowledge_root.resolve())
            target_dir.mkdir(parents=True, exist_ok=True)
    except (OSError, ValueError) as exc:
        raise HTTPException(status_code=503, detail="合同档案目录暂不可写入") from exc
    temporary = target_dir / f".{secrets.token_hex(12)}.part"
    maximum = settings.inbox_max_file_bytes
    size = 0
    try:
        with temporary.open("wb") as destination:
            while True:
                block = await file.read(1024 * 1024)
                if not block:
                    break
                size += len(block)
                if size > maximum:
                    raise HTTPException(
                        status_code=413,
                        detail="文件超过当前单文件大小上限",
                    )
                destination.write(block)
        if size == 0:
            raise HTTPException(status_code=422, detail="不能上传空文件")
        final_path = _commit_uploaded_file(temporary, target_dir, filename)
        stat = final_path.stat()
        entry = prefill_entry(
            final_path,
            settings.knowledge_root,
            domain_override=contract_category.domain,
        )
        entry.update(
            {
                "project_name": f"{contract_category.name} · {final_path.stem}",
                "role": "contract",
                "confidentiality": contract_category.confidentiality,
                "knowledge_status": "candidate",
                "force_review": True,
            }
        )
        result = ingest_document(
            db,
            entry,
            expected_stat=(stat.st_size, stat.st_mtime_ns),
        )
        document = db.get(Document, result["document_id"])
        if result.get("duplicate_filtered"):
            canonical_path = (
                Path(document.file_blob.source_path).resolve()
                if document and document.file_blob
                else None
            )
            assigned_path = (
                Path(document.contract_source.source_path).resolve()
                if document and document.contract_source
                else canonical_path
            )
            if assigned_path is not None and final_path.resolve() != assigned_path:
                final_path.unlink(missing_ok=True)
                final_path = assigned_path
        filing = {
            "mode": "preserved" if relative_parent.parts else "existing",
            "folder_path": relative_parent.as_posix() if relative_parent.parts else "",
            "confidence": None,
            "matched_keywords": [],
        }
        if not relative_parent.parts and not result.get("duplicate_filtered"):
            reference_count = db.scalar(
                select(func.count(Document.id)).where(
                    Document.content_hash == document.content_hash
                )
            ) if document else 0
            if (
                document is not None
                and document.file_blob is not None
                and int(reference_count or 0) == 1
                and Path(document.file_blob.source_path).resolve() == final_path.resolve()
            ):
                chunk_texts = db.scalars(
                    select(Chunk.text)
                    .where(Chunk.document_id == document.id)
                    .order_by(Chunk.chunk_index.asc())
                    .limit(60)
                ).all()
                suggestion = suggest_contract_folder(
                    category,
                    title=document.title or final_path.name,
                    text="\n".join(chunk_texts),
                )
                category_root = ensure_contract_layout(settings.knowledge_root)[category]
                target_folder = _safe_contract_folder_path(suggestion.folder)
                moved_path = _move_nas_file_without_overwrite(
                    final_path,
                    category_root / target_folder,
                )
                document.file_blob.source_path = str(moved_path)
                final_path = moved_path
                filing = {
                    "mode": "local_ai",
                    "folder_path": target_folder.as_posix(),
                    "confidence": suggestion.confidence,
                    "matched_keywords": list(suggestion.matched_keywords),
                }
                db.add(
                    AuditLog(
                        user_id=user.id,
                        action="contract_auto_filed",
                        document_ids_json=json.dumps([document.id]),
                        details_json=json.dumps(
                            {
                                "category": category,
                                "folder_path": target_folder.as_posix(),
                                "confidence": suggestion.confidence,
                                "matched_keywords": list(suggestion.matched_keywords),
                                "classifier": "local_rules_v1",
                            },
                            ensure_ascii=False,
                        ),
                    )
                )
        if document is not None and final_path.is_file():
            _set_contract_document_source(db, document, final_path)
        existing_owner = db.scalar(
            select(ContractDocumentOwner).where(
                ContractDocumentOwner.document_id == result["document_id"],
                ContractDocumentOwner.uploaded_by_user_id == user.id,
            )
        )
        if existing_owner is None:
            db.add(
                ContractDocumentOwner(
                    document_id=result["document_id"],
                    uploaded_by_user_id=user.id,
                )
            )
        db.add(
            AuditLog(
                user_id=user.id,
                action="contract_upload",
                document_ids_json=json.dumps([result["document_id"]]),
                details_json=json.dumps(
                    {
                        "category": contract_category.key,
                        "confidentiality": contract_category.confidentiality,
                        "size_bytes": size,
                        "review_required": True,
                        "relative_path": relative_path,
                        "filing": filing,
                    },
                    ensure_ascii=False,
                ),
            )
        )
        db.commit()
        return {
            "filename": final_path.name,
            "category": contract_category.key,
            "category_name": contract_category.name,
            "confidentiality": contract_category.confidentiality,
            "size_bytes": size,
            "status": result["status"],
            "document_id": result["document_id"],
            "review_required": True,
            "relative_path": relative_path,
            "filing": filing,
            "notice": (
                f"合同原件已保存至NAS，并已在本地自动归档到“{filing['folder_path']}”；审核通过后可检索。"
                if filing["mode"] == "local_ai"
                else "合同原件已保存至NAS并保留原文件夹层级；审核通过后可检索。"
                if filing["mode"] == "preserved"
                else "合同原件已保存至NAS；审核通过后可检索。"
            ),
            "duplicate_filtered": result.get("duplicate_filtered", False),
        }
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    finally:
        await file.close()


@app.get("/readyz")
def readiness(db: Session = Depends(get_db)):
    payload = collect_operations_status(db)
    public = _public_health(payload)
    return JSONResponse(
        public,
        status_code=503 if payload["status"] == "critical" else 200,
    )


@app.get("/monitorz")
def monitoring(db: Session = Depends(get_db)):
    payload = collect_operations_status(db)
    public = _public_health(payload)
    return JSONResponse(
        public,
        status_code=503 if payload["status"] != "ok" else 200,
    )


@app.post("/v1/auth/login")
def login(
    payload: LoginRequest,
    request: Request,
    db: Session = Depends(get_db),
) -> dict:
    username = payload.username.strip().lower()
    username_hash = hashlib.sha256(username.encode()).hexdigest()
    retry_after = login_rate_limiter.retry_after(username)
    if retry_after:
        db.add(
            AuditLog(
                user_id=None,
                action="login_rate_limited",
                query_hash=username_hash,
                details_json=json.dumps({"retry_after": retry_after}),
            )
        )
        db.commit()
        raise HTTPException(
            status_code=429,
            detail="登录尝试过多，请稍后再试",
            headers={"Retry-After": str(retry_after)},
        )
    user = authenticate(db, payload.username, payload.password)
    if not user:
        retry_after = login_rate_limiter.record_failure(username)
        db.add(
            AuditLog(
                user_id=None,
                action="login_failed",
                query_hash=username_hash,
                details_json=json.dumps(
                    {"locked": bool(retry_after)},
                    ensure_ascii=False,
                ),
            )
        )
        db.commit()
        if retry_after:
            raise HTTPException(
                status_code=429,
                detail="登录尝试过多，请稍后再试",
                headers={"Retry-After": str(retry_after)},
            )
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="账号或密码错误",
        )
    login_rate_limiter.record_success(username)
    token, expires_at = create_session(db, user)
    db.add(
        AuditLog(
            user_id=user.id,
            action="login",
            details_json=json.dumps(
                {
                    "environment": settings.environment,
                    "client": request.client.host if request.client else None,
                }
            ),
        )
    )
    db.commit()
    return {
        "access_token": token,
        "token_type": "bearer",
        "expires_at": expires_at,
        "user": user_payload(user),
    }


@app.post("/v1/auth/logout")
def logout(
    token: str = Depends(raw_bearer_token),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    revoke_session(db, token)
    db.add(AuditLog(user_id=user.id, action="logout"))
    db.commit()
    return {"status": "ok"}


@app.get("/v1/me")
def me(user: User = Depends(current_user)) -> dict:
    return user_payload(user)


def _category_payload(category: KnowledgeCategory) -> dict:
    return {
        "id": category.id,
        "key": category.key,
        "name": category.name,
        "active": category.active,
        "sort_order": category.sort_order,
        "created_at": category.created_at,
        "updated_at": category.updated_at,
    }


def _validate_active_categories(db: Session, keys: list[str]) -> list[str]:
    normalized = list(dict.fromkeys(str(item).strip() for item in keys))
    if not normalized or any(not item for item in normalized):
        raise HTTPException(status_code=422, detail="至少选择一个有效资料分类")
    found = set(
        db.scalars(
            select(KnowledgeCategory.key).where(
                KnowledgeCategory.key.in_(normalized),
                KnowledgeCategory.active.is_(True),
            )
        ).all()
    )
    if found != set(normalized):
        raise HTTPException(status_code=422, detail="资料分类不存在或已停用")
    return normalized


@app.get("/v1/categories")
def list_categories(
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> list[dict]:
    query = select(KnowledgeCategory)
    if user.role != "founder":
        query = query.where(KnowledgeCategory.active.is_(True))
    categories = db.scalars(
        query.order_by(
            KnowledgeCategory.active.desc(),
            KnowledgeCategory.sort_order.asc(),
            KnowledgeCategory.created_at.asc(),
        )
    ).all()
    return [_category_payload(item) for item in categories]


@app.post("/v1/admin/categories")
def create_category(
    payload: KnowledgeCategoryCreate,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    require_founder(user)
    try:
        name = validate_category_folder_name(payload.name)
    except StorageLayoutError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    if db.scalar(
        select(KnowledgeCategory.id).where(
            func.lower(KnowledgeCategory.name) == name.lower()
        )
    ):
        raise HTTPException(status_code=409, detail="资料分类名称已存在")
    highest = db.scalar(select(func.max(KnowledgeCategory.sort_order))) or 0
    category = KnowledgeCategory(
        key=f"cat_{secrets.token_hex(6)}",
        name=name,
        active=True,
        sort_order=highest + 10,
    )
    db.add(category)
    db.flush()
    try:
        ensure_category_layout(settings.knowledge_root, category.name)
    except (OSError, StorageLayoutError) as exc:
        db.rollback()
        raise HTTPException(status_code=503, detail="无法创建资料分类文件夹") from exc
    db.add(
        AuditLog(
            user_id=user.id,
            action="create_knowledge_category",
            details_json=json.dumps(
                {"category_id": category.id, "category_key": category.key},
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    return _category_payload(category)


@app.patch("/v1/admin/categories/{category_id}")
def update_category(
    category_id: str,
    payload: KnowledgeCategoryUpdate,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    require_founder(user)
    category = db.get(KnowledgeCategory, category_id)
    if category is None:
        raise HTTPException(status_code=404, detail="资料分类不存在")
    changes = payload.model_dump(exclude_unset=True)
    source_path_updates = 0
    if "name" in changes:
        try:
            name = validate_category_folder_name(changes["name"])
        except StorageLayoutError as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc
        duplicate = db.scalar(
            select(KnowledgeCategory.id).where(
                func.lower(KnowledgeCategory.name) == name.lower(),
                KnowledgeCategory.id != category.id,
            )
        )
        if duplicate:
            raise HTTPException(status_code=409, detail="资料分类名称已存在")
        try:
            source_path_updates = rename_category_layout(
                db,
                settings.knowledge_root,
                category.name,
                name,
            )
        except (OSError, StorageLayoutError) as exc:
            db.rollback()
            raise HTTPException(status_code=409, detail=str(exc)) from exc
        category.name = name
    if changes.get("active") is False and category.active:
        active_count = db.scalar(
            select(func.count(KnowledgeCategory.id)).where(
                KnowledgeCategory.active.is_(True)
            )
        ) or 0
        if active_count <= 1:
            raise HTTPException(status_code=409, detail="至少保留一个启用的资料分类")
    if "active" in changes:
        category.active = changes["active"]
    if "sort_order" in changes:
        category.sort_order = changes["sort_order"]
    db.add(
        AuditLog(
            user_id=user.id,
            action="update_knowledge_category",
            details_json=json.dumps(
                {
                    "category_id": category.id,
                    "changed_fields": sorted(changes),
                    "active": category.active,
                    "source_path_updates": source_path_updates,
                },
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    return _category_payload(category)


@app.post("/v1/me/password")
def change_own_password(
    payload: OwnPasswordChange,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    if not verify_password(payload.current_password, user.password_hash):
        raise HTTPException(status_code=400, detail="当前密码错误")
    if payload.current_password == payload.new_password:
        raise HTTPException(status_code=400, detail="新密码不能与当前密码相同")
    user.password_hash = hash_password(payload.new_password)
    db.add(
        AuditLog(
            user_id=user.id,
            action="change_own_password",
            details_json=json.dumps({"sessions_revoked": True}),
        )
    )
    revoke_user_sessions(db, user.id)
    return {"status": "password_changed", "login_required": True}


@app.get("/v1/admin/users")
def list_local_users(
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> list[dict]:
    require_founder(user)
    users = db.scalars(
        select(User).order_by(User.active.desc(), User.created_at.asc())
    ).all()
    return [admin_user_payload(item) for item in users]


@app.post("/v1/admin/users")
def create_local_user(
    payload: AdminUserCreate,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    require_founder(user)
    username = payload.username.strip().lower()
    existing = db.scalar(
        select(User).where(func.lower(User.username) == username)
    )
    if existing:
        raise HTTPException(status_code=409, detail="账号已存在")
    created = User(
        username=username,
        display_name=payload.display_name.strip(),
        password_hash=hash_password(payload.password),
        role=ORGANIZATION_ACCESS_ROLE[payload.role],
        organization_role=payload.role,
        confidentiality_ceiling="L1" if payload.role == "education" else payload.confidentiality_ceiling,
        departments_json='["*"]',
        active=True,
    )
    db.add(created)
    db.flush()
    db.add(
        AuditLog(
            user_id=user.id,
            action="create_local_user",
            details_json=json.dumps(
                {
                    "target_user_id": created.id,
                    "username": created.username,
                    "role": created.organization_role,
                    "access_role": created.role,
                    "confidentiality_ceiling": created.confidentiality_ceiling,
                    "departments": json.loads(created.departments_json),
                },
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    return admin_user_payload(created)


@app.patch("/v1/admin/users/{target_user_id}")
def update_local_user(
    target_user_id: str,
    payload: AdminUserUpdate,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    require_founder(user)
    target = db.get(User, target_user_id)
    if not target:
        raise HTTPException(status_code=404, detail="账号不存在")
    before = admin_user_payload(target)
    changes = payload.model_dump(exclude_unset=True)
    removes_management_access = (
        target.role == "founder"
        and target.active
        and (
            changes.get("active") is False
            or (
                "role" in changes
                and ORGANIZATION_ACCESS_ROLE[changes["role"]] != "founder"
            )
        )
    )
    if removes_management_access:
        other_active_managers = db.scalar(
            select(func.count(User.id)).where(
                User.id != target.id,
                User.active.is_(True),
                User.role == "founder",
            )
        ) or 0
        if other_active_managers < 1:
            raise HTTPException(
                status_code=409,
                detail=(
                    "这是最后一个启用的管理账号。请先创建或提升另一个管理账号，"
                    "再停用或调整当前账号。"
                ),
            )
    if "display_name" in changes:
        target.display_name = changes["display_name"].strip()
    if "role" in changes:
        target.organization_role = changes["role"]
        target.role = ORGANIZATION_ACCESS_ROLE[changes["role"]]
    if "confidentiality_ceiling" in changes:
        target.confidentiality_ceiling = changes["confidentiality_ceiling"]
    if target.organization_role == "education":
        target.confidentiality_ceiling = "L1"
    if "departments" in changes:
        changes["departments"] = ["*"]
        target.departments_json = '["*"]'
    if "active" in changes:
        target.active = changes["active"]
    db.flush()
    if not target.active:
        revoke_user_sessions(db, target.id)
    db.add(
        AuditLog(
            user_id=user.id,
            action="update_local_user",
            details_json=json.dumps(
                {
                    "target_user_id": target.id,
                    "before": {
                        "role": before["organization_role"],
                        "access_role": before["role"],
                        "confidentiality_ceiling": before[
                            "confidentiality_ceiling"
                        ],
                        "departments": before["departments"],
                        "active": before["active"],
                    },
                    "after": {
                        "role": target.organization_role,
                        "access_role": target.role,
                        "confidentiality_ceiling": target.confidentiality_ceiling,
                        "departments": json.loads(target.departments_json),
                        "active": target.active,
                    },
                },
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    return admin_user_payload(target)


@app.post("/v1/admin/users/{target_user_id}/password")
def reset_local_user_password(
    target_user_id: str,
    payload: PasswordReset,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    require_founder(user)
    target = db.get(User, target_user_id)
    if not target:
        raise HTTPException(status_code=404, detail="账号不存在")
    target.password_hash = hash_password(payload.password)
    db.add(
        AuditLog(
            user_id=user.id,
            action="reset_local_user_password",
            details_json=json.dumps(
                {"target_user_id": target.id, "sessions_revoked": True}
            ),
        )
    )
    revoke_user_sessions(db, target.id)
    return {"status": "password_reset", "target_user_id": target.id}


@app.get("/v1/status")
def service_status(
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    projects = db.scalars(select(Project)).all()
    visible_projects = [
        project
        for project in projects
        if any(is_authorized(user, document, project) for document in project.documents)
        or not project.documents
    ]
    documents = db.scalars(
        select(Document).options(joinedload(Document.file_blob))
    ).all()
    source_health = {
        row.content_hash: row
        for row in db.scalars(select(SourceHealth)).all()
    }
    document_count = len(documents)
    source_available_count = sum(
        1
        for document in documents
        if source_is_available(
            document.file_blob,
            source_health.get(document.content_hash),
        )
    )
    chunk_count = db.scalar(select(func.count(Chunk.id))) or 0
    embedding_count = db.scalar(select(func.count(ChunkEmbedding.id))) or 0
    current_count = sum(
        1 for document in documents if document.knowledge_status == "current"
    )
    candidate_count = sum(
        1 for document in documents if document.knowledge_status == "candidate"
    )
    return {
        "status": "ready",
        "environment": settings.environment,
        "database": "sqlite" if settings.database_url.startswith("sqlite") else "postgresql",
        "retrieval": {
            "exact": True,
            "semantic": bool(
                settings.embedding_enabled
                and settings.gateway_api_key
                and embedding_count
            ),
            "llm_generation": bool(
                settings.llm_enabled and settings.gateway_api_key
            ),
            "outbound_enabled": bool(
                settings.gateway_api_key
                and (settings.embedding_enabled or settings.llm_enabled)
            ),
        },
        "counts": {
            "projects": len(visible_projects),
            "documents": document_count if user.role == "founder" else None,
            "chunks": chunk_count if user.role == "founder" else None,
            "embeddings": embedding_count if user.role == "founder" else None,
            "current_documents": current_count if user.role == "founder" else None,
            "candidate_documents": candidate_count if user.role == "founder" else None,
            "source_available_documents": (
                source_available_count if user.role == "founder" else None
            ),
            "source_missing_documents": (
                document_count - source_available_count
                if user.role == "founder"
                else None
            ),
        },
        "policy_version": settings.policy_version,
        "generated_at": datetime.now(timezone.utc),
    }


@app.get("/v1/operations/status")
def operations_status(
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    if user.role not in {"founder", "knowledge_admin"}:
        raise HTTPException(status_code=403, detail="无权查看系统运行状态")
    return collect_operations_status(db)


@app.post("/v1/operations/reconcile-sources")
def trigger_source_reconciliation(
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    if user.role not in {"founder", "knowledge_admin"}:
        raise HTTPException(status_code=403, detail="无权执行原件核验")
    result = reconcile_sources(db, settings.knowledge_root)
    db.add(
        AuditLog(
            user_id=user.id,
            action="manual_source_reconcile",
            details_json=json.dumps(
                {
                    "status": result["status"],
                    "counts": result["counts"],
                },
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    return result


@app.get("/v1/evaluations/latest")
def latest_evaluation(
    user: User = Depends(current_user),
) -> dict:
    if user.role not in {"founder", "knowledge_admin"}:
        raise HTTPException(status_code=403, detail="无权查看评测结果")
    report = load_latest_evaluation()
    if not report:
        raise HTTPException(status_code=404, detail="尚未运行技术评测")
    return report


@app.get("/v1/evaluations/business-gold")
def get_business_gold_registry(
    user: User = Depends(current_user),
) -> dict:
    if user.role not in {"founder", "knowledge_admin"}:
        raise HTTPException(status_code=403, detail="无权查看业务金标")
    return business_gold_registry()


@app.get("/v1/evaluations/business-gold/sources")
def list_business_gold_sources(
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> list[dict]:
    if user.role not in {"founder", "knowledge_admin"}:
        raise HTTPException(status_code=403, detail="无权查看业务金标来源")
    health_by_hash = {
        row.content_hash: row
        for row in db.scalars(select(SourceHealth)).all()
    }
    documents = db.scalars(
        select(Document)
        .options(
            joinedload(Document.project),
            joinedload(Document.file_blob),
        )
        .where(
            Document.knowledge_status.in_({
                "current",
                "approved",
                "superseded",
                "archived",
            })
        )
        .order_by(Document.title)
    ).all()
    return [
        {
            "document_id": document.id,
            "title": document.title,
            "version": document.version,
            "page_count": document.page_count,
            "knowledge_status": document.knowledge_status,
            "confidentiality": document.confidentiality,
            "project_id": document.project.id,
            "project": document.project.name,
        }
        for document in documents
        if is_authorized(user, document, document.project)
        and source_is_available(
            document.file_blob,
            health_by_hash.get(document.content_hash),
        )
    ]


@app.post("/v1/evaluations/business-gold/import")
def import_business_gold_drafts(
    payload: BusinessGoldImportRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    if user.role not in {"founder", "knowledge_admin"}:
        raise HTTPException(status_code=403, detail="无权导入业务金标草稿")
    try:
        result = upsert_business_gold_drafts(
            payload.cases,
            user=user,
        )
    except ValueError as exc:
        details = {
            "duplicate_id": "导入内容存在重复题目 ID",
            "registry_invalid": "现有业务金标题库格式异常，请先修复",
            "invalid_id": "题目 ID 仅支持字母、数字、点、横线、下划线和冒号",
            "invalid_question": "问题长度必须为 5–1000 个字符",
            "invalid_expected": "每道题必须设置预期结果",
            "invalid_expected_mode": "引用、应拒答、权限无泄露三种预期只能选择一种",
        }
        raise HTTPException(
            status_code=422,
            detail=details.get(str(exc), f"业务金标格式无效：{exc}"),
        ) from exc
    db.add(
        AuditLog(
            user_id=user.id,
            action="business_gold_drafts_imported",
            details_json=json.dumps(
                {
                    "imported_ids": [str(item.get("id", "")) for item in payload.cases],
                    "total": result["total"],
                    "approved": result["approved"],
                },
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    return result


@app.post("/v1/evaluations/business-gold/generate")
def generate_business_gold_drafts(
    payload: BusinessGoldGenerateRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    if user.role not in {"founder", "knowledge_admin"}:
        raise HTTPException(status_code=403, detail="无权生成业务金标候选题")
    result = generate_business_gold_candidates(
        db,
        user=user,
        target=payload.target,
    )
    generation = result.get("generation", {})
    db.add(
        AuditLog(
            user_id=user.id,
            action="business_gold_candidates_generated",
            details_json=json.dumps(
                {
                    "requested": generation.get("requested", payload.target),
                    "created": generation.get("created", 0),
                    "shortfall": generation.get("shortfall", 0),
                    "counts": generation.get("counts", {}),
                },
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    return result


@app.post("/v1/evaluations/business-gold/approve")
def approve_business_gold(
    payload: BusinessGoldApprovalRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    if user.role != "founder":
        raise HTTPException(status_code=403, detail="仅创始人可批准业务金标")
    try:
        result = approve_business_gold_cases(
            db,
            payload.case_ids,
            user=user,
            owner_note=payload.owner_note,
        )
    except ValueError as exc:
        detail = str(exc)
        if detail == "registry_invalid":
            message = "现有业务金标题库格式异常，请先修复"
        elif detail == "case_not_found":
            message = "所选题目不存在，请刷新后重试"
        elif detail.startswith("answer_key_incomplete:"):
            message = "命中来源题必须填写标准答案和至少一个关键点后才能批准"
        elif detail.startswith("source_not_eligible:"):
            message = "所选题目的来源不是可用正式资料，不能批准"
        else:
            message = "业务金标无法批准"
        raise HTTPException(status_code=409, detail=message) from exc
    db.add(
        AuditLog(
            user_id=user.id,
            action="business_gold_cases_approved",
            details_json=json.dumps(
                {
                    "case_ids": list(dict.fromkeys(payload.case_ids)),
                    "approved": result["approved"],
                },
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    return result


@app.post("/v1/evaluations/run")
def run_evaluation(
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    if user.role not in {"founder", "knowledge_admin"}:
        raise HTTPException(status_code=403, detail="无权运行技术评测")
    report = run_technical_evaluation(db)
    db.add(
        AuditLog(
            user_id=user.id,
            action="evaluation_run",
            details_json=json.dumps(
                {
                    "technical_status": report["technical_status"],
                    "technical_case_count": report["technical_case_count"],
                    "business_gold_total": report["business_gold_total"],
                    "metrics": report["metrics"],
                },
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    return report


@app.post("/v1/evaluations/concurrency")
def run_concurrency(
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    if user.role not in {"founder", "knowledge_admin"}:
        raise HTTPException(status_code=403, detail="无权运行并发测试")
    report = run_concurrency_evaluation(SessionLocal)
    db.add(
        AuditLog(
            user_id=user.id,
            action="evaluation_concurrency",
            details_json=json.dumps(
                {
                    "status": report["status"],
                    "concurrent_users": report["concurrent_users"],
                    "request_count": report["request_count"],
                    "error_count": report["error_count"],
                    "side_channel_leak_count": (
                        report["side_channel_leak_count"]
                    ),
                    "p95_ms": report["p95_ms"],
                },
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    return report


@app.post("/v1/search")
def knowledge_search(
    payload: SearchRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    result = search(
        db,
        user=user,
        query=payload.query,
        requested_scope=payload.scope,
        requested_retrieval=payload.retrieval,
        limit=payload.limit,
        offset=payload.offset,
        category=payload.category,
    )
    return {
        **result,
        "user_id": user.id,
        "policy_version": settings.policy_version,
        "generated_at": datetime.now(timezone.utc),
    }


@app.get("/v1/contracts")
def list_contract_documents(
    category: str = Query(..., min_length=1, max_length=40),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    contract_category = _contract_category_for_user(user, category)
    access_scope = _contract_scope_for_user(user, category)
    owned_document_ids = (
        _owned_contract_document_ids(db, user)
        if access_scope == "own"
        else None
    )
    if owned_document_ids is not None and not owned_document_ids:
        documents = []
    else:
        statement = (
            select(Document)
            .join(Document.project)
            .options(
                joinedload(Document.project),
                joinedload(Document.file_blob),
                joinedload(Document.contract_source),
            )
            .where(
                Project.domain == contract_category.domain,
                Document.knowledge_status.in_(
                    ("approved", "current", "superseded", "archived")
                ),
            )
            .order_by(Document.ingested_at.desc())
            .limit(200)
        )
        if owned_document_ids is not None:
            statement = statement.where(Document.id.in_(owned_document_ids))
        documents = db.scalars(statement).all()
    pending_statement = (
        select(Document)
        .join(Document.project)
        .options(joinedload(Document.project))
        .where(
            Project.domain == contract_category.domain,
            Document.knowledge_status == "candidate",
        )
    )
    if owned_document_ids is not None:
        pending_statement = pending_statement.where(Document.id.in_(owned_document_ids))
    pending_documents = db.scalars(pending_statement).all()
    pending_count = sum(
        1
        for document in pending_documents
        if access_scope == "own" or is_authorized(user, document, document.project)
    )
    items: list[dict] = []
    for document in documents:
        if access_scope != "own" and not is_authorized(
            user, document, document.project
        ):
            continue
        source = _contract_document_source_path(document, category)
        if source is None or not source.is_file():
            continue
        items.append(
            {
                "document_id": document.id,
                "title": document.title,
                "version": document.version,
                "page_count": document.page_count,
                "citation_basis": document.citation_basis,
                "knowledge_status": document.knowledge_status,
                "confidentiality": document.confidentiality,
                "created_at": document.ingested_at,
                "folder_path": _contract_document_folder_path(
                    str(source), category
                ),
            }
        )
    db.add(
        AuditLog(
            user_id=user.id,
            action="contract_archive_list",
            document_ids_json=json.dumps([item["document_id"] for item in items]),
            details_json=json.dumps(
                {
                    "category": contract_category.key,
                    "result_count": len(items),
                },
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    return {
        "category": contract_category.key,
        "category_name": contract_category.name,
        "confidentiality": contract_category.confidentiality,
        "search_scope": access_scope,
        "pending_count": pending_count,
        "items": items,
    }


@app.post("/v1/contracts/search")
def search_contract_documents(
    payload: ContractSearchRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    contract_category = _contract_category_for_user(user, payload.category)
    access_scope = _contract_scope_for_user(user, payload.category)
    owned_document_ids = (
        _owned_contract_document_ids(db, user)
        if access_scope == "own"
        else None
    )
    result = search(
        db,
        user=user,
        query=payload.query,
        requested_scope="history",
        requested_retrieval="exact",
        limit=payload.limit,
        category=contract_category.domain,
        generate=False,
        include_contracts=True,
        restricted_document_ids=owned_document_ids,
        audit_action="contract_search",
    )
    return {
        **result,
        "category": contract_category.key,
        "category_name": contract_category.name,
        "confidentiality": contract_category.confidentiality,
        "local_only": True,
        "search_scope": access_scope,
        "generated_at": datetime.now(timezone.utc),
    }


WRITING_CATEGORY_ALIASES = {
    "youth": ("青训", "选秀", "试训", "新秀"),
    "training": ("培训", "课程", "教学", "实训", "训练营"),
    "team": ("战队", "俱乐部", "选手", "教练"),
    "tournament": ("赛事", "比赛", "联赛", "赛制"),
    "venue": ("场馆", "场地", "电竞馆"),
    "company": ("公司介绍", "企业介绍", "资质", "公司能力"),
}


def _infer_writing_category(
    instruction: str,
    categories: list[KnowledgeCategory],
) -> KnowledgeCategory | None:
    value = instruction.casefold()
    scored: list[tuple[int, KnowledgeCategory]] = []
    for category in categories:
        score = 0
        compact_name = category.name.casefold().replace("电竞", "").replace("资料", "")
        if len(compact_name) >= 2 and compact_name in value:
            score += 10
        for alias in WRITING_CATEGORY_ALIASES.get(category.key, ()):
            if alias.casefold() in value:
                score += 4
        if score:
            scored.append((score, category))
    if not scored:
        return None
    return max(scored, key=lambda item: (item[0], -item[1].sort_order))[1]


def _merge_writing_results(*groups: list[dict], limit: int = 10) -> list[dict]:
    merged: list[dict] = []
    seen_documents: set[str] = set()
    for group in groups:
        for item in group:
            if item["document_id"] in seen_documents:
                continue
            seen_documents.add(item["document_id"])
            merged.append(item)
            if len(merged) >= limit:
                return merged
    return merged


@app.post("/v1/writing/draft")
def writing_draft(
    payload: WritingDraftRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    if payload.allow_l3_generation and user.role != "founder":
        raise HTTPException(status_code=403, detail="仅创始人可逐次授权L3资料用于AI生成")
    active_categories = db.scalars(
        select(KnowledgeCategory)
        .where(KnowledgeCategory.active.is_(True))
        .order_by(KnowledgeCategory.sort_order, KnowledgeCategory.name)
    ).all()
    category = next(
        (item for item in active_categories if item.key == payload.category),
        None,
    ) if payload.category else None
    if payload.category and category is None:
        raise HTTPException(status_code=422, detail="资料分类不存在或已停用")
    inferred_category = (
        None
        if category is not None
        else _infer_writing_category(payload.instruction, active_categories)
    )
    effective_category = category or inferred_category
    evidence_result = search(
        db,
        user=user,
        query=payload.instruction,
        requested_scope=payload.scope,
        requested_retrieval="hybrid",
        limit=8 if effective_category else 10,
        audit=False,
        category=effective_category.key if effective_category else None,
        generate=False,
    )
    primary_results = evidence_result["results"]
    broad_result = None
    if inferred_category is not None:
        broad_result = search(
            db,
            user=user,
            query=payload.instruction,
            requested_scope=payload.scope,
            requested_retrieval="hybrid",
            limit=4,
            audit=False,
            category=None,
            generate=False,
        )
    results = _merge_writing_results(
        primary_results,
        broad_result["results"] if broad_result else [],
        limit=10,
    )
    founder = user.role == "founder"
    allow_l3 = bool(payload.allow_l3_generation and founder)
    allowed_pairs: list[tuple[dict, GenerationEvidence]] = []
    restricted_source_count = 0
    for item in results:
        evidence = GenerationEvidence(
            document_id=item["document_id"],
            title=item["title"],
            version=item["version"],
            page=item["page"],
            status=item["knowledge_status"],
            confidentiality=item["effective_confidentiality"],
            excerpt=item["excerpt"],
        )
        if evidence_allowed(
            evidence.confidentiality,
            founder=founder,
            allow_l3=allow_l3,
        ):
            allowed_pairs.append((item, evidence))
        else:
            restricted_source_count += 1
    allowed_pairs = allowed_pairs[:8]
    generation_sources = [item for item, _evidence in allowed_pairs]
    outbound_evidence = [evidence for _item, evidence in allowed_pairs]
    draft = ""
    generation_mode = "failed"
    generation_model = None
    fallback_used = False
    generation_degraded = True
    outbound_document_ids: list[str] = []
    outbound_characters = 0
    error_code = None
    failure_status = 0
    failure_detail = ""
    if not results:
        error_code = "writing_evidence_not_found"
        failure_status = 422
        failure_detail = "智库中未找到足够相关的资料，无法生成可靠初稿。请补充关键词或先上传资料。"
    elif not outbound_evidence:
        generation_mode = "local_only"
        error_code = "writing_no_authorized_outbound_evidence"
        failure_status = 409
        failure_detail = (
            "命中的资料均为L3敏感资料。创始人可勾选“允许本次使用L3资料生成”后重试。"
            if founder
            and any(
                item["effective_confidentiality"] == "L3"
                for item in results
            )
            else "命中的资料受密级保护，不能发送至AI生成模型；L4/L5资料永不出网。"
        )
    elif not settings.llm_enabled or not settings.gateway_api_key:
        error_code = "writing_generation_service_disabled"
        failure_status = 503
        failure_detail = "AI生成服务当前不可用，系统不会用检索片段冒充初稿。请稍后重试。"
    else:
        try:
            generated = generate_grounded_draft(
                payload.instruction,
                outbound_evidence,
                founder=founder,
                allow_l3=allow_l3,
            )
            draft = generated.answer
            generation_mode = "llm"
            generation_model = generated.model
            fallback_used = generated.fallback_used
            generation_degraded = False
            outbound_document_ids = generated.evidence_document_ids
            outbound_characters = generated.outbound_characters
        except GenerationServiceError as exc:
            error_code = str(exc)
            failure_status = 502
            if error_code in {
                "gateway_timeout",
                "gateway_unreachable",
                "gateway_response_invalid",
                "gateway_http_408",
                "gateway_http_425",
                "gateway_http_500",
                "gateway_http_502",
                "gateway_http_503",
                "gateway_http_504",
            }:
                failure_detail = (
                    "AI模型网关本次没有正常响应。系统已经依次尝试主模型和备用模型，"
                    "没有用检索片段冒充初稿；请直接再次点击生成。"
                )
            elif error_code == "gateway_http_429":
                failure_detail = (
                    "AI模型当前请求较多，主模型和备用模型都触发了限流。"
                    "输入内容已经保留，请稍后再次点击生成。"
                )
            else:
                failure_detail = (
                    "AI返回的内容没有通过引用或完整性校验，因此没有展示低质量初稿。"
                    "输入内容已经保留，请再次生成。"
                )

    instruction_hash = hashlib.sha256(payload.instruction.encode()).hexdigest()
    document_ids = list(dict.fromkeys(item["document_id"] for item in results))
    db.add(
        AuditLog(
            user_id=user.id,
            action="writing_draft",
            query_hash=instruction_hash,
            document_ids_json=json.dumps(document_ids),
            details_json=json.dumps(
                {
                    "category": payload.category,
                    "scope": payload.scope,
                    "effective_category": (
                        effective_category.key if effective_category else None
                    ),
                    "category_auto_inferred": inferred_category is not None,
                    "result_count": len(results),
                    "generation_source_count": len(generation_sources),
                    "restricted_source_count": restricted_source_count,
                    "l3_authorized_for_request": allow_l3,
                    "generation_mode": generation_mode,
                    "generation_degraded": generation_degraded,
                    "fallback_used": fallback_used,
                },
                ensure_ascii=False,
            ),
        )
    )
    if results:
        db.add(
            AuditLog(
                user_id=user.id,
                action=(
                    "writing_llm_generation"
                    if generation_mode == "llm"
                    else "writing_llm_generation_skipped"
                    if generation_mode == "local_only"
                    else "writing_llm_generation_failed"
                ),
                query_hash=instruction_hash,
                document_ids_json=json.dumps(outbound_document_ids),
                details_json=json.dumps(
                    {
                        "model": generation_model,
                        "fallback_used": fallback_used,
                        "outbound_characters": outbound_characters,
                        "error_code": error_code,
                        "l3_authorized_for_request": allow_l3,
                    },
                    ensure_ascii=False,
                ),
            )
        )
    db.commit()
    if failure_status:
        raise HTTPException(status_code=failure_status, detail=failure_detail)
    response_payload = {
        "draft": draft,
        "sources": generation_sources,
        "generation_mode": generation_mode,
        "generation_model": generation_model,
        "generation_degraded": generation_degraded,
        "retrieval_mode": evidence_result["retrieval_mode"],
        "retrieval_degraded": (
            evidence_result["retrieval_degraded"]
            or bool(broad_result and broad_result["retrieval_degraded"])
        ),
        "effective_category": effective_category.key if effective_category else None,
        "effective_category_name": (
            effective_category.name if effective_category else None
        ),
        "category_auto_inferred": inferred_category is not None,
        "restricted_source_count": restricted_source_count,
        "l3_authorized_for_request": allow_l3,
        "notice": (
            f"已生成可编辑初稿并保留 {len(generation_sources)} 条原页引用。"
            + (
                f" 已自动聚焦“{effective_category.name}”资料。"
                if inferred_category is not None
                else ""
            )
            + (
                f" 另有 {restricted_source_count} 条受限证据未发送至模型。"
                if restricted_source_count
                else ""
            )
        ),
    }
    latest = upsert_latest_writing_draft(
        db,
        user_id=user.id,
        instruction=payload.instruction,
        content=draft,
        category=payload.category,
        scope=payload.scope,
        generation_mode=generation_mode,
        generation_model=generation_model,
        sources=generation_sources,
        response=response_payload,
    )
    db.add(
        AuditLog(
            user_id=user.id,
            action="writing_draft_auto_saved",
            query_hash=instruction_hash,
            document_ids_json=json.dumps(outbound_document_ids),
            details_json=json.dumps(
                {
                    "draft_id": latest.id,
                    "content_characters": len(draft),
                    "generation_mode": generation_mode,
                    "model": generation_model,
                },
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    response_payload["auto_saved"] = True
    response_payload["latest_draft_id"] = latest.id
    return response_payload


def _writing_save_values(payload: WritingDraftSaveRequest) -> tuple[list[dict], dict]:
    response = dict(payload.response or {})
    response_sources = response.get("sources")
    sources = list(payload.sources)
    if not sources and isinstance(response_sources, list):
        sources = [item for item in response_sources if isinstance(item, dict)]
    response["draft"] = payload.draft
    response["sources"] = sources
    if payload.generation_mode is not None:
        response["generation_mode"] = payload.generation_mode
    if payload.generation_model is not None:
        response["generation_model"] = payload.generation_model
    return sources, response


@app.get("/v1/writing/drafts")
def writing_drafts_registry(
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    return list_writing_drafts(db, user_id=user.id)


@app.post("/v1/writing/drafts")
def save_writing_draft(
    payload: WritingDraftSaveRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    sources, response = _writing_save_values(payload)
    generation_mode = payload.generation_mode or response.get("generation_mode")
    generation_model = payload.generation_model or response.get("generation_model")
    try:
        record, saved_count, duplicate = save_writing_snapshot(
            db,
            user_id=user.id,
            title=payload.title,
            instruction=payload.instruction,
            content=payload.draft,
            category=payload.category,
            scope=payload.scope,
            time_scope=payload.time_scope,
            generation_mode=(str(generation_mode) if generation_mode else None),
            generation_model=(str(generation_model) if generation_model else None),
            sources=sources,
            response=response,
        )
    except OverflowError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    content_hash = hashlib.sha256(payload.draft.encode("utf-8")).hexdigest()
    db.add(
        AuditLog(
            user_id=user.id,
            action="writing_draft_saved",
            query_hash=content_hash,
            document_ids_json=json.dumps(
                list(
                    dict.fromkeys(
                        str(item.get("document_id"))
                        for item in sources
                        if item.get("document_id")
                    )
                )
            ),
            details_json=json.dumps(
                {
                    "draft_id": record.id,
                    "content_characters": len(payload.draft),
                    "saved_count": saved_count,
                    "idempotent_duplicate": duplicate,
                },
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    return {
        "draft": serialize_writing_draft(record),
        "saved_count": saved_count,
        "max_saved": MAX_SAVED_WRITING_DRAFTS,
        "duplicate": duplicate,
    }


@app.patch("/v1/writing/drafts/latest")
def update_latest_writing_draft(
    payload: WritingDraftSaveRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    existing = db.scalar(
        select(WritingDraft).where(
            WritingDraft.user_id == user.id,
            WritingDraft.kind == "latest",
        )
    )
    if existing is None:
        raise HTTPException(status_code=404, detail="尚无最近一次创作，请先成功生成一份初稿。")
    sources, response = _writing_save_values(payload)
    generation_mode = payload.generation_mode or response.get("generation_mode")
    generation_model = payload.generation_model or response.get("generation_model")
    record = upsert_latest_writing_draft(
        db,
        user_id=user.id,
        title=payload.title,
        instruction=payload.instruction,
        content=payload.draft,
        category=payload.category,
        scope=payload.scope,
        generation_mode=(str(generation_mode) if generation_mode else None),
        generation_model=(str(generation_model) if generation_model else None),
        sources=sources,
        response=response,
    )
    db.add(
        AuditLog(
            user_id=user.id,
            action="writing_latest_updated",
            query_hash=hashlib.sha256(payload.draft.encode("utf-8")).hexdigest(),
            document_ids_json=json.dumps([]),
            details_json=json.dumps(
                {"draft_id": record.id, "content_characters": len(payload.draft)},
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    return {"draft": serialize_writing_draft(record), "auto_saved": True}


@app.delete("/v1/writing/drafts/{draft_id}")
def delete_writing_draft(
    draft_id: str,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    record = db.get(WritingDraft, draft_id)
    if record is None or record.user_id != user.id:
        raise HTTPException(status_code=404, detail="未找到这条已保存创作。")
    if record.kind == "latest":
        raise HTTPException(
            status_code=409,
            detail="最近一次创作是自动保存槽，生成新内容时会自动覆盖，无需删除。",
        )
    db.delete(record)
    db.add(
        AuditLog(
            user_id=user.id,
            action="writing_draft_deleted",
            query_hash=record.content_hash,
            document_ids_json="[]",
            details_json=json.dumps({"draft_id": draft_id}, ensure_ascii=False),
        )
    )
    db.commit()
    saved_count = db.scalar(
        select(func.count(WritingDraft.id)).where(
            WritingDraft.user_id == user.id,
            WritingDraft.kind == "saved",
        )
    ) or 0
    return {
        "deleted": True,
        "draft_id": draft_id,
        "saved_count": int(saved_count),
        "max_saved": MAX_SAVED_WRITING_DRAFTS,
    }


@app.post("/v1/evolution/digest")
def evolution_digest(
    payload: EvolutionDigestRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    if user.role not in {"founder", "knowledge_admin"}:
        raise HTTPException(status_code=403, detail="无权生成知识进化摘要")
    if payload.reviewed_through <= payload.cutoff_date:
        raise HTTPException(
            status_code=422,
            detail="复核截止日必须晚于当前资料截止日",
        )
    if payload.reviewed_through > datetime.now(timezone.utc).date():
        raise HTTPException(
            status_code=422,
            detail="复核截止日不能晚于今天",
        )
    return create_evolution_digest(
        db,
        user=user,
        artifact=payload.artifact,
        cutoff_date=payload.cutoff_date,
        reviewed_through=payload.reviewed_through,
        audience=payload.audience,
        limit=payload.limit,
    )


@app.get("/v1/evolution/artifacts")
def evolution_artifacts(
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> list[dict]:
    return list_artifacts(db, user=user)


@app.post("/v1/evolution/artifacts")
def register_evolution_artifact(
    payload: MaintainedArtifactCreate,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    return create_artifact(
        db,
        user=user,
        name=payload.name,
        artifact_type=payload.artifact_type,
        current_document_id=payload.current_document_id,
        audience=payload.audience,
        cutoff_date=payload.cutoff_date,
        review_cadence_days=payload.review_cadence_days,
        section_map=payload.section_map,
    )


@app.get("/v1/evolution/baseline-documents")
def evolution_baseline_documents(
    artifact_type: str = Query(default="company_profile"),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> list[dict]:
    return list_baseline_documents(
        db,
        user=user,
        artifact_type=artifact_type,
    )


@app.get("/v1/evolution/artifacts/{artifact_id}")
def evolution_artifact_detail(
    artifact_id: str,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    return get_artifact(db, artifact_id=artifact_id, user=user)


@app.post("/v1/evolution/artifacts/{artifact_id}/runs")
def start_evolution_artifact_review(
    artifact_id: str,
    payload: EvolutionArtifactRunRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    return create_artifact_review(
        db,
        user=user,
        artifact_id=artifact_id,
        reviewed_through=(
            payload.reviewed_through
            or datetime.now(timezone.utc).date()
        ),
        limit=payload.limit,
    )


@app.post("/v1/evolution/artifacts/{artifact_id}/baseline")
def replace_evolution_artifact_baseline(
    artifact_id: str,
    payload: MaintainedArtifactBaselineUpdate,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    return update_artifact_baseline(
        db,
        user=user,
        artifact_id=artifact_id,
        current_document_id=payload.current_document_id,
        cutoff_date=payload.cutoff_date,
        confirmation=payload.confirmation,
    )


@app.get("/v1/evolution/runs/{run_id}")
def evolution_review_run_detail(
    run_id: str,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    return get_review_run(db, run_id=run_id, user=user)


@app.post("/v1/evolution/candidates/{candidate_id}/decision")
def evolution_candidate_decision(
    candidate_id: str,
    payload: EvolutionCandidateDecisionRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    return decide_candidate(
        db,
        user=user,
        candidate_id=candidate_id,
        decision=payload.decision,
        confirmation=payload.confirmation,
        note=payload.note,
    )


@app.post("/v1/evolution/runs/{run_id}/change-plan")
def evolution_change_plan(
    run_id: str,
    payload: EvolutionChangePlanRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    return create_change_plan(
        db,
        user=user,
        run_id=run_id,
        confirmation=payload.confirmation,
    )


@app.post("/v1/evolution/runs/{run_id}/close")
def complete_evolution_review(
    run_id: str,
    payload: EvolutionReviewCloseRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    return close_review(
        db,
        user=user,
        run_id=run_id,
        confirmation=payload.confirmation,
    )


@app.get("/v1/projects")
def list_projects(
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> list[dict]:
    projects = db.scalars(
        select(Project).options(joinedload(Project.documents)).order_by(Project.year.desc())
    ).unique().all()
    response: list[dict] = []
    for project in projects:
        if project.domain in CONTRACT_DOMAINS:
            continue
        visible = [
            document
            for document in project.documents
            if _document_visible_to_user(db, user, document, project)
        ]
        if not visible:
            continue
        has_closing_report = any(
            document.role == "closing_report"
            and document.knowledge_status in {"approved", "current"}
            and document.valid_to is None
            for document in visible
        )
        response.append(
            {
                "id": project.id,
                "name": project.name,
                "client": project.client,
                "year": project.year,
                "domain": project.domain,
                "status": project.status,
                "knowledge_status": project.knowledge_status,
                "confirmed": project.confirmed,
                "document_count": len(visible),
                "has_closing_report": has_closing_report,
                "closing_report_reminder": (
                    project.knowledge_status in {"approved", "current"}
                    and not has_closing_report
                ),
            }
        )
    return response


@app.get("/v1/projects/{project_id}")
def get_project(
    project_id: str,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    project = db.scalar(
        select(Project)
        .options(joinedload(Project.documents))
        .where(Project.id == project_id)
    )
    if not project:
        raise HTTPException(status_code=404, detail="项目不存在")
    if project.domain in CONTRACT_DOMAINS:
        raise HTTPException(status_code=404, detail="项目不存在")
    documents = [
        document
        for document in project.documents
        if _document_visible_to_user(db, user, document, project)
    ]
    if not documents:
        raise HTTPException(status_code=404, detail="项目不存在")
    has_closing_report = any(
        document.role == "closing_report"
        and document.knowledge_status in {"approved", "current"}
        and document.valid_to is None
        for document in documents
    )
    return {
        "id": project.id,
        "name": project.name,
        "client": project.client,
        "year": project.year,
        "domain": project.domain,
        "status": project.status,
        "knowledge_status": project.knowledge_status,
        "confirmed": project.confirmed,
        "has_closing_report": has_closing_report,
        "closing_report_reminder": (
            project.knowledge_status in {"approved", "current"}
            and not has_closing_report
        ),
        "documents": [
            {
                "id": document.id,
                "title": document.title,
                "role": document.role,
                "version": document.version,
                "is_final": document.is_final,
                "knowledge_status": document.knowledge_status,
                "confidentiality": document.confidentiality,
                "page_count": document.page_count,
                "citation_basis": document.citation_basis,
            }
            for document in documents
        ],
    }


REVIEWER_ROLES = {"founder", "knowledge_admin", "department_owner"}
EMPLOYEE_VISIBLE_STATUSES = {"approved", "current", "superseded", "archived"}
IMAGE_PREVIEW_MEDIA_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
}
VIDEO_PREVIEW_MEDIA_TYPES = {
    ".mp4": "video/mp4",
    ".webm": "video/webm",
    ".mov": "video/quicktime",
    ".m4v": "video/x-m4v",
}
TEXT_PREVIEW_SUFFIXES = {".txt", ".md", ".xls"}
SPREADSHEET_PREVIEW_SUFFIXES = {".xlsx", ".xlsm"}
SPREADSHEET_PREVIEW_MAX_ROWS = 500
SPREADSHEET_PREVIEW_MAX_COLUMNS = 80
PREVIEW_TICKET_TTL_SECONDS = 5 * 60
_preview_tickets: dict[str, tuple[str, str, int, float]] = {}
_preview_ticket_lock = Lock()


def _document_visible_to_user(
    db: Session,
    user: User,
    document: Document,
    project: Project,
) -> bool:
    contract_category = CONTRACT_CATEGORIES_BY_DOMAIN.get(project.domain)
    if contract_category is not None:
        organization_role = user.organization_role or "business"
        if contract_category.key == "executive_office":
            scope = _contract_scope_for_user(user, contract_category.key)
            if scope == "own":
                return bool(
                    document.id in _owned_contract_document_ids(db, user)
                    and document.knowledge_status in EMPLOYEE_VISIBLE_STATUSES
                )
            if scope != "all":
                return False
        if not is_authorized(user, document, project):
            return False
        # Contract retrieval is intentionally narrower than the account's
        # confidentiality ceiling. Administrative users also remain able to
        # preview every ceiling-permitted category because they are an
        # upload/review role; personnel and management receive only their
        # explicit archive scopes. Finance receives only its own uploaded
        # executive-office documents through the special branch above.
        allowed_categories = {
            "administrative": set(CONTRACT_CATEGORIES),
            "personnel": {"personnel"},
            "business": set(),
            "finance": set(),
            "management": set(CONTRACT_CATEGORIES),
        }
        if contract_category.key not in allowed_categories.get(
            organization_role,
            set(),
        ):
            return False
    elif not is_authorized(user, document, project):
        return False
    if user.role in REVIEWER_ROLES:
        return document.knowledge_status not in {"deleted", "quarantined", "rejected"}
    return document.knowledge_status in EMPLOYEE_VISIBLE_STATUSES


def _ready_reference_pdf(document: Document):
    return next(
        (
            item
            for item in document.artifacts
            if item.kind == "reference_pdf"
            and item.status == "ready"
            and item.path
            and Path(item.path).is_file()
        ),
        None,
    )


def _preview_kind(document: Document) -> str | None:
    if not document.file_blob:
        return None
    if _ready_reference_pdf(document):
        return "pdf"
    suffix = Path(document.file_blob.source_path).suffix.lower()
    if suffix == ".pdf":
        return "pdf"
    if suffix in IMAGE_PREVIEW_MEDIA_TYPES:
        return "image"
    if suffix in VIDEO_PREVIEW_MEDIA_TYPES:
        return "video"
    if suffix in SPREADSHEET_PREVIEW_SUFFIXES:
        return "spreadsheet"
    if suffix in TEXT_PREVIEW_SUFFIXES and document.chunks:
        return "text"
    return None


def _spreadsheet_cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    return str(value)


def _spreadsheet_preview_html(
    source: Path,
    document: Document,
    page: int,
    download_url: str | None = None,
) -> str:
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(
            status_code=409,
            detail="Excel 原件暂时无法读取，请下载原件查看",
        ) from exc

    try:
        worksheets = workbook.worksheets
        if page > len(worksheets):
            raise HTTPException(status_code=422, detail="工作表序号超出文件范围")
        worksheet = worksheets[page - 1]
        max_row = min(worksheet.max_row or 1, SPREADSHEET_PREVIEW_MAX_ROWS)
        max_column = min(
            worksheet.max_column or 1,
            SPREADSHEET_PREVIEW_MAX_COLUMNS,
        )
        rows = list(
            worksheet.iter_rows(
                min_row=1,
                max_row=max_row,
                min_col=1,
                max_col=max_column,
                values_only=True,
            )
        )
        sheet_count = len(worksheets)
        sheet_name = worksheet.title
        truncated = (
            (worksheet.max_row or 1) > SPREADSHEET_PREVIEW_MAX_ROWS
            or (worksheet.max_column or 1) > SPREADSHEET_PREVIEW_MAX_COLUMNS
        )
    finally:
        workbook.close()

    column_headers = "".join(
        f'<th class="column-heading">{get_column_letter(index)}</th>'
        for index in range(1, max_column + 1)
    )
    table_rows: list[str] = []
    for row_number, row in enumerate(rows, start=1):
        cells = "".join(
            f"<td>{html.escape(_spreadsheet_cell_text(value))}</td>"
            for value in row
        )
        table_rows.append(
            f'<tr><th class="row-heading">{row_number}</th>{cells}</tr>'
        )

    confidentiality = html.escape(document.confidentiality or "")
    project_name = html.escape(document.project.name if document.project else "")
    file_name = html.escape(document.title or source.name)
    safe_sheet_name = html.escape(sheet_name)
    download_link = (
        f'<a class="download" href="{html.escape(download_url, quote=True)}">下载原件</a>'
        if download_url
        else ""
    )
    truncation_notice = (
        '<div class="notice">该工作表内容较多，在线预览仅显示前 '
        f"{SPREADSHEET_PREVIEW_MAX_ROWS} 行、{SPREADSHEET_PREVIEW_MAX_COLUMNS} 列；"
        "完整内容请下载原件。</div>"
        if truncated
        else ""
    )
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width,initial-scale=1">
  <title>{file_name} · {safe_sheet_name}</title>
  <style>
    :root {{ color-scheme: light; font-family: Inter, "Microsoft YaHei", sans-serif; }}
    * {{ box-sizing: border-box; }}
    body {{ margin: 0; color: #182133; background: #f3f5f8; }}
    .header {{ position: sticky; top: 0; z-index: 5; padding: 16px 20px; color: #fff;
      background: #151b28; box-shadow: 0 2px 14px rgba(0,0,0,.16); }}
    .brand {{ color: #f3a21b; font-size: 13px; font-weight: 800; letter-spacing: .14em; }}
    .file-row {{ display: flex; align-items: center; gap: 12px; justify-content: space-between; margin-top: 7px; }}
    h1 {{ margin: 0; overflow: hidden; font-size: 20px; text-overflow: ellipsis; white-space: nowrap; }}
    .download {{ flex: 0 0 auto; padding: 9px 14px; color: #fff; text-decoration: none;
      font-weight: 700; border-radius: 8px; background: #e21d2b; }}
    .meta {{ margin-top: 7px; color: #cbd3e0; font-size: 13px; }}
    .meta span {{ margin-right: 14px; }}
    .notice {{ margin: 12px 16px 0; padding: 11px 14px; color: #805200;
      border: 1px solid #f4d39c; border-radius: 8px; background: #fff6e7; }}
    .sheet {{ margin: 16px; overflow: auto; max-height: calc(100vh - 150px);
      border: 1px solid #d8dee8; border-radius: 10px; background: #fff; }}
    table {{ border-collapse: separate; border-spacing: 0; min-width: 100%; width: max-content; }}
    th, td {{ min-width: 110px; max-width: 360px; padding: 9px 11px; overflow-wrap: anywhere;
      border-right: 1px solid #e2e6ed; border-bottom: 1px solid #e2e6ed; background: #fff;
      font-size: 14px; line-height: 1.45; vertical-align: top; }}
    .corner, .column-heading, .row-heading {{ color: #526173; font-weight: 700; text-align: center; background: #edf1f6; }}
    .corner {{ position: sticky; top: 0; left: 0; z-index: 4; min-width: 54px; }}
    .column-heading {{ position: sticky; top: 0; z-index: 2; min-width: 110px; }}
    .row-heading {{ position: sticky; left: 0; z-index: 1; min-width: 54px; width: 54px; }}
    tbody tr:nth-child(even) td {{ background: #fafbfc; }}
    @media (max-width: 640px) {{
      .header {{ padding: 13px 14px; }} h1 {{ font-size: 17px; }}
      .meta {{ line-height: 1.7; }} .sheet {{ margin: 10px; max-height: calc(100vh - 145px); }}
      th, td {{ min-width: 100px; padding: 8px; font-size: 13px; }}
    }}
  </style>
</head>
<body>
  <header class="header">
    <div class="brand">JAOS · EXCEL 原表预览</div>
    <div class="file-row"><h1>{file_name}</h1>{download_link}</div>
    <div class="meta"><span>{project_name}</span><span>{safe_sheet_name}</span><span>工作表 {page}/{sheet_count}</span><span>{confidentiality}</span></div>
  </header>
  {truncation_notice}
  <main class="sheet">
    <table aria-label="{safe_sheet_name}">
      <thead><tr><th class="corner"></th>{column_headers}</tr></thead>
      <tbody>{''.join(table_rows)}</tbody>
    </table>
  </main>
</body>
</html>"""


@app.get("/v1/documents/{document_id}")
def get_document(
    document_id: str,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    document = db.execute(
        select(Document)
        .options(
            joinedload(Document.project),
            joinedload(Document.chunks),
            joinedload(Document.file_blob),
            joinedload(Document.artifacts),
        )
        .where(Document.id == document_id)
    ).unique().scalar_one_or_none()
    if not document or not _document_visible_to_user(
        db,
        user,
        document,
        document.project,
    ):
        raise HTTPException(status_code=404, detail="文档不存在")
    health = db.get(SourceHealth, document.content_hash)
    return {
        "id": document.id,
        "project_id": document.project_id,
        "title": document.title,
        "role": document.role,
        "version": document.version,
        "is_final": document.is_final,
        "knowledge_status": document.knowledge_status,
        "confidentiality": document.confidentiality,
        "page_count": document.page_count,
        "citation_basis": document.citation_basis,
        "source_available": source_is_available(
            document.file_blob,
            health,
        ),
        "artifacts": [
            {
                "kind": artifact.kind,
                "status": artifact.status,
                "page_count": artifact.page_count,
                "error_code": artifact.error_code,
            }
            for artifact in document.artifacts
        ],
        "pages": [
            {
                "page": chunk.page,
                "section": chunk.section,
                "excerpt": chunk.text[:500],
            }
            for chunk in sorted(document.chunks, key=lambda item: item.chunk_index)
        ],
    }


@app.get("/v1/documents/{document_id}/preview")
def preview_document(
    document_id: str,
    page: int = Query(default=1, ge=1),
    preview_ticket: str | None = None,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> Response:
    document = db.execute(
        select(Document)
        .options(
            joinedload(Document.project),
            joinedload(Document.file_blob),
            joinedload(Document.artifacts),
            joinedload(Document.chunks),
        )
        .where(Document.id == document_id)
    ).unique().scalar_one_or_none()
    if not document or not _document_visible_to_user(
        db,
        user,
        document,
        document.project,
    ):
        raise HTTPException(status_code=404, detail="文档不存在")
    if document.page_count and page > document.page_count:
        raise HTTPException(status_code=422, detail="页码超出文档范围")
    if not document.file_blob:
        raise HTTPException(status_code=409, detail="原件记录缺失")
    source = Path(document.file_blob.source_path)
    health = verify_and_record(
        db,
        document.file_blob,
        commit=False,
    )
    if health.status != "verified":
        db.add(
            AuditLog(
                user_id=user.id,
                action="source_integrity_failed",
                document_ids_json=json.dumps([document.id]),
                details_json=json.dumps(
                    {"status": health.status, "error_code": health.error_code},
                    ensure_ascii=False,
                ),
            )
        )
        db.commit()
        raise HTTPException(
            status_code=409,
            detail="原件失联或内容已变化，无法核验引用",
        )

    artifact = _ready_reference_pdf(document)
    suffix = source.suffix.lower()
    if artifact:
        preview_path = Path(artifact.path)
        preview_kind = "reference_pdf"
        response: Response = FileResponse(
            preview_path,
            media_type="application/pdf",
            content_disposition_type="inline",
        )
    elif suffix == ".pdf":
        preview_path = source
        preview_kind = "source_pdf"
        response = FileResponse(
            preview_path,
            media_type="application/pdf",
            content_disposition_type="inline",
        )
    elif suffix in IMAGE_PREVIEW_MEDIA_TYPES:
        preview_kind = "source_image"
        response = FileResponse(
            source,
            media_type=IMAGE_PREVIEW_MEDIA_TYPES[suffix],
            content_disposition_type="inline",
        )
    elif suffix in VIDEO_PREVIEW_MEDIA_TYPES:
        preview_kind = "source_video"
        response = FileResponse(
            source,
            media_type=VIDEO_PREVIEW_MEDIA_TYPES[suffix],
            content_disposition_type="inline",
        )
    elif suffix in SPREADSHEET_PREVIEW_SUFFIXES:
        preview_kind = "spreadsheet_table"
        download_url = (
            f"/v1/previews/{preview_ticket}?download=true"
            if preview_ticket
            else None
        )
        response = Response(
            content=_spreadsheet_preview_html(
                source,
                document,
                page,
                download_url,
            ),
            media_type="text/html; charset=utf-8",
        )
    elif suffix in TEXT_PREVIEW_SUFFIXES and document.chunks:
        preview_kind = "extracted_text"
        extracted = "\n\n".join(
            chunk.text.strip()
            for chunk in sorted(
                document.chunks,
                key=lambda item: (item.page, item.chunk_index),
            )
            if chunk.text.strip()
        )
        response = Response(
            content=extracted,
            media_type="text/plain; charset=utf-8",
        )
    else:
        raise HTTPException(
            status_code=409,
            detail="该文件尚未生成可预览内容",
        )

    db.add(
        AuditLog(
            user_id=user.id,
            action="open_source",
            document_ids_json=json.dumps([document.id]),
            details_json=json.dumps(
                {"page": page, "preview_kind": preview_kind},
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    response.headers.update({
        "Cache-Control": "private, no-store",
        "X-Content-Type-Options": "nosniff",
    })
    if preview_kind == "extracted_text":
        response.headers["Content-Security-Policy"] = (
            "default-src 'none'; sandbox"
        )
    elif preview_kind == "spreadsheet_table":
        response.headers["Content-Security-Policy"] = (
            "default-src 'none'; style-src 'unsafe-inline'; "
            "sandbox allow-downloads; base-uri 'none'; form-action 'none'"
        )
        response.headers["Referrer-Policy"] = "no-referrer"
    return response


@app.post("/v1/documents/{document_id}/preview-ticket")
def create_preview_ticket(
    document_id: str,
    page: int = Query(default=1, ge=1),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    """Issue a short-lived URL that native PDF viewers can load with ranges."""
    document = db.execute(
        select(Document)
        .options(
            joinedload(Document.project),
            joinedload(Document.file_blob),
            joinedload(Document.artifacts),
            joinedload(Document.chunks),
        )
        .where(Document.id == document_id)
    ).unique().scalar_one_or_none()
    if not document or not _document_visible_to_user(
        db,
        user,
        document,
        document.project,
    ):
        raise HTTPException(status_code=404, detail="文档不存在")
    if document.page_count and page > document.page_count:
        raise HTTPException(status_code=422, detail="页码超出文档范围")
    if _preview_kind(document) is None:
        raise HTTPException(
            status_code=409,
            detail="该文件尚未生成可预览内容",
        )

    now = datetime.now(timezone.utc).timestamp()
    expires_at = now + PREVIEW_TICKET_TTL_SECONDS
    ticket = secrets.token_urlsafe(32)
    with _preview_ticket_lock:
        expired = [
            key
            for key, value in _preview_tickets.items()
            if value[3] <= now
        ]
        for key in expired:
            _preview_tickets.pop(key, None)
        _preview_tickets[ticket] = (
            user.id,
            document.id,
            page,
            expires_at,
        )
    return {
        "path": f"/v1/previews/{ticket}",
        "expires_in_seconds": PREVIEW_TICKET_TTL_SECONDS,
    }


@app.get("/v1/previews/{ticket}")
def preview_with_ticket(
    ticket: str,
    download: bool = Query(default=False),
    db: Session = Depends(get_db),
) -> Response:
    now = datetime.now(timezone.utc).timestamp()
    with _preview_ticket_lock:
        record = _preview_tickets.get(ticket)
        if record is not None and record[3] <= now:
            _preview_tickets.pop(ticket, None)
            record = None
    if record is None:
        raise HTTPException(status_code=404, detail="预览链接已失效，请重新打开")

    user = db.get(User, record[0])
    if user is None or not user.active:
        raise HTTPException(status_code=404, detail="预览链接已失效，请重新打开")
    if download:
        document = db.execute(
            select(Document)
            .options(
                joinedload(Document.project),
                joinedload(Document.file_blob),
            )
            .where(Document.id == record[1])
        ).unique().scalar_one_or_none()
        if (
            not document
            or not document.file_blob
            or not _document_visible_to_user(db, user, document, document.project)
        ):
            raise HTTPException(status_code=404, detail="预览链接已失效，请重新打开")
        health = verify_and_record(db, document.file_blob, commit=False)
        if health.status != "verified":
            raise HTTPException(status_code=409, detail="原件失联或内容已变化，暂时无法下载")
        source = Path(document.file_blob.source_path)
        db.add(
            AuditLog(
                user_id=user.id,
                action="download_source",
                document_ids_json=json.dumps([document.id]),
                details_json=json.dumps({"via": "preview_ticket"}, ensure_ascii=False),
            )
        )
        db.commit()
        response = FileResponse(
            source,
            filename=source.name,
            media_type="application/octet-stream",
            content_disposition_type="attachment",
        )
        response.headers.update({
            "Cache-Control": "private, no-store",
            "X-Content-Type-Options": "nosniff",
        })
        return response
    return preview_document(
        document_id=record[1],
        page=record[2],
        preview_ticket=ticket,
        user=user,
        db=db,
    )


REVIEW_FIELD_NAMES = [
    "project_name",
    "client",
    "year",
    "domain",
    "document_role",
    "version",
    "confidentiality",
    "is_final",
]

CONFIDENTIALITY_RANK = {"L1": 1, "L2": 2, "L3": 3, "L4": 4, "L5": 5}


def _lineage_text(value: str | None) -> str:
    return " ".join((value or "").casefold().split())


def _time_sort_key(value: datetime) -> float:
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.timestamp()


def _same_publication_lineage(
    candidate: Document,
    current: Document,
) -> bool:
    if candidate.role == "company_profile":
        return (
            current.role == "company_profile"
            and candidate.project.domain == current.project.domain
        )
    same_base = (
        candidate.role == current.role
        and candidate.project.domain == current.project.domain
        and _lineage_text(candidate.project.name)
        == _lineage_text(current.project.name)
        and _lineage_text(candidate.project.client)
        == _lineage_text(current.project.client)
    )
    if candidate.role == "asset":
        return same_base and _lineage_text(candidate.title) == _lineage_text(
            current.title
        )
    return same_base


def _matching_current_documents(
    candidate: Document,
    current_documents: list[Document],
) -> list[Document]:
    return [
        document
        for document in current_documents
        if document.id != candidate.id
        and document.knowledge_status == "current"
        and document.valid_to is None
        and _same_publication_lineage(candidate, document)
    ]


def _publication_blockers(
    document: Document,
    *,
    metadata_confirmed: bool,
    pending_metadata: bool,
    source_available: bool | None = None,
) -> list[str]:
    blockers: list[str] = []
    if pending_metadata:
        blockers.append("尚有待确认的元数据建议")
    if not metadata_confirmed:
        blockers.append("元数据尚未由创始人确认")
    if not document.project.confirmed:
        blockers.append("项目归属尚未确认")
    if not document.is_final:
        blockers.append("尚未标记为定稿")
    if source_available is None:
        source_available = source_is_available(document.file_blob)
    if not source_available:
        blockers.append("源文件失联")
    if document.role != "asset" and not document.chunks:
        blockers.append("尚无可引用的解析内容")
    if document.page_count < 1:
        blockers.append("页数或素材页记录无效")
    source_suffix = (
        Path(document.file_blob.source_path).suffix.lower()
        if document.file_blob
        else ""
    )
    if source_suffix == ".pdf":
        ocr_status = next(
            (
                artifact.status
                for artifact in document.artifacts
                if artifact.kind == "ocr_text"
            ),
            None,
        )
        if ocr_status == "failed":
            blockers.append("PDF 本地 OCR 质量检查失败")
        elif ocr_status != "ready":
            blockers.append("PDF 尚未完成逐页 OCR 质量检查")
    return blockers


def _require_review_role(user: User) -> None:
    if user.role not in REVIEWER_ROLES:
        raise HTTPException(status_code=403, detail="无权访问审核队列")


def _can_review_document(
    user: User,
    document: Document,
    project: Project,
) -> bool:
    """Apply reviewer scope before the document is returned or changed."""
    if not is_authorized(user, document, project):
        return False
    if user.role in {"founder", "knowledge_admin"}:
        return True
    if user.role != "department_owner":
        return False
    required = max(
        CONFIDENTIALITY_RANK.get(document.confidentiality, 99),
        CONFIDENTIALITY_RANK.get(project.confidentiality, 99),
    )
    return required <= 2


def _proposal_within_review_scope(user: User, proposal: ReviewProposal) -> bool:
    if user.role != "department_owner":
        return True
    return (
        CONFIDENTIALITY_RANK.get(proposal.confidentiality, 99) <= 2
        and proposal.document_role != "company_profile"
    )


def _review_document_or_404(
    document_id: str,
    user: User,
    db: Session,
    *,
    allowed_statuses: set[str] | None = None,
) -> Document:
    allowed_statuses = allowed_statuses or {"candidate"}
    document = db.scalar(
        select(Document)
        .options(
            joinedload(Document.project),
            joinedload(Document.file_blob),
        )
        .where(Document.id == document_id)
    )
    if (
        document is None
        or document.knowledge_status not in allowed_statuses
        or not _can_review_document(user, document, document.project)
    ):
        # Unauthorized callers receive the same response as a missing document.
        raise HTTPException(status_code=404, detail="待审核文档不存在")
    return document


def _review_proposal_payload(
    proposal: ReviewProposal,
    db: Session,
) -> dict:
    submitted_by = db.get(User, proposal.submitted_by_user_id)
    decided_by = (
        db.get(User, proposal.decided_by_user_id)
        if proposal.decided_by_user_id
        else None
    )
    return {
        "id": proposal.id,
        "document_id": proposal.document_id,
        "project_name": proposal.project_name,
        "client": proposal.client,
        "year": proposal.year,
        "domain": proposal.domain,
        "document_role": proposal.document_role,
        "version": proposal.version,
        "confidentiality": proposal.confidentiality,
        "is_final": proposal.is_final,
        "note": proposal.note,
        "decision_note": proposal.decision_note,
        "status": proposal.status,
        "submitted_by": (
            submitted_by.display_name if submitted_by else "账号已停用"
        ),
        "decided_by": decided_by.display_name if decided_by else None,
        "created_at": proposal.created_at,
        "decided_at": proposal.decided_at,
    }


GOVERNANCE_STATUSES = {"approved", "current", "superseded", "archived"}


def _require_governance_role(user: User) -> None:
    if user.role not in {"founder", "knowledge_admin"}:
        raise HTTPException(status_code=403, detail="无权进行资料密级治理")


def _recalculate_project_confidentiality(project: Project) -> None:
    active_levels = [
        document.confidentiality
        for document in project.documents
        if document.knowledge_status not in {"deleted", "quarantined"}
    ]
    project.confidentiality = max(
        active_levels or ["L1"],
        key=lambda level: CONFIDENTIALITY_RANK.get(level, 99),
    )


def _remove_disallowed_embeddings(
    db: Session,
    document: Document,
    target_level: str,
) -> None:
    if target_level not in {"L3", "L4", "L5"}:
        return
    if target_level == "L3" and settings.embedding_l3_enabled:
        return
    chunk_ids = select(Chunk.id).where(Chunk.document_id == document.id)
    db.execute(
        delete(ChunkEmbedding).where(ChunkEmbedding.chunk_id.in_(chunk_ids))
    )


def _governance_documents(user: User, db: Session) -> list[Document]:
    # Joining project siblings, chunks and artifacts in one query multiplies
    # each document into thousands of rows for large photo/video projects.
    # Fetch collections separately so memory grows with actual records.
    documents = db.scalars(
        select(Document)
        .join(Document.project)
        .options(
            joinedload(Document.project).selectinload(Project.documents),
            selectinload(Document.chunks),
            joinedload(Document.file_blob),
            selectinload(Document.artifacts),
        )
        .where(
            Document.knowledge_status.in_(sorted(GOVERNANCE_STATUSES)),
            Project.domain.not_in(CONTRACT_DOMAINS),
            Document.role != "contract",
        )
        .order_by(Document.ingested_at.desc())
    ).unique().all()
    return [
        document
        for document in documents
        if _can_review_document(user, document, document.project)
    ]


@app.get("/v1/governance/documents")
def governance_documents(
    q: str = Query(default="", max_length=200),
    domain: str = Query(default="", max_length=40),
    confidentiality: str = Query(default="", max_length=2),
    role: str = Query(default="", max_length=40),
    knowledge_status: str = Query(default="", max_length=24),
    manual_only: bool = False,
    offset: int = Query(default=0, ge=0),
    limit: int = Query(default=50, ge=1, le=100),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    _require_governance_role(user)
    documents = _governance_documents(user, db)
    counts = {
        level: sum(document.confidentiality == level for document in documents)
        for level in CONFIDENTIALITY_RANK
    }
    rows: list[tuple[Document, dict]] = []
    normalized_query = q.strip().casefold()
    for document in documents:
        suggestion = suggest_document_confidentiality(document).payload()
        needs_review = (
            suggestion["requires_manual_review"]
            or suggestion["confidence"] == "low"
            or suggestion["level"] != document.confidentiality
        )
        if normalized_query and normalized_query not in (
            f"{document.title} {document.project.name}".casefold()
        ):
            continue
        if domain and document.project.domain != domain:
            continue
        if confidentiality and document.confidentiality != confidentiality:
            continue
        if role and document.role != role:
            continue
        if knowledge_status and document.knowledge_status != knowledge_status:
            continue
        if manual_only and not needs_review:
            continue
        rows.append((document, {**suggestion, "needs_review": needs_review}))

    total = len(rows)
    page = rows[offset:offset + limit]
    source_health = {
        item.content_hash: item
        for item in db.scalars(select(SourceHealth)).all()
    }
    return {
        "items": [
            {
                "document_id": document.id,
                "title": document.title,
                "project": document.project.name,
                "domain": document.project.domain,
                "role": document.role,
                "version": document.version,
                "knowledge_status": document.knowledge_status,
                "confidentiality": document.confidentiality,
                "page_count": document.page_count,
                "preview_available": _preview_kind(document) is not None,
                "source_available": source_is_available(
                    document.file_blob,
                    source_health.get(document.content_hash),
                ),
                "ai_suggestion": suggestion,
            }
            for document, suggestion in page
        ],
        "total": total,
        "offset": offset,
        "limit": limit,
        "counts": counts,
        "manual_review_count": sum(
            suggestion["needs_review"] for _, suggestion in rows
        ),
    }


@app.post("/v1/governance/documents/confidentiality")
def update_governance_confidentiality(
    payload: GovernanceConfidentialityUpdateRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    _require_governance_role(user)
    document_ids = list(dict.fromkeys(payload.document_ids))
    if payload.confidentiality in {"L4", "L5"} and (
        user.role != "founder" or len(document_ids) != 1
    ):
        raise HTTPException(status_code=403, detail="L4/L5只能由创始人逐份调整")
    if CONFIDENTIALITY_RANK[payload.confidentiality] > CONFIDENTIALITY_RANK.get(
        user.confidentiality_ceiling,
        0,
    ):
        raise HTTPException(status_code=403, detail="目标密级超过账号权限")

    documents = db.scalars(
        select(Document)
        .options(
            joinedload(Document.project).selectinload(Project.documents),
            selectinload(Document.chunks),
        )
        .where(
            Document.id.in_(document_ids),
            Document.knowledge_status.in_(sorted(GOVERNANCE_STATUSES)),
        )
    ).unique().all()
    if len(documents) != len(document_ids) or any(
        not _can_review_document(user, document, document.project)
        for document in documents
    ):
        raise HTTPException(status_code=404, detail="资料不存在或无权调整")
    if any(_contract_policy_for_document(document) for document in documents):
        raise HTTPException(
            status_code=409,
            detail="合同档案密级由分类固定，不能在资料密级治理中调整",
        )
    if user.role != "founder" and any(
        document.confidentiality in {"L4", "L5"} for document in documents
    ):
        raise HTTPException(status_code=403, detail="L4/L5资料只能由创始人调整")

    changed: list[dict] = []
    projects: dict[str, Project] = {}
    for document in documents:
        previous = document.confidentiality
        projects[document.project_id] = document.project
        if previous == payload.confidentiality:
            continue
        document.confidentiality = payload.confidentiality
        _remove_disallowed_embeddings(db, document, payload.confidentiality)
        changed.append({
            "document_id": document.id,
            "from": previous,
            "to": payload.confidentiality,
        })
        db.add(
            AuditLog(
                user_id=user.id,
                action="confidentiality_changed",
                document_ids_json=json.dumps([document.id]),
                details_json=json.dumps(
                    {
                        "from": previous,
                        "to": payload.confidentiality,
                        "reason": payload.reason,
                        "source": "manual_governance",
                    },
                    ensure_ascii=False,
                ),
            )
        )
    for project in projects.values():
        _recalculate_project_confidentiality(project)
    db.commit()
    return {
        "changed_count": len(changed),
        "unchanged_count": len(document_ids) - len(changed),
        "changes": changed,
        "notice": f"已调整 {len(changed)} 份资料；原密级和原因已记录。",
    }


@app.post("/v1/governance/ai-classify")
def ai_classify_governance_documents(
    payload: GovernanceAIClassifyRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    require_founder(user)
    documents = _governance_documents(user, db)
    suggestions: list[tuple[Document, dict]] = [
        (document, suggest_document_confidentiality(document).payload())
        for document in documents
    ]
    manual = [
        {
            "document_id": document.id,
            "title": document.title,
            "current": document.confidentiality,
            "suggested": suggestion["level"],
            "reason": suggestion["reason"],
        }
        for document, suggestion in suggestions
        if document.role == "contract"
        or document.confidentiality in {"L4", "L5"}
        or suggestion["level"] == "L4"
        or suggestion["requires_manual_review"]
        or suggestion["confidence"] == "low"
    ]
    proposed = [
        (document, suggestion)
        for document, suggestion in suggestions
        if document.role != "contract"
        and document.confidentiality not in {"L4", "L5"}
        and suggestion["level"] in {"L1", "L2", "L3"}
        and suggestion["level"] != document.confidentiality
        and not suggestion["requires_manual_review"]
        and suggestion["confidence"] != "low"
    ]
    proposed_counts = {
        level: sum(suggestion["level"] == level for _, suggestion in proposed)
        for level in ("L1", "L2", "L3")
    }
    if payload.dry_run:
        return {
            "dry_run": True,
            "scanned_count": len(documents),
            "proposed_change_count": len(proposed),
            "proposed_counts": proposed_counts,
            "manual_review_count": len(manual),
            "manual_review_samples": manual[:20],
        }

    projects: dict[str, Project] = {}
    for document, suggestion in proposed:
        previous = document.confidentiality
        document.confidentiality = suggestion["level"]
        projects[document.project_id] = document.project
        _remove_disallowed_embeddings(db, document, suggestion["level"])
        db.add(
            AuditLog(
                user_id=user.id,
                action="ai_confidentiality_applied",
                document_ids_json=json.dumps([document.id]),
                details_json=json.dumps(
                    {
                        "from": previous,
                        "to": suggestion["level"],
                        "confidence": suggestion["confidence"],
                        "reason": suggestion["reason"],
                        "source": "local_policy_classifier",
                    },
                    ensure_ascii=False,
                ),
            )
        )
    for project in projects.values():
        _recalculate_project_confidentiality(project)
    db.commit()
    return {
        "dry_run": False,
        "scanned_count": len(documents),
        "changed_count": len(proposed),
        "changed_counts": proposed_counts,
        "manual_review_count": len(manual),
        "manual_review_samples": manual[:20],
        "notice": "AI本地初筛已应用；合同、L4/L5建议和低置信度资料仍需人工逐份复核。",
    }


def _publication_document_payload(document: Document) -> dict:
    return {
        "document_id": document.id,
        "title": document.title,
        "project": document.project.name,
        "version": document.version,
        "valid_from": document.valid_from,
    }


def _review_uploader_payloads(
    db: Session,
    documents: list[Document],
) -> dict[str, dict[str, str | None]]:
    """Resolve original Web uploaders without adding an N+1 query to review."""
    if not documents:
        return {}

    document_ids = [document.id for document in documents]
    documents_by_id = {document.id: document for document in documents}
    target_ids = set(document_ids)
    uploader_ids: dict[str, str] = {}
    upload_logs = db.scalars(
        select(AuditLog)
        .where(
            AuditLog.action.in_(["web_upload", "contract_upload"]),
            AuditLog.user_id.is_not(None),
        )
        .order_by(AuditLog.created_at.asc())
    ).all()
    for audit in upload_logs:
        try:
            logged_document_ids = json.loads(audit.document_ids_json or "[]")
        except (TypeError, ValueError):
            continue
        if not isinstance(logged_document_ids, list) or not audit.user_id:
            continue
        for document_id in logged_document_ids:
            if document_id in target_ids and document_id not in uploader_ids:
                uploader_ids[document_id] = audit.user_id

    # Project proposal/closing files are first written by JAOS and discovered by
    # the NAS scanner afterwards, so they do not have a web_upload audit row.
    # Preserve the project creator as the human uploader instead of displaying
    # the container-owned NAS file as "root".
    source_paths = {
        document.file_blob.source_path
        for document in documents
        if document.file_blob and document.file_blob.source_path
    }
    if source_paths:
        managed_projects = db.scalars(
            select(ManagedProject).where(
                or_(
                    ManagedProject.proposal_path.in_(source_paths),
                    ManagedProject.closing_report_path.in_(source_paths),
                )
            )
        ).all()
        project_uploaders_by_path: dict[str, str] = {}
        for project in managed_projects:
            if project.proposal_path:
                project_uploaders_by_path[project.proposal_path] = (
                    project.created_by_user_id
                )
            if project.closing_report_path:
                project_uploaders_by_path[project.closing_report_path] = (
                    project.manager_user_id
                )
        for document in documents:
            if document.id in uploader_ids or not document.file_blob:
                continue
            project_uploader_id = project_uploaders_by_path.get(
                document.file_blob.source_path
            )
            if project_uploader_id:
                uploader_ids[document.id] = project_uploader_id

    user_ids = set(uploader_ids.values())
    users_by_id = {
        item.id: item
        for item in db.scalars(select(User).where(User.id.in_(user_ids))).all()
    } if user_ids else {}

    payloads: dict[str, dict[str, str | None]] = {}
    for document_id in document_ids:
        uploader_id = uploader_ids.get(document_id)
        uploader = users_by_id.get(uploader_id) if uploader_id else None
        if uploader:
            payloads[document_id] = {
                "uploader_name": uploader.display_name,
                "uploader_username": uploader.username,
                "upload_source": "web",
            }
        elif uploader_id:
            payloads[document_id] = {
                "uploader_name": "原上传账号已停用",
                "uploader_username": None,
                "upload_source": "web",
            }
        else:
            source_owner = (
                documents_by_id[document_id].file_blob.source_owner_name
                if documents_by_id[document_id].file_blob
                else None
            )
            payloads[document_id] = {
                "uploader_name": source_owner or "NAS用户未识别",
                "uploader_username": source_owner,
                "upload_source": "nas",
            }
    return payloads


@app.get("/v1/review/queue")
def review_queue(
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> list[dict]:
    _require_review_role(user)
    allowed_levels = [
        level
        for level, rank in CONFIDENTIALITY_RANK.items()
        if rank <= CONFIDENTIALITY_RANK.get(
            user.confidentiality_ceiling,
            0,
        )
    ]
    if user.role == "department_owner":
        allowed_levels = [level for level in allowed_levels if level in {"L1", "L2"}]
    document_query = (
        select(Document)
        .join(Document.project)
        .options(
            joinedload(Document.project),
            joinedload(Document.chunks),
            joinedload(Document.file_blob),
            joinedload(Document.artifacts),
        )
        .where(
            Document.knowledge_status.in_(["candidate", "approved"]),
            Document.confidentiality.in_(allowed_levels),
            Project.confidentiality.in_(allowed_levels),
        )
        .order_by(Document.ingested_at.desc())
    )
    documents = db.scalars(document_query).unique().all()
    executive_scope = _contract_scope_for_user(user, "executive_office")
    if executive_scope == "own":
        owned_contract_ids = _owned_contract_document_ids(db, user)
        documents = [
            document
            for document in documents
            if document.project.domain != CONTRACT_CATEGORIES["executive_office"].domain
            or document.id in owned_contract_ids
        ]
    document_ids = [document.id for document in documents]
    uploaders_by_document = _review_uploader_payloads(db, documents)
    proposals = (
        db.scalars(
            select(ReviewProposal)
            .where(
                ReviewProposal.document_id.in_(document_ids),
                ReviewProposal.status.in_(["pending_founder", "applied"]),
            )
            .order_by(ReviewProposal.created_at.desc())
        ).all()
        if document_ids
        else []
    )
    pending_by_document: dict[str, ReviewProposal] = {}
    applied_document_ids: set[str] = set()
    for proposal in proposals:
        if proposal.status == "pending_founder":
            pending_by_document.setdefault(proposal.document_id, proposal)
        elif proposal.status == "applied":
            applied_document_ids.add(proposal.document_id)

    current_documents = (
        db.scalars(
            select(Document)
            .join(Document.project)
            .options(joinedload(Document.project))
            .where(
                Document.knowledge_status == "current",
                Document.valid_to.is_(None),
                Document.confidentiality.in_(allowed_levels),
                Project.confidentiality.in_(allowed_levels),
            )
        ).unique().all()
        if user.role == "founder"
        else []
    )
    source_health = {
        row.content_hash: row
        for row in db.scalars(select(SourceHealth)).all()
    }

    response: list[dict] = []
    for document in documents:
        if not _can_review_document(user, document, document.project):
            continue
        # Confirmed company introductions stay visible to the founder until
        # they are explicitly promoted to the current company-wide version.
        if document.knowledge_status == "approved" and not (
            user.role == "founder" and document.role == "company_profile"
        ):
            continue
        source_available = source_is_available(
            document.file_blob,
            source_health.get(document.content_hash),
        )
        artifact_status = {item.kind: item.status for item in document.artifacts}
        preview_kind = _preview_kind(document) if source_available else None
        preview_available = preview_kind is not None
        issue = None
        if not source_available:
            issue = "源文件失联"
        elif (
            Path(document.file_blob.source_path).suffix.lower() == ".pdf"
            and artifact_status.get("ocr_text") == "failed"
        ):
            issue = "PDF 本地 OCR 质量检查失败"
        elif (
            Path(document.file_blob.source_path).suffix.lower() == ".pdf"
            and artifact_status.get("ocr_text") != "ready"
        ):
            issue = "等待 PDF 逐页 OCR 质量检查"
        elif (
            document.role != "asset"
            and document.page_count
            and not document.chunks
        ):
            issue = "需要 OCR"
        elif (
            Path(document.file_blob.source_path).suffix.lower() == ".docx"
            and artifact_status.get("reference_pdf") != "ready"
        ):
            issue = "DOCX需生成引用版PDF"
        pending = pending_by_document.get(document.id)
        metadata_confirmed = (
            document.knowledge_status == "approved"
            or (
                document.id in applied_document_ids
                and pending is None
                and document.project.confirmed
            )
        )
        publication_blockers = _publication_blockers(
            document,
            metadata_confirmed=metadata_confirmed,
            pending_metadata=pending is not None,
            source_available=source_available,
        )
        replacement_candidates = _matching_current_documents(
            document,
            current_documents,
        )
        if document.knowledge_status == "approved":
            review_state = "已确认入库，待发布当前版本"
        elif pending:
            review_state = "信息已修改，待确认"
        elif document.id in applied_document_ids:
            review_state = "待确认入库"
        else:
            review_state = "待预览确认"
        response.append({
            "document_id": document.id,
            "title": document.title,
            "project_id": document.project_id,
            "project": document.project.name,
            "client": document.project.client,
            "year": document.project.year,
            "domain": document.project.domain,
            "role": document.role,
            "version": document.version,
            "is_final": document.is_final,
            "knowledge_status": document.knowledge_status,
            "confidentiality": document.confidentiality,
            "page_count": document.page_count,
            "extracted_chunk_count": len(document.chunks),
            "citation_basis": document.citation_basis,
            "source_available": source_available,
            "preview_available": preview_available,
            "preview_kind": preview_kind,
            "issue": issue,
            "review_state": review_state,
            **uploaders_by_document[document.id],
            "pending_proposal": (
                _review_proposal_payload(pending, db) if pending else None
            ),
            "metadata_confirmed": metadata_confirmed,
            "confirm_eligible": (
                document.knowledge_status == "candidate"
                and preview_available
                and issue is None
            ),
            "confirm_blockers": (
                []
                if preview_available and issue is None
                else [issue or "资料尚未生成可预览内容"]
            ),
            "publication_eligible": (
                user.role == "founder" and not publication_blockers
            ),
            "publication_blockers": publication_blockers,
            "replacement_candidates": [
                _publication_document_payload(item)
                for item in replacement_candidates
            ],
        })
    return response


@app.post("/v1/review/{document_id}/confirm")
def confirm_review_document(
    document_id: str,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    """Confirm a previewed candidate into the historical knowledge base."""
    _require_review_role(user)
    document = db.execute(
        select(Document)
        .options(
            joinedload(Document.project).selectinload(Project.documents),
            joinedload(Document.file_blob),
            selectinload(Document.artifacts),
            selectinload(Document.chunks),
        )
        .where(Document.id == document_id)
    ).unique().scalar_one_or_none()
    if (
        document is None
        or document.knowledge_status != "candidate"
        or not _can_review_document(user, document, document.project)
    ):
        raise HTTPException(status_code=404, detail="待审核文档不存在")
    if not document.file_blob:
        raise HTTPException(status_code=409, detail="原件记录缺失，不能确认入库")

    contract_policy = _contract_policy_for_document(document)
    if contract_policy is not None:
        # Contract metadata is derived from its fixed NAS directory. Never
        # accept a staged proposal that could move it into a normal category
        # or lower its confidentiality during review.
        document.role = "contract"
        document.confidentiality = contract_policy.confidentiality
        document.project.confidentiality = contract_policy.confidentiality

    health = verify_and_record(db, document.file_blob, commit=False)
    if health.status != "verified":
        db.add(
            AuditLog(
                user_id=user.id,
                action="source_integrity_failed",
                document_ids_json=json.dumps([document.id]),
                details_json=json.dumps(
                    {"status": health.status, "error_code": health.error_code},
                    ensure_ascii=False,
                ),
            )
        )
        db.commit()
        raise HTTPException(
            status_code=409,
            detail="原件失联或内容已变化，请恢复原件后再确认",
        )
    if _preview_kind(document) is None:
        raise HTTPException(
            status_code=409,
            detail="资料尚未生成可预览内容，不能确认入库",
        )

    proposal = db.scalar(
        select(ReviewProposal)
        .where(
            ReviewProposal.document_id == document.id,
            ReviewProposal.status == "pending_founder",
        )
        .order_by(ReviewProposal.created_at.desc())
    )
    now = datetime.now(timezone.utc)
    applied_proposal_id = None
    if proposal is not None:
        if contract_policy is not None:
            raise HTTPException(
                status_code=409,
                detail="合同档案分类和密级由固定目录决定，不能应用普通资料修改建议",
            )
        if not _proposal_within_review_scope(user, proposal):
            raise HTTPException(
                status_code=403,
                detail="修改后的资料归属或密级超出业务负责人权限",
            )
        document.project.name = proposal.project_name
        document.project.client = proposal.client
        document.project.year = proposal.year
        document.project.domain = proposal.domain
        document.role = proposal.document_role
        document.version = proposal.version
        document.confidentiality = proposal.confidentiality
        document.is_final = proposal.is_final
        proposal.status = "applied"
        proposal.decided_by_user_id = user.id
        proposal.decided_at = now
        proposal.updated_at = now
        applied_proposal_id = proposal.id

    document.project.confirmed = True
    document.project.confidentiality = max(
        [
            document.project.confidentiality,
            *(item.confidentiality for item in document.project.documents),
        ],
        key=lambda level: CONFIDENTIALITY_RANK.get(level, 99),
    )
    document.knowledge_status = "approved"
    if document.project.knowledge_status != "current":
        document.project.knowledge_status = "approved"

    db.add(
        AuditLog(
            user_id=user.id,
            action="review_confirmed",
            document_ids_json=json.dumps([document.id]),
            details_json=json.dumps(
                {
                    "previous_status": "candidate",
                    "knowledge_status": "approved",
                    "applied_proposal_id": applied_proposal_id,
                    "source_verified": True,
                    "confirmed_by_role": user.role,
                },
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    return {
        "document_id": document.id,
        "project_id": document.project.id,
        "knowledge_status": "approved",
        "metadata_confirmed": True,
        "notice": "已确认入库，可在历史知识中检索使用。",
    }


@app.post("/v1/review/confirm-batch")
def confirm_review_documents_batch(
    payload: BatchReviewConfirmRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    """Confirm authorized L1-L3 files with idempotent retry behavior."""
    _require_review_role(user)
    confirmed: list[str] = []
    already_confirmed: list[str] = []
    failed: list[dict[str, str]] = []
    for document_id in dict.fromkeys(payload.document_ids):
        document = db.execute(
            select(Document)
            .options(joinedload(Document.project))
            .where(Document.id == document_id)
        ).unique().scalar_one_or_none()
        if document is None or not _can_review_document(
            user,
            document,
            document.project,
        ):
            failed.append({"document_id": document_id, "detail": "资料不存在或无权确认"})
            continue
        required = max(
            CONFIDENTIALITY_RANK.get(document.confidentiality, 99),
            CONFIDENTIALITY_RANK.get(document.project.confidentiality, 99),
        )
        if required > 3:
            failed.append(
                {
                    "document_id": document_id,
                    "detail": "L4/L5 高敏资料必须逐份确认",
                }
            )
            continue
        # A browser tab may retry after the first request has succeeded but
        # before its refreshed queue is visible.  Treat an authorized retry as
        # success without writing a duplicate confirmation audit event.
        if document.knowledge_status in {"approved", "current"}:
            already_confirmed.append(document_id)
            continue
        try:
            confirm_review_document(document_id, user=user, db=db)
            confirmed.append(document_id)
        except HTTPException as exc:
            db.rollback()
            failed.append({"document_id": document_id, "detail": str(exc.detail)})

    accepted = [*confirmed, *already_confirmed]
    return {
        "confirmed_count": len(accepted),
        "newly_confirmed_count": len(confirmed),
        "already_confirmed_count": len(already_confirmed),
        "failed_count": len(failed),
        "confirmed_document_ids": accepted,
        "newly_confirmed_document_ids": confirmed,
        "already_confirmed_document_ids": already_confirmed,
        "failed": failed,
        "notice": f"已确认 {len(accepted)} 份资料",
    }


def _reject_ingestion_candidate(
    document_id: str,
    *,
    user: User,
    db: Session,
    reason: str | None,
) -> dict:
    document = _review_document_or_404(document_id, user, db)
    project = document.project
    now = datetime.now(timezone.utc)
    decision_note = (reason or "审核人拒绝入库").strip()
    pending_proposals = db.scalars(
        select(ReviewProposal).where(
            ReviewProposal.document_id == document.id,
            ReviewProposal.status == "pending_founder",
        )
    ).all()
    for proposal in pending_proposals:
        proposal.status = "rejected"
        proposal.decision_note = decision_note
        proposal.decided_by_user_id = user.id
        proposal.decided_at = now
        proposal.updated_at = now

    document.knowledge_status = "rejected"
    db.flush()
    remaining_active = db.scalar(
        select(func.count())
        .select_from(Document)
        .where(
            Document.project_id == project.id,
            Document.knowledge_status.in_(["candidate", "approved", "current"]),
        )
    ) or 0
    if remaining_active == 0 and project.knowledge_status != "current":
        project.knowledge_status = "rejected"
        project.confirmed = False

    db.add(
        AuditLog(
            user_id=user.id,
            action="ingestion_rejected",
            document_ids_json=json.dumps([document.id]),
            details_json=json.dumps(
                {
                    "previous_status": "candidate",
                    "knowledge_status": "rejected",
                    "reason": decision_note,
                    "source_retained": True,
                    "rejected_by_role": user.role,
                },
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    return {
        "document_id": document.id,
        "knowledge_status": "rejected",
        "source_retained": True,
        "notice": "已拒绝入库；NAS 原文件保留，但不会进入智库检索。",
    }


@app.post("/v1/review/{document_id}/decline")
def reject_ingestion_candidate(
    document_id: str,
    payload: ReviewIngestionRejectRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    _require_review_role(user)
    return _reject_ingestion_candidate(
        document_id,
        user=user,
        db=db,
        reason=payload.reason,
    )


@app.post("/v1/review/reject-batch")
def reject_ingestion_candidates_batch(
    payload: BatchReviewRejectRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    """Reject authorized L1-L3 candidates while retaining their NAS originals."""
    _require_review_role(user)
    rejected: list[str] = []
    already_rejected: list[str] = []
    failed: list[dict[str, str]] = []
    for document_id in dict.fromkeys(payload.document_ids):
        document = db.execute(
            select(Document)
            .options(joinedload(Document.project))
            .where(Document.id == document_id)
        ).unique().scalar_one_or_none()
        if document is None or not _can_review_document(
            user,
            document,
            document.project,
        ):
            failed.append({"document_id": document_id, "detail": "资料不存在或无权拒绝"})
            continue
        required = max(
            CONFIDENTIALITY_RANK.get(document.confidentiality, 99),
            CONFIDENTIALITY_RANK.get(document.project.confidentiality, 99),
        )
        if required > 3:
            failed.append(
                {
                    "document_id": document_id,
                    "detail": "L4/L5 高敏资料必须逐份处理",
                }
            )
            continue
        if document.knowledge_status == "rejected":
            already_rejected.append(document_id)
            continue
        if document.knowledge_status != "candidate":
            failed.append({"document_id": document_id, "detail": "资料已不在待审核状态"})
            continue
        try:
            _reject_ingestion_candidate(
                document_id,
                user=user,
                db=db,
                reason=payload.reason,
            )
            rejected.append(document_id)
        except HTTPException as exc:
            db.rollback()
            failed.append({"document_id": document_id, "detail": str(exc.detail)})

    accepted = [*rejected, *already_rejected]
    return {
        "rejected_count": len(accepted),
        "newly_rejected_count": len(rejected),
        "already_rejected_count": len(already_rejected),
        "failed_count": len(failed),
        "rejected_document_ids": accepted,
        "failed": failed,
        "notice": f"已拒绝 {len(accepted)} 份资料；NAS 原文件均保留。",
    }


@app.get("/v1/review/inbox/issues")
def inbox_issues(
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> list[dict]:
    require_founder(user)
    issues = db.scalars(
        select(InboxIssue)
        .where(InboxIssue.resolved_at.is_(None))
        .order_by(InboxIssue.last_seen_at.desc())
    ).all()
    return [
        {
            "id": issue.id,
            "relative_path": issue.relative_path,
            "status": issue.status,
            "error_code": issue.error_code,
            "message": issue.message,
            "size_bytes": issue.size_bytes,
            "last_seen_at": issue.last_seen_at,
        }
        for issue in issues
    ]


@app.post("/v1/review/inbox/issues/{issue_id}/ignore")
def ignore_inbox_issue(
    issue_id: str,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    """Acknowledge a non-searchable NAS file without deleting its source."""
    require_founder(user)
    issue = db.get(InboxIssue, issue_id)
    if issue is None or issue.resolved_at is not None:
        raise HTTPException(status_code=404, detail="待处理目录问题不存在")
    now = datetime.now(timezone.utc)
    issue.status = "ignored"
    issue.resolved_at = now
    issue.last_seen_at = now
    db.add(
        AuditLog(
            user_id=user.id,
            action="inbox_issue_ignored",
            details_json=json.dumps(
                {
                    "issue_id": issue.id,
                    "relative_path": issue.relative_path,
                    "error_code": issue.error_code,
                    "source_deleted": False,
                },
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    return {
        "status": "ignored",
        "issue_id": issue.id,
        "notice": "已忽略该文件；NAS 原文件保留，文件内容变化后会重新检查。",
    }


@app.post("/v1/review/inbox/scan")
def trigger_inbox_scan(
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    if user.role not in {"founder", "knowledge_admin"}:
        raise HTTPException(status_code=403, detail="无权扫描全局待审核目录")
    # Scan the complete managed library so files copied directly through fnOS/SMB
    # are registered as well as files uploaded through the Web intake flow.
    result = scan_inbox(db, settings.knowledge_root)
    db.add(
        AuditLog(
            user_id=user.id,
            action="manual_inbox_scan",
            details_json=json.dumps(
                {
                    "status": result["status"],
                    "counts": result["counts"],
                },
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    return {
        "status": result["status"],
        "counts": result["counts"],
        "duration_ms": result.get("duration_ms"),
        "checked_at": result["checked_at"],
    }


@app.post("/v1/review/{document_id}/proposal")
def submit_review_proposal(
    document_id: str,
    payload: ReviewProposalRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    _require_review_role(user)
    document = _review_document_or_404(document_id, user, db)
    if _contract_policy_for_document(document) is not None:
        raise HTTPException(
            status_code=409,
            detail="合同档案分类和密级由固定目录决定，无需修改普通资料元数据",
        )
    if not db.scalar(
        select(KnowledgeCategory.id).where(
            KnowledgeCategory.key == payload.domain,
            KnowledgeCategory.active.is_(True),
        )
    ):
        raise HTTPException(status_code=422, detail="资料分类不存在或已停用")
    if user.role == "department_owner" and (
        CONFIDENTIALITY_RANK.get(payload.confidentiality, 99) > 2
        or payload.document_role == "company_profile"
    ):
        raise HTTPException(
            status_code=403,
            detail="修改后的资料归属或密级超出业务负责人权限",
        )
    now = datetime.now(timezone.utc)
    previous = db.scalars(
        select(ReviewProposal).where(
            ReviewProposal.document_id == document.id,
            ReviewProposal.status == "pending_founder",
        )
    ).all()
    for proposal in previous:
        proposal.status = "superseded"
        proposal.decided_by_user_id = user.id
        proposal.decided_at = now
        proposal.updated_at = now

    proposal = ReviewProposal(
        document_id=document.id,
        project_name=payload.project_name.strip(),
        client=payload.client.strip() if payload.client else None,
        year=payload.year,
        domain=payload.domain,
        document_role=payload.document_role,
        version=payload.version.strip(),
        confidentiality=payload.confidentiality,
        is_final=payload.is_final,
        note=payload.note.strip() if payload.note else None,
        status="pending_founder",
        submitted_by_user_id=user.id,
    )
    db.add(proposal)
    db.flush()
    db.add(
        AuditLog(
            user_id=user.id,
            action="review_proposed",
            document_ids_json=json.dumps([document.id]),
            details_json=json.dumps(
                {
                    "proposal_id": proposal.id,
                    "field_names": REVIEW_FIELD_NAMES,
                    "superseded_count": len(previous),
                },
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    return {
        "proposal": _review_proposal_payload(proposal, db),
        "document_unchanged": True,
        "knowledge_status": document.knowledge_status,
        "notice": "建议已提交；创始人确认前不会修改资料元数据。",
    }


@app.post("/v1/review/{document_id}/apply")
def apply_review_proposal(
    document_id: str,
    payload: ReviewApplyRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    require_founder(user)
    document = _review_document_or_404(document_id, user, db)
    if _contract_policy_for_document(document) is not None:
        raise HTTPException(
            status_code=409,
            detail="合同档案分类和密级由固定目录决定，不能应用普通资料修改建议",
        )
    proposal = db.scalar(
        select(ReviewProposal).where(
            ReviewProposal.id == payload.proposal_id,
            ReviewProposal.document_id == document.id,
            ReviewProposal.status == "pending_founder",
        )
    )
    if proposal is None:
        raise HTTPException(status_code=409, detail="确认建议已失效或已处理")

    document.project.name = proposal.project_name
    document.project.client = proposal.client
    document.project.year = proposal.year
    document.project.domain = proposal.domain
    document.project.confirmed = True
    document.role = proposal.document_role
    document.version = proposal.version
    document.confidentiality = proposal.confidentiality
    document.is_final = proposal.is_final
    document.project.confidentiality = max(
        [
            document.project.confidentiality,
            *(item.confidentiality for item in document.project.documents),
        ],
        key=lambda level: CONFIDENTIALITY_RANK.get(level, 99),
    )
    # Metadata confirmation is deliberately separate from truth promotion.
    document.knowledge_status = "candidate"
    document.project.knowledge_status = "candidate"

    now = datetime.now(timezone.utc)
    proposal.status = "applied"
    proposal.decided_by_user_id = user.id
    proposal.decided_at = now
    proposal.updated_at = now
    db.add(
        AuditLog(
            user_id=user.id,
            action="review_applied",
            document_ids_json=json.dumps([document.id]),
            details_json=json.dumps(
                {
                    "proposal_id": proposal.id,
                    "field_names": REVIEW_FIELD_NAMES,
                    "knowledge_status": "candidate",
                    "current_promotion": False,
                },
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    return {
        "proposal": _review_proposal_payload(proposal, db),
        "document_id": document.id,
        "project_id": document.project.id,
        "metadata_confirmed": True,
        "knowledge_status": "candidate",
        "current_promotion": False,
        "notice": "元数据已确认；资料仍为候选，不代表公司当前事实。",
    }


@app.post("/v1/review/{document_id}/publish")
def publish_candidate_document(
    document_id: str,
    payload: PublicationRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    require_founder(user)
    document = _review_document_or_404(
        document_id,
        user,
        db,
        allowed_statuses={"candidate", "approved"},
    )
    proposal_statuses = db.scalars(
        select(ReviewProposal.status).where(
            ReviewProposal.document_id == document.id,
            ReviewProposal.status.in_(["pending_founder", "applied"]),
        )
    ).all()
    pending_metadata = "pending_founder" in proposal_statuses
    metadata_confirmed = (
        document.knowledge_status == "approved"
        or (
            "applied" in proposal_statuses
            and not pending_metadata
            and document.project.confirmed
        )
    )
    source_verified = False
    if document.file_blob:
        source_verified = (
            verify_and_record(
                db,
                document.file_blob,
                commit=False,
            ).status
            == "verified"
        )
    blockers = _publication_blockers(
        document,
        metadata_confirmed=metadata_confirmed,
        pending_metadata=pending_metadata,
        source_available=source_verified,
    )
    if blockers:
        raise HTTPException(
            status_code=409,
            detail=f"暂不能发布：{'；'.join(blockers)}",
        )

    current_documents = db.scalars(
        select(Document)
        .join(Document.project)
        .options(joinedload(Document.project))
        .where(
            Document.knowledge_status == "current",
            Document.valid_to.is_(None),
        )
    ).unique().all()
    replacement_candidates = _matching_current_documents(
        document,
        [
            item
            for item in current_documents
            if is_authorized(user, item, item.project)
        ],
    )
    provided_ids = payload.supersedes_document_ids
    if len(provided_ids) != len(set(provided_ids)):
        raise HTTPException(status_code=422, detail="被替代版本不可重复")
    expected_ids = {item.id for item in replacement_candidates}
    if set(provided_ids) != expected_ids:
        raise HTTPException(
            status_code=409,
            detail="待替代版本范围已变化，请刷新审核队列后重新确认",
        )

    now = datetime.now(timezone.utc)
    predecessor = (
        max(
            replacement_candidates,
            key=lambda item: _time_sort_key(
                item.valid_from or item.ingested_at
            ),
        )
        if replacement_candidates
        else None
    )
    touched_project_ids = {
        item.project_id for item in replacement_candidates
    }
    for previous in replacement_candidates:
        previous.knowledge_status = "superseded"
        previous.valid_to = now

    document.knowledge_status = "current"
    document.valid_from = now
    document.valid_to = None
    document.supersedes_document_id = predecessor.id if predecessor else None
    document.project.knowledge_status = "current"
    document.project.confirmed = True
    db.flush()

    for project_id in touched_project_ids - {document.project_id}:
        project = db.get(Project, project_id)
        if project is None:
            continue
        has_current = db.scalar(
            select(Document.id)
            .where(
                Document.project_id == project_id,
                Document.knowledge_status == "current",
                Document.valid_to.is_(None),
            )
            .limit(1)
        )
        has_candidate = db.scalar(
            select(Document.id)
            .where(
                Document.project_id == project_id,
                Document.knowledge_status == "candidate",
            )
            .limit(1)
        )
        project.knowledge_status = (
            "current"
            if has_current
            else "candidate"
            if has_candidate
            else "superseded"
        )

    replaced_ids = sorted(expected_ids)
    db.add(
        AuditLog(
            user_id=user.id,
            action="knowledge_published",
            document_ids_json=json.dumps(
                [document.id, *replaced_ids],
                ensure_ascii=False,
            ),
            details_json=json.dumps(
                {
                    "published_document_id": document.id,
                    "superseded_document_ids": replaced_ids,
                    "supersedes_document_id": document.supersedes_document_id,
                    "valid_from": now.isoformat(),
                    "note_provided": bool(payload.note and payload.note.strip()),
                },
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    return {
        "document_id": document.id,
        "project_id": document.project_id,
        "knowledge_status": "current",
        "valid_from": document.valid_from,
        "supersedes_document_id": document.supersedes_document_id,
        "superseded_documents": [
            _publication_document_payload(item)
            for item in replacement_candidates
        ],
        "notice": (
            f"已发布为当前版本，并将 {len(replacement_candidates)} 个旧版本转为历史。"
        ),
    }


@app.post("/v1/review/{document_id}/reject")
def reject_review_proposal(
    document_id: str,
    payload: ReviewRejectRequest,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    require_founder(user)
    document = _review_document_or_404(document_id, user, db)
    proposal = db.scalar(
        select(ReviewProposal).where(
            ReviewProposal.id == payload.proposal_id,
            ReviewProposal.document_id == document.id,
            ReviewProposal.status == "pending_founder",
        )
    )
    if proposal is None:
        raise HTTPException(status_code=409, detail="确认建议已失效或已处理")
    now = datetime.now(timezone.utc)
    proposal.status = "rejected"
    proposal.decision_note = payload.reason.strip() if payload.reason else None
    proposal.decided_by_user_id = user.id
    proposal.decided_at = now
    proposal.updated_at = now
    db.add(
        AuditLog(
            user_id=user.id,
            action="review_rejected",
            document_ids_json=json.dumps([document.id]),
            details_json=json.dumps(
                {
                    "proposal_id": proposal.id,
                    "reason_provided": bool(proposal.decision_note),
                },
                ensure_ascii=False,
            ),
        )
    )
    db.commit()
    return {
        "proposal": _review_proposal_payload(proposal, db),
        "document_unchanged": True,
        "knowledge_status": document.knowledge_status,
        "notice": "建议已退回，原资料元数据未修改。",
    }


def _proposal_context(
    brief: ProposalBrief,
    user: User,
    db: Session,
) -> tuple[list[str], str, dict, list[dict]]:
    fields = {
        "客户": brief.client,
        "目标": brief.objective,
        "目标受众": brief.audience,
        "预算": brief.budget,
        "地区": brief.geography,
        "周期": brief.duration,
        "交付物": brief.deliverables,
    }
    missing = [name for name, value in fields.items() if not value]
    query = " ".join(
        value for value in [brief.title, brief.client, brief.objective, brief.audience] if value
    )
    evidence = search(db, user=user, query=query, requested_scope="history", limit=5)
    outline = [
        {"page": 1, "title": "项目封面", "purpose": brief.title},
        {"page": 2, "title": "需求理解", "purpose": "复述目标、受众与约束，待客户确认"},
        {"page": 3, "title": "项目目标", "purpose": "将业务目标拆为可验收指标"},
        {"page": 4, "title": "总体方案", "purpose": "说明培训路径、组织方式和核心价值"},
        {"page": 5, "title": "课程与活动设计", "purpose": "按阶段列出课程、实训和成果"},
        {"page": 6, "title": "执行计划", "purpose": "明确里程碑、人员、场地和协同机制"},
        {"page": 7, "title": "评估与结案", "purpose": "定义过程数据、验收和复盘材料"},
        {"page": 8, "title": "风险与待确认项", "purpose": "集中列出缺失信息和边界条件"},
    ]
    return missing, query, evidence, outline


@app.post("/v1/proposals/draft")
def proposal_draft(
    brief: ProposalBrief,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    missing, query, evidence, outline = _proposal_context(brief, user, db)
    audit = AuditLog(
        user_id=user.id,
        action="proposal_draft",
        query_hash=hashlib.sha256(query.encode()).hexdigest(),
        document_ids_json=json.dumps(
            list(dict.fromkeys(item["document_id"] for item in evidence["results"]))
        ),
        details_json=json.dumps({"missing": missing}, ensure_ascii=False),
    )
    db.add(audit)
    db.commit()
    return {
        "draft_id": audit.id,
        "mode": "local-rules-with-retrieved-evidence",
        "missing_information": missing,
        "similar_cases": evidence["results"],
        "outline": outline,
        "notice": "这是本地规则生成的文案骨架；候选资料尚未确认，不得直接作为对外事实。",
    }


@app.post("/v1/proposals/pptx")
def proposal_pptx(
    brief: ProposalBrief,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> Response:
    missing, query, evidence, _outline = _proposal_context(brief, user, db)
    try:
        payload = render_proposal_pptx(brief, evidence["results"], missing)
    except ProposalPptxError as exc:
        raise HTTPException(
            status_code=503,
            detail="PPTX 模板暂不可用，请联系管理员",
        ) from exc
    audit = AuditLog(
        user_id=user.id,
        action="proposal_pptx",
        query_hash=hashlib.sha256(query.encode()).hexdigest(),
        document_ids_json=json.dumps(
            list(dict.fromkeys(item["document_id"] for item in evidence["results"]))
        ),
        details_json=json.dumps(
            {
                "missing": missing,
                "format": "pptx",
                "bytes": len(payload),
            },
            ensure_ascii=False,
        ),
    )
    db.add(audit)
    db.commit()
    safe_title = re.sub(r'[<>:"/\\|?*]+', "-", brief.title).strip(" .-")
    filename = f"{safe_title or '京奥电竞提案'}-初稿.pptx"
    return Response(
        content=payload,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "presentationml.presentation"
        ),
        headers={
            "Content-Disposition": (
                f"attachment; filename*=UTF-8''{quote(filename)}"
            ),
            "Cache-Control": "no-store",
        },
    )


# 产品当前只保留资料治理、检索、创作与运行状态。评测题库和知识进化
# 的历史数据及内部实现继续保留，但对应 HTTP 路由暂不发布，也不会出现在
# OpenAPI 文档中。以后若重新启用，可在经过业务验证后恢复路由注册。
_RETIRED_PRODUCT_API_PREFIXES = ("/v1/evaluations", "/v1/evolution")
app.router.routes[:] = [
    route
    for route in app.router.routes
    if not any(
        getattr(route, "path", "").startswith(prefix)
        for prefix in _RETIRED_PRODUCT_API_PREFIXES
    )
]
